from __future__ import annotations

import csv
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from django.test import SimpleTestCase, override_settings

from apps.services.knowledge_ingestion import KnowledgeIngestionService

try:  # pragma: no cover - optional dependency
    from openpyxl import Workbook
except ImportError:  # pragma: no cover
    Workbook = None


class IngestionNormalizationParityTests(SimpleTestCase):
    @override_settings(INGEST_NORMALIZE_TABLES=True)
    def test_csv_and_xlsx_ingestion_align_after_normalization(self) -> None:
        if Workbook is None:
            self.skipTest("openpyxl not installed")

        rows = [
            ["Name", "Status", "Notes", "EmptyColumn"],
            ["Alice", "NULL", "First record", ""],
            ["Bob", "", "Second row", "N/A"],
            ["", "", "", ""],
            ["Cara", "Ready", "Uses #REF! token", "#REF!"],
        ]

        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            csv_path = root / "sample.csv"
            xlsx_path = root / "sample.xlsx"

            with csv_path.open("w", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerows(rows)

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Plans"
            for row in rows:
                sheet.append(row)
            workbook.save(xlsx_path)

            service = KnowledgeIngestionService(media_root=root)
            csv_detail = SimpleNamespace(
                filename="sample.csv",
                content_type="text/csv",
                storage_path=str(csv_path),
            )
            xlsx_detail = SimpleNamespace(
                filename="sample.xlsx",
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                storage_path=str(xlsx_path),
            )

            csv_result = service._extract_csv(csv_path, format_hint="csv", file_detail=csv_detail, upload=None)
            xlsx_result = service._extract_xlsx(xlsx_path, file_detail=xlsx_detail, upload=None)

            self.assertTrue(csv_result.tables)
            self.assertTrue(xlsx_result.tables)
            csv_table = csv_result.tables[0]
            xlsx_table = xlsx_result.tables[0]

            self.assertEqual(csv_table.column_schema, xlsx_table.column_schema)
            self.assertEqual(self._table_rows(csv_table), self._table_rows(xlsx_table))
            self.assertEqual(csv_result.metadata.get("table_stats"), xlsx_result.metadata.get("table_stats"))

            csv_norm = csv_result.metadata.get("normalization") or {}
            xlsx_norm = xlsx_result.metadata.get("normalization") or {}
            self.assertEqual(
                (csv_norm.get("rows_dropped") or {}).get("total"),
                (xlsx_norm.get("rows_dropped") or {}).get("total"),
            )
            self.assertEqual(
                (csv_norm.get("columns_trimmed") or {}).get("total"),
                (xlsx_norm.get("columns_trimmed") or {}).get("total"),
            )

    @staticmethod
    def _table_rows(table) -> list[list[str]]:
        return [[cell.raw_text for cell in row.cells] for row in table.rows]

