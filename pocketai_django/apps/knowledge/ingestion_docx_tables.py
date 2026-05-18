from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Any, Sequence

from apps.accounts.models import KnowledgeIssueSeverity
from apps.knowledge.ingestion_contracts import IssuePayload, TableCellPayload, TablePayload, TableRowPayload
from apps.knowledge.ingestion_page_renderer import PageRenderer
from apps.knowledge.ingestion_signals import _column_numeric_signal
from apps.knowledge.ingestion_table_detection import TableDetector
from apps.knowledge.ingestion_text_utils import IngestionTextUtilsMixin

try:  # pragma: no cover - dependency failure should be surfaced at runtime
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover - fallback handled via runtime check
    DocxDocument = None  # type: ignore


class IngestionDocxTablesMixin:

    def _extract_docx_table_candidates(
        self,
        path: Path,
        *,
        filename: str = "",
    ) -> tuple[list[TablePayload], list[IssuePayload], dict[str, Any]]:
        """
        Parse DOCX tables directly from `document.tables` and emit structured TablePayloads.

        This is the primary DOCX table path. Paragraph extraction remains supplemental.
        """
        if DocxDocument is None:
            return [], [
                IssuePayload(
                    code="docx_tables_missing_dependency",
                    severity=KnowledgeIssueSeverity.WARNING.value,
                    description="python-docx is unavailable; DOCX table extraction skipped.",
                    page_number=1,
                )
            ], {"enabled": False, "reason": "python_docx_missing"}

        try:
            document = DocxDocument(str(path))
        except Exception as exc:
            return [], [
                IssuePayload(
                    code="docx_tables_parse_failed",
                    severity=KnowledgeIssueSeverity.ERROR.value,
                    description=f"DOCX table extraction failed: {exc}",
                    page_number=1,
                )
            ], {"enabled": False, "reason": "parse_failed"}

        heading_map = self._docx_table_heading_map(document)
        tables: list[TablePayload] = []
        issues: list[IssuePayload] = []
        merged_regions_total = 0
        header_rows_total = 0
        compacted_blank_rows_total = 0
        section_rows_total = 0
        series_columns_total = 0

        for order_index, table in enumerate(document.tables, start=1):
            key_grid, text_by_key, span_by_key = self._docx_table_grid(table)
            key_grid, span_by_key, grid_meta = self._docx_compact_table_grid(
                key_grid,
                text_by_key,
                span_by_key,
            )
            provisional_header_rows = self._docx_detect_header_rows(key_grid, text_by_key)
            key_grid, text_by_key, span_by_key, collapse_meta = self._docx_collapse_helper_columns(
                key_grid,
                text_by_key,
                span_by_key,
                provisional_header_rows,
            )
            key_grid, text_by_key, span_by_key, series_meta = self._docx_normalize_sparse_series_columns(
                key_grid,
                text_by_key,
                span_by_key,
            )
            row_count = len(key_grid)
            col_count = max((len(row) for row in key_grid), default=0)
            if row_count <= 0 or col_count <= 0:
                issues.append(
                    IssuePayload(
                        code="docx_table_empty",
                        severity=KnowledgeIssueSeverity.INFO.value,
                        description=f"DOCX table {order_index} had no readable cells.",
                        page_number=1,
                        table_order_index=order_index,
                    )
                )
                continue

            header_rows = self._docx_detect_header_rows(key_grid, text_by_key)
            header_row_set = set(header_rows)
            section_rows = self._docx_detect_section_rows(key_grid, text_by_key, header_row_set)
            section_row_set = set(section_rows)
            if header_rows:
                header_rows_total += len(header_rows)
            compacted_blank_rows_total += int(grid_meta.get("removed_blank_rows") or 0)
            if section_rows:
                section_rows_total += len(section_rows)
            series_columns_total += int(series_meta.get("merged_columns") or 0)

            column_schema: list[str] = []
            for col_idx in range(col_count):
                labels: list[str] = []
                seen_labels: set[str] = set()
                for row_idx in header_rows:
                    key = key_grid[row_idx][col_idx] if col_idx < len(key_grid[row_idx]) else None
                    if key is None:
                        continue
                    label = self._docx_canonicalize_header_label(
                        self._sanitize_text(text_by_key.get(key, "")).strip()
                    )
                    if not label:
                        continue
                    dedupe_key = re.sub(r"\s+", " ", label).strip().lower()
                    if dedupe_key in seen_labels:
                        continue
                    seen_labels.add(dedupe_key)
                    labels.append(label)
                merged_label = " | ".join(labels).strip()
                column_schema.append(self._docx_normalize_column_key(merged_label, col_idx))

            if not any(column_schema):
                column_schema = [f"column_{idx + 1}" for idx in range(col_count)]
            else:
                column_schema = self._docx_refine_grouped_column_schema(
                    column_schema,
                    key_grid,
                    text_by_key,
                    header_rows,
                )

            table_rows: list[TableRowPayload] = []
            merged_regions = 0
            for row_idx, row_keys in enumerate(key_grid):
                row_type = "data"
                if row_idx in header_row_set:
                    row_type = "header"
                elif row_idx in section_row_set:
                    row_type = "section_header"

                initial_row_values = [
                    self._sanitize_text(
                        text_by_key.get(row_keys[col_idx], "")
                        if col_idx < len(row_keys) and row_keys[col_idx] is not None
                        else ""
                    ).strip()
                    for col_idx in range(col_count)
                ]
                repeated_section_label = ""
                repeated_section_col_idx: int | None = None
                if row_type == "section_header":
                    repeated_section_label, repeated_section_col_idx = self._docx_repeated_row_label_info(
                        initial_row_values
                    )

                row_cells: list[TableCellPayload] = []
                row_values: list[str] = []
                for col_idx in range(col_count):
                    key = row_keys[col_idx] if col_idx < len(row_keys) else None
                    raw_text = self._sanitize_text(text_by_key.get(key, "") if key is not None else "").strip()
                    if repeated_section_label and raw_text:
                        canonical_cell = self._docx_canonicalize_header_label(raw_text)
                        if canonical_cell.strip().lower() == repeated_section_label.strip().lower():
                            raw_text = repeated_section_label if col_idx == repeated_section_col_idx else ""
                    row_values.append(raw_text)
                    cell_metadata: dict[str, Any] = {}
                    if key is not None:
                        span = span_by_key.get(key) or {}
                        row_span = int(span.get("row_span") or 1)
                        col_span = int(span.get("column_span") or 1)
                        if row_span > 1 or col_span > 1:
                            is_anchor = (
                                row_idx == int(span.get("row_start") or 0)
                                and col_idx == int(span.get("col_start") or 0)
                            )
                            cell_metadata["row_span"] = row_span
                            cell_metadata["column_span"] = col_span
                            cell_metadata["merged_anchor"] = is_anchor
                            if is_anchor:
                                merged_regions += 1
                            else:
                                cell_metadata["merged_from"] = {
                                    "row_index": int(span.get("row_start") or 0),
                                    "column_index": int(span.get("col_start") or 0),
                                }

                    row_cells.append(
                        TableCellPayload(
                            row_index=row_idx,
                            column_index=col_idx,
                            column_key=column_schema[col_idx] if col_idx < len(column_schema) else f"column_{col_idx + 1}",
                            raw_text=raw_text,
                            normalized_value=self.table_detector._normalize_cell_value(raw_text),
                            metadata=cell_metadata,
                        )
                    )

                row_metadata: dict[str, Any] = {"row_type": row_type}
                if row_idx in header_row_set:
                    row_metadata["header_source"] = "docx_detected"
                    row_metadata["header_level"] = header_rows.index(row_idx) + 1
                elif row_idx in section_row_set:
                    row_metadata["section_source"] = "docx_detected"
                table_rows.append(
                    TableRowPayload(
                        row_index=row_idx,
                        page_number=1,
                        raw_text="\t".join(row_values),
                        metadata=row_metadata,
                        cells=row_cells,
                    )
                )

            merged_regions_total += merged_regions
            section_heading = self._sanitize_text(heading_map.get(order_index, "")).strip()
            title = section_heading or f"Table {order_index}"
            table_metadata: dict[str, Any] = {
                "detected_via": "docx:table_xml",
                "extractor": "docx_table_parser",
                "table_index": order_index,
                "row_count": row_count,
                "column_count": col_count,
                "header_rows": list(header_rows),
                "section_rows": list(section_rows),
                "merged_regions": merged_regions,
                "structure_confidence": 0.95,
            }
            if grid_meta.get("removed_blank_rows"):
                table_metadata["blank_rows_compacted"] = int(grid_meta.get("removed_blank_rows") or 0)
            if collapse_meta.get("removed_columns"):
                table_metadata["helper_columns_removed"] = int(collapse_meta.get("removed_columns") or 0)
            if collapse_meta.get("merged_columns"):
                table_metadata["helper_columns_merged"] = int(collapse_meta.get("merged_columns") or 0)
            if series_meta.get("merged_columns"):
                table_metadata["series_columns_merged"] = int(series_meta.get("merged_columns") or 0)
            if filename:
                table_metadata["filename"] = filename
            if section_heading:
                table_metadata["section_heading"] = section_heading
            style_name = self._sanitize_text(getattr(getattr(table, "style", None), "name", "")).strip()
            if style_name:
                table_metadata["style_name"] = style_name

            tables.append(
                TablePayload(
                    order_index=order_index,
                    title=title,
                    section_heading=section_heading,
                    page_number=1,
                    column_schema=column_schema,
                    data_dictionary={},
                    metadata=table_metadata,
                    rows=table_rows,
                )
            )

        meta: dict[str, Any] = {
            "enabled": True,
            "table_count": len(tables),
            "header_rows_detected": header_rows_total,
            "section_rows_detected": section_rows_total,
            "blank_rows_compacted": compacted_blank_rows_total,
            "merged_regions": merged_regions_total,
            "series_columns_merged": series_columns_total,
        }
        return tables, issues, meta

    def _docx_table_heading_map(self, document: Any) -> dict[int, str]:
        body = getattr(getattr(document, "element", None), "body", None)
        if body is None:
            return {}
        try:
            from docx.text.paragraph import Paragraph as DocxParagraphClass  # type: ignore
        except Exception:
            return {}

        current_heading = ""
        heading_map: dict[int, str] = {}
        table_index = 0
        for child in body.iterchildren():
            child_tag = str(getattr(child, "tag", "") or "")
            if child_tag.endswith("}p"):
                paragraph = DocxParagraphClass(child, document)
                text = self._sanitize_text(getattr(paragraph, "text", "")).strip()
                if text and PageRenderer._looks_like_heading(text):
                    current_heading = text
                continue
            if child_tag.endswith("}tbl"):
                table_index += 1
                if current_heading:
                    heading_map[table_index] = current_heading
        return heading_map

    @staticmethod
    def _docx_table_grid(
        table: Any,
    ) -> tuple[list[list[int | None]], dict[int, str], dict[int, dict[str, int]]]:
        rows = [list(getattr(row, "cells", []) or []) for row in getattr(table, "rows", [])]
        if not rows:
            return [], {}, {}
        column_count = max((len(cells) for cells in rows), default=0)
        if column_count <= 0:
            return [], {}, {}

        key_grid: list[list[int | None]] = []
        text_by_key: dict[int, str] = {}
        positions_by_key: dict[int, list[tuple[int, int]]] = {}

        for row_idx, row_cells in enumerate(rows):
            row_keys: list[int | None] = []
            for col_idx in range(column_count):
                if col_idx >= len(row_cells):
                    row_keys.append(None)
                    continue
                cell = row_cells[col_idx]
                tc = getattr(cell, "_tc", None)
                key = id(tc) if tc is not None else id(cell)
                row_keys.append(key)
                positions_by_key.setdefault(key, []).append((row_idx, col_idx))
                if key not in text_by_key:
                    text_by_key[key] = IngestionTextUtilsMixin._sanitize_text(getattr(cell, "text", "")).strip()
            key_grid.append(row_keys)

        span_by_key: dict[int, dict[str, int]] = {}
        for key, positions in positions_by_key.items():
            row_start = min(pos[0] for pos in positions)
            row_end = max(pos[0] for pos in positions)
            col_start = min(pos[1] for pos in positions)
            col_end = max(pos[1] for pos in positions)
            span_by_key[key] = {
                "row_start": row_start,
                "row_end": row_end,
                "col_start": col_start,
                "col_end": col_end,
                "row_span": (row_end - row_start) + 1,
                "column_span": (col_end - col_start) + 1,
            }
        return key_grid, text_by_key, span_by_key

    @staticmethod
    def _docx_compact_table_grid(
        key_grid: Sequence[Sequence[int | None]],
        text_by_key: Mapping[int, str],
        span_by_key: Mapping[int, Mapping[str, int]],
    ) -> tuple[list[list[int | None]], dict[int, dict[str, int]], dict[str, Any]]:
        if not key_grid:
            return [], {}, {"removed_blank_rows": 0}

        compacted_grid: list[list[int | None]] = []
        positions_by_key: dict[int, list[tuple[int, int]]] = {}
        removed_blank_rows = 0

        for row in key_grid:
            has_readable_value = any(
                key is not None and str(text_by_key.get(key, "")).strip()
                for key in row
            )
            if not has_readable_value:
                removed_blank_rows += 1
                continue
            new_row = list(row)
            new_row_index = len(compacted_grid)
            compacted_grid.append(new_row)
            for col_idx, key in enumerate(new_row):
                if key is None:
                    continue
                positions_by_key.setdefault(key, []).append((new_row_index, col_idx))

        if not compacted_grid:
            return [], {}, {"removed_blank_rows": removed_blank_rows}

        compacted_spans: dict[int, dict[str, int]] = {}
        for key, positions in positions_by_key.items():
            prior_span = span_by_key.get(key) or {}
            row_start = min(pos[0] for pos in positions)
            row_end = max(pos[0] for pos in positions)
            col_start = min(pos[1] for pos in positions)
            col_end = max(pos[1] for pos in positions)
            compacted_spans[key] = {
                "row_start": row_start,
                "row_end": row_end,
                "col_start": col_start,
                "col_end": col_end,
                "row_span": (row_end - row_start) + 1,
                "column_span": int(prior_span.get("column_span") or ((col_end - col_start) + 1)),
            }

        return compacted_grid, compacted_spans, {"removed_blank_rows": removed_blank_rows}

    @staticmethod
    def _docx_cell_is_helper_token(value: str) -> bool:
        sample = str(value or "").strip()
        if not sample:
            return False
        if re.fullmatch(r"[$€£¥₹]", sample):
            return True
        if sample in {"(", ")", "[", "]"}:
            return True
        return False

    @staticmethod
    def _docx_combine_cell_texts(parts: Sequence[str]) -> str:
        cleaned = [str(part or "").strip() for part in parts if str(part or "").strip()]
        if not cleaned:
            return ""
        combined = cleaned[0]
        currency_tokens = {"$", "€", "£", "¥", "₹"}
        for part in cleaned[1:]:
            if not combined:
                combined = part
                continue
            if part in {")", "]", "%"}:
                combined = combined.rstrip() + part
                continue
            if part in {"(", "["}:
                combined = combined.rstrip() + part
                continue
            if combined.endswith(tuple(currency_tokens)) or combined.endswith(("(", "[")):
                combined = combined.rstrip() + part
                continue
            combined = combined.rstrip() + " " + part
        return combined.strip()

    @staticmethod
    def _docx_collapse_repeated_sequence(parts: Sequence[str]) -> list[str]:
        items = [str(part or "").strip() for part in parts if str(part or "").strip()]
        size = len(items)
        if size <= 1:
            return items
        for chunk_size in range(1, (size // 2) + 1):
            if size % chunk_size != 0:
                continue
            chunk = items[:chunk_size]
            if chunk * (size // chunk_size) == items:
                return chunk
        return items

    @staticmethod
    def _docx_canonicalize_header_label(label: str) -> str:
        cleaned = re.sub(r"\s+", " ", str(label or "").strip())
        if not cleaned:
            return ""
        tokens = cleaned.split(" ")
        collapsed = IngestionDocxTablesMixin._docx_collapse_repeated_sequence(tokens)
        return " ".join(collapsed).strip()

    @staticmethod
    def _docx_repeated_row_label_info(values: Sequence[str]) -> tuple[str, int | None]:
        normalized: list[tuple[int, str, str]] = []
        for idx, value in enumerate(values):
            raw = str(value or "").strip()
            if not raw:
                continue
            canonical = IngestionDocxTablesMixin._docx_canonicalize_header_label(raw)
            if not canonical:
                continue
            normalized.append((idx, canonical, canonical.strip().lower()))
        if len(normalized) < 2:
            return "", None
        lowered = {entry[2] for entry in normalized if entry[2]}
        if len(lowered) != 1:
            return "", None
        first_idx, first_label, _ = normalized[0]
        return first_label, first_idx

    @staticmethod
    def _docx_normalize_column_key(label: str, index: int) -> str:
        normalized = TableDetector._normalize_header_cell(label, index)
        parts = [part for part in str(normalized or "").split("_") if part]
        collapsed = IngestionDocxTablesMixin._docx_collapse_repeated_sequence(parts)
        if collapsed:
            normalized = "_".join(collapsed)
        return normalized or f"column_{index + 1}"

    @staticmethod
    def _docx_column_key_looks_period_like(key: str) -> bool:
        sample = str(key or "").strip().lower()
        if not sample or sample.startswith("column_"):
            return False
        if re.fullmatch(r"(?:19|20)\d{2}", sample):
            return True
        if re.fullmatch(r"\d{1,2}_\d{2,4}", sample):
            return True
        if re.fullmatch(r"(?:q[1-4]|[1-4]q)(?:_(?:fy)?(?:19|20)?\d{2,4})?", sample):
            return True
        if "thereafter" in sample:
            return True
        return bool(
            re.search(r"(?:19|20)\d{2}", sample)
            and re.search(
                r"\b(?:year|years|quarter|quarters|month|months|ended|ending|june|march|september|december)\b",
                sample.replace("_", " "),
            )
        )

    @staticmethod
    def _docx_column_key_is_generic(key: str) -> bool:
        return bool(re.fullmatch(r"column_\d+", str(key or "").strip().lower()))

    @classmethod
    def _docx_refine_grouped_column_schema(
        cls,
        column_schema: Sequence[str],
        key_grid: Sequence[Sequence[int | None]],
        text_by_key: Mapping[int, str],
        header_rows: Sequence[int],
    ) -> list[str]:
        schema = list(column_schema or [])
        if not schema or not key_grid:
            return schema

        row_count = len(key_grid)
        header_row_set = set(header_rows)

        def _value_at(row_idx: int, col_idx: int) -> str:
            if row_idx >= row_count:
                return ""
            row = key_grid[row_idx]
            if col_idx >= len(row):
                return ""
            key = row[col_idx]
            if key is None:
                return ""
            return str(text_by_key.get(key, "")).strip()

        profiles: list[dict[str, Any]] = []
        for col_idx, key in enumerate(schema):
            data_values = [
                _value_at(row_idx, col_idx)
                for row_idx in range(row_count)
                if row_idx not in header_row_set and _value_at(row_idx, col_idx)
            ]
            helper_only = bool(data_values) and all(cls._docx_cell_is_helper_token(value) for value in data_values)
            substantive = any(
                value
                and not cls._docx_cell_is_helper_token(value)
                and (_column_numeric_signal(value) or re.search(r"[A-Za-z\u0600-\u06FF]", value))
                for value in data_values
            )
            profiles.append(
                {
                    "key": str(key or "").strip(),
                    "data_count": len(data_values),
                    "helper_only": helper_only,
                    "substantive": substantive,
                    "period_like": cls._docx_column_key_looks_period_like(str(key or "").strip()),
                    "generic": cls._docx_column_key_is_generic(str(key or "").strip()),
                }
            )

        refined = list(schema)
        for col_idx, profile in enumerate(profiles):
            key = profile["key"]
            if not key or not profile["period_like"] or profile["data_count"] <= 0:
                continue
            if col_idx + 1 >= len(profiles):
                continue
            right = profiles[col_idx + 1]
            right_key = str(right["key"] or "").strip()
            if (
                right_key
                and not right["period_like"]
                and not right["generic"]
                and not right["helper_only"]
                and right["data_count"] == 0
            ):
                refined[col_idx] = cls._docx_normalize_column_key(f"{right_key} {key}", col_idx)

        for col_idx, profile in enumerate(profiles[:-1]):
            key = str(refined[col_idx] or "").strip()
            next_key = str(refined[col_idx + 1] or "").strip()
            if not key or key != next_key or not cls._docx_column_key_looks_period_like(key):
                continue
            next_profile = profiles[col_idx + 1]
            if profile["helper_only"] and next_profile["substantive"]:
                refined[col_idx] = cls._docx_normalize_column_key(f"helper {key}", col_idx)
            elif next_profile["helper_only"] and profile["substantive"]:
                refined[col_idx + 1] = cls._docx_normalize_column_key(f"helper {key}", col_idx + 1)

        return refined

    @staticmethod
    def _docx_value_looks_like_period_label(value: str) -> bool:
        sample = IngestionDocxTablesMixin._docx_canonicalize_header_label(str(value or "").strip())
        if not sample:
            return False
        lowered = sample.lower()
        if re.fullmatch(r"(?:19|20)\d{2}", lowered):
            return True
        if re.fullmatch(r"\d{1,2}/\d{2,4}", lowered):
            return True
        if re.fullmatch(r"(?:q[1-4]|[1-4]q)(?:\s*(?:fy)?\s*(?:19|20)?\d{2,4})?", lowered):
            return True
        if re.fullmatch(
            r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*[\s\-]+(?:19|20)\d{2}",
            lowered,
        ):
            return True
        if (
            re.search(r"\b(?:year|years|quarter|quarters|month|months|ended|ending)\b", lowered)
            and re.search(r"(?:19|20)\d{2}", lowered)
        ):
            return True
        return False

    @staticmethod
    def _docx_normalize_sparse_series_columns(
        key_grid: Sequence[Sequence[int | None]],
        text_by_key: Mapping[int, str],
        span_by_key: Mapping[int, Mapping[str, int]],
    ) -> tuple[list[list[int | None]], dict[int, str], dict[int, dict[str, int]], dict[str, int]]:
        if not key_grid:
            return [], dict(text_by_key), dict(span_by_key), {"merged_columns": 0}

        row_count = len(key_grid)
        column_count = max((len(row) for row in key_grid), default=0)
        if row_count < 2 or column_count < 4:
            return list(map(list, key_grid)), dict(text_by_key), {int(key): dict(value) for key, value in span_by_key.items()}, {"merged_columns": 0}

        def _value_at(row_idx: int, col_idx: int) -> str:
            if row_idx >= row_count:
                return ""
            row = key_grid[row_idx]
            if col_idx >= len(row):
                return ""
            key = row[col_idx]
            if key is None:
                return ""
            return str(text_by_key.get(key, "")).strip()

        series_row_idx: int | None = None
        for candidate_row_idx in range(min(2, row_count)):
            row_values = [_value_at(candidate_row_idx, col_idx) for col_idx in range(column_count)]
            non_empty_values = [value for value in row_values if value]
            if len(non_empty_values) < 3:
                continue
            period_like_count = sum(
                1 for value in non_empty_values if IngestionDocxTablesMixin._docx_value_looks_like_period_label(value)
            )
            if period_like_count < max(2, int(math.ceil(len(non_empty_values) * 0.6))):
                continue
            duplicate_pairs = 0
            for col_idx in range(1, column_count - 1):
                current_value = IngestionDocxTablesMixin._docx_canonicalize_header_label(row_values[col_idx])
                next_value = IngestionDocxTablesMixin._docx_canonicalize_header_label(row_values[col_idx + 1])
                if (
                    current_value
                    and next_value
                    and current_value == next_value
                    and IngestionDocxTablesMixin._docx_value_looks_like_period_label(current_value)
                ):
                    duplicate_pairs += 1
            if duplicate_pairs >= 2:
                series_row_idx = candidate_row_idx
                break

        if series_row_idx is None:
            return list(map(list, key_grid)), dict(text_by_key), {int(key): dict(value) for key, value in span_by_key.items()}, {"merged_columns": 0}

        actions = ["keep"] * column_count
        merge_targets: dict[int, int] = {}
        data_row_count = max(1, row_count - (series_row_idx + 1))

        for col_idx in range(1, column_count - 1):
            current_label = IngestionDocxTablesMixin._docx_canonicalize_header_label(_value_at(series_row_idx, col_idx))
            next_label = IngestionDocxTablesMixin._docx_canonicalize_header_label(_value_at(series_row_idx, col_idx + 1))
            if (
                not current_label
                or not next_label
                or current_label != next_label
                or not IngestionDocxTablesMixin._docx_value_looks_like_period_label(current_label)
            ):
                continue

            left_values = [_value_at(row_idx, col_idx) for row_idx in range(series_row_idx + 1, row_count)]
            right_values = [_value_at(row_idx, col_idx + 1) for row_idx in range(series_row_idx + 1, row_count)]
            left_non_empty = sum(1 for value in left_values if value)
            right_non_empty = sum(1 for value in right_values if value)
            if left_non_empty == right_non_empty:
                continue

            if left_non_empty < right_non_empty:
                sparse_idx, data_idx = col_idx, col_idx + 1
                sparse_non_empty, data_non_empty = left_non_empty, right_non_empty
            else:
                sparse_idx, data_idx = col_idx + 1, col_idx
                sparse_non_empty, data_non_empty = right_non_empty, left_non_empty

            if sparse_non_empty > 1:
                continue
            if data_non_empty < max(2, int(math.ceil(data_row_count * 0.5))):
                continue
            actions[sparse_idx] = "merge"
            merge_targets[sparse_idx] = data_idx

        if not merge_targets:
            return list(map(list, key_grid)), dict(text_by_key), {int(key): dict(value) for key, value in span_by_key.items()}, {"merged_columns": 0}

        groups: dict[int, list[int]] = {}
        for col_idx, action in enumerate(actions):
            if action == "merge":
                target_idx = merge_targets.get(col_idx)
                if target_idx is None:
                    continue
                groups.setdefault(target_idx, []).append(col_idx)
                groups[target_idx].append(target_idx)
                continue
            groups.setdefault(col_idx, []).append(col_idx)

        ordered_targets = [col_idx for col_idx, action in enumerate(actions) if action == "keep"]
        new_key_grid: list[list[int | None]] = []
        new_text_by_key: dict[int, str] = {}
        positions_by_key: dict[int, list[tuple[int, int]]] = {}
        next_synthetic_key = -1

        for row_idx in range(row_count):
            new_row: list[int | None] = []
            for out_col_idx, target_idx in enumerate(ordered_targets):
                source_indices = sorted(set(groups.get(target_idx, [target_idx])))
                parts: list[str] = []
                source_keys: list[int] = []
                for source_idx in source_indices:
                    if source_idx >= len(key_grid[row_idx]):
                        continue
                    key = key_grid[row_idx][source_idx]
                    if key is None:
                        continue
                    source_keys.append(int(key))
                    value = str(text_by_key.get(key, "")).strip()
                    if value:
                        parts.append(value)

                combined = IngestionDocxTablesMixin._docx_combine_cell_texts(parts)
                if row_idx == series_row_idx:
                    combined = IngestionDocxTablesMixin._docx_canonicalize_header_label(combined)
                if not combined:
                    new_row.append(None)
                    continue

                if len(source_keys) == 1 and combined == str(text_by_key.get(source_keys[0], "")).strip():
                    key_to_use = source_keys[0]
                else:
                    key_to_use = next_synthetic_key
                    next_synthetic_key -= 1
                new_text_by_key[key_to_use] = combined
                positions_by_key.setdefault(key_to_use, []).append((row_idx, out_col_idx))
                new_row.append(key_to_use)
            new_key_grid.append(new_row)

        new_span_by_key: dict[int, dict[str, int]] = {}
        for key, positions in positions_by_key.items():
            row_start = min(pos[0] for pos in positions)
            row_end = max(pos[0] for pos in positions)
            col_start = min(pos[1] for pos in positions)
            col_end = max(pos[1] for pos in positions)
            prior_span = span_by_key.get(key) or {}
            new_span_by_key[key] = {
                "row_start": row_start,
                "row_end": row_end,
                "col_start": col_start,
                "col_end": col_end,
                "row_span": int(prior_span.get("row_span") or ((row_end - row_start) + 1)),
                "column_span": int(prior_span.get("column_span") or ((col_end - col_start) + 1)),
            }

        return new_key_grid, new_text_by_key, new_span_by_key, {
            "merged_columns": len(merge_targets),
        }

    @staticmethod
    def _docx_collapse_helper_columns(
        key_grid: Sequence[Sequence[int | None]],
        text_by_key: Mapping[int, str],
        span_by_key: Mapping[int, Mapping[str, int]],
        header_rows: Sequence[int],
    ) -> tuple[list[list[int | None]], dict[int, str], dict[int, dict[str, int]], dict[str, int]]:
        if not key_grid:
            return [], dict(text_by_key), dict(span_by_key), {"removed_columns": 0, "merged_columns": 0}

        row_count = len(key_grid)
        column_count = max((len(row) for row in key_grid), default=0)
        if column_count <= 0:
            return list(map(list, key_grid)), dict(text_by_key), dict(span_by_key), {"removed_columns": 0, "merged_columns": 0}

        header_row_set = set(header_rows)

        def _value_at(row_idx: int, col_idx: int) -> str:
            if row_idx >= row_count:
                return ""
            row = key_grid[row_idx]
            if col_idx >= len(row):
                return ""
            key = row[col_idx]
            if key is None:
                return ""
            return str(text_by_key.get(key, "")).strip()

        column_profiles: list[dict[str, Any]] = []
        for col_idx in range(column_count):
            all_values = [_value_at(row_idx, col_idx) for row_idx in range(row_count)]
            non_empty_values = [value for value in all_values if value]
            header_values = [_value_at(row_idx, col_idx) for row_idx in range(row_count) if row_idx in header_row_set]
            header_values = [value for value in header_values if value]
            data_values = [_value_at(row_idx, col_idx) for row_idx in range(row_count) if row_idx not in header_row_set]
            data_values = [value for value in data_values if value]
            header_fingerprint = tuple(
                re.sub(r"\s+", " ", value).strip().lower()
                for value in header_values
                if value.strip()
            )
            column_profiles.append(
                {
                    "all_values": non_empty_values,
                    "header_values": header_values,
                    "data_values": data_values,
                    "header_fingerprint": header_fingerprint,
                    "helper_only_data": bool(data_values)
                    and all(IngestionDocxTablesMixin._docx_cell_is_helper_token(value) for value in data_values),
                    "substantive_data": any(
                        value
                        and not IngestionDocxTablesMixin._docx_cell_is_helper_token(value)
                        and (
                            _column_numeric_signal(value)
                            or re.search(r"[A-Za-z\u0600-\u06FF]", value)
                        )
                        for value in data_values
                    ),
                }
            )

        actions = ["keep"] * column_count
        merge_targets: dict[int, int] = {}

        for col_idx, profile in enumerate(column_profiles):
            all_values = profile["all_values"]
            data_values = profile["data_values"]
            header_fingerprint = profile["header_fingerprint"]

            if not all_values:
                actions[col_idx] = "drop"
                continue

            if not data_values and not header_fingerprint:
                actions[col_idx] = "drop"
                continue

            if not data_values and header_fingerprint:
                prev_fingerprint = column_profiles[col_idx - 1]["header_fingerprint"] if col_idx > 0 else ()
                next_fingerprint = column_profiles[col_idx + 1]["header_fingerprint"] if col_idx + 1 < column_count else ()
                if header_fingerprint == prev_fingerprint or header_fingerprint == next_fingerprint:
                    actions[col_idx] = "drop"
                    continue

            if not profile["helper_only_data"]:
                continue

            helper_values = {value for value in data_values if value}
            prefer_right = helper_values <= {"$", "€", "£", "¥", "₹", "(", "["}
            prefer_left = helper_values <= {")", "]"}

            target_idx: int | None = None
            candidate_indices: list[int] = []
            if prefer_left and col_idx > 0:
                candidate_indices.append(col_idx - 1)
            if prefer_right and col_idx + 1 < column_count:
                candidate_indices.append(col_idx + 1)
            if not candidate_indices:
                if col_idx + 1 < column_count:
                    candidate_indices.append(col_idx + 1)
                if col_idx > 0:
                    candidate_indices.append(col_idx - 1)

            for candidate_idx in candidate_indices:
                if actions[candidate_idx] == "drop":
                    continue
                if column_profiles[candidate_idx]["substantive_data"]:
                    target_idx = candidate_idx
                    break
            if target_idx is None:
                continue
            actions[col_idx] = "merge"
            merge_targets[col_idx] = target_idx

        groups: dict[int, list[int]] = {}
        for col_idx, action in enumerate(actions):
            if action == "drop":
                continue
            if action == "merge":
                target_idx = merge_targets.get(col_idx)
                if target_idx is None or actions[target_idx] == "drop":
                    actions[col_idx] = "drop"
                    continue
                groups.setdefault(target_idx, []).append(col_idx)
                continue
            groups.setdefault(col_idx, []).append(col_idx)

        for target_idx, source_indices in list(groups.items()):
            unique_sources = sorted(set(source_indices + [target_idx]))
            groups[target_idx] = unique_sources

        ordered_targets = [col_idx for col_idx, action in enumerate(actions) if action == "keep"]
        if len(ordered_targets) == column_count and not any(action != "keep" for action in actions):
            return (
                [list(row) for row in key_grid],
                dict(text_by_key),
                {int(key): dict(value) for key, value in span_by_key.items()},
                {"removed_columns": 0, "merged_columns": 0},
            )

        new_key_grid: list[list[int | None]] = []
        new_text_by_key: dict[int, str] = {}
        positions_by_key: dict[int, list[tuple[int, int]]] = {}
        next_synthetic_key = -1

        for row_idx in range(row_count):
            new_row: list[int | None] = []
            for out_col_idx, target_idx in enumerate(ordered_targets):
                source_indices = groups.get(target_idx, [target_idx])
                parts: list[str] = []
                source_keys: list[int] = []
                for source_idx in source_indices:
                    if source_idx >= len(key_grid[row_idx]):
                        continue
                    key = key_grid[row_idx][source_idx]
                    if key is None:
                        continue
                    value = str(text_by_key.get(key, "")).strip()
                    if value:
                        parts.append(value)
                    source_keys.append(int(key))

                combined = IngestionDocxTablesMixin._docx_combine_cell_texts(parts)
                if row_idx in header_row_set:
                    combined = IngestionDocxTablesMixin._docx_canonicalize_header_label(combined)
                if not combined:
                    new_row.append(None)
                    continue

                if len(source_keys) == 1 and combined == str(text_by_key.get(source_keys[0], "")).strip():
                    key_to_use = source_keys[0]
                else:
                    key_to_use = next_synthetic_key
                    next_synthetic_key -= 1
                new_text_by_key[key_to_use] = combined
                positions_by_key.setdefault(key_to_use, []).append((row_idx, out_col_idx))
                new_row.append(key_to_use)
            new_key_grid.append(new_row)

        new_span_by_key: dict[int, dict[str, int]] = {}
        for key, positions in positions_by_key.items():
            if key in span_by_key and key >= 0:
                prior_span = span_by_key.get(key) or {}
                row_start = min(pos[0] for pos in positions)
                row_end = max(pos[0] for pos in positions)
                col_start = min(pos[1] for pos in positions)
                col_end = max(pos[1] for pos in positions)
                new_span_by_key[key] = {
                    "row_start": row_start,
                    "row_end": row_end,
                    "col_start": col_start,
                    "col_end": col_end,
                    "row_span": int(prior_span.get("row_span") or ((row_end - row_start) + 1)),
                    "column_span": int(prior_span.get("column_span") or ((col_end - col_start) + 1)),
                }
                continue
            row_start = min(pos[0] for pos in positions)
            row_end = max(pos[0] for pos in positions)
            col_start = min(pos[1] for pos in positions)
            col_end = max(pos[1] for pos in positions)
            new_span_by_key[key] = {
                "row_start": row_start,
                "row_end": row_end,
                "col_start": col_start,
                "col_end": col_end,
                "row_span": (row_end - row_start) + 1,
                "column_span": (col_end - col_start) + 1,
            }

        removed_columns = sum(1 for action in actions if action == "drop")
        merged_columns = sum(1 for action in actions if action == "merge")
        return new_key_grid, new_text_by_key, new_span_by_key, {
            "removed_columns": removed_columns,
            "merged_columns": merged_columns,
        }

    @staticmethod
    def _docx_row_has_financial_data_signal(values: Sequence[str]) -> bool:
        for value in values:
            sample = str(value or "").strip()
            if not sample:
                continue
            if "$" in sample or "%" in sample:
                return True
            if re.search(r"\b\d{1,3}(?:,\d{3})+(?:\.\d+)?\b", sample):
                return True
            if re.search(r"\(\s*\d", sample):
                return True
        return False

    @staticmethod
    def _docx_row_looks_like_header_band(values: Sequence[str]) -> bool:
        non_empty = [
            IngestionDocxTablesMixin._docx_canonicalize_header_label(str(value or "").strip())
            for value in values
            if str(value or "").strip()
        ]
        if not non_empty:
            return False
        if IngestionDocxTablesMixin._docx_row_has_financial_data_signal(non_empty):
            return False

        short_cell_ratio = sum(
            1
            for value in non_empty
            if len(re.findall(r"\w+", value)) <= 4 and len(value) <= 40
        ) / float(max(1, len(non_empty)))
        if short_cell_ratio < 0.6:
            return False

        lowered = [value.lower() for value in non_empty]
        period_or_header_terms = sum(
            1
            for value in lowered
            if re.search(
                r"\b(year|years|ended|ending|quarter|quarters|fiscal|period|periods|date|dates|record|payment|declaration|month|months|june|march|september|december|thereafter)\b",
                value,
            )
        )
        explicit_year_cells = sum(
            1
            for value in non_empty
            if re.fullmatch(r"(?:19|20)\d{2}", value)
        )
        period_label_cells = sum(
            1
            for value in non_empty
            if IngestionDocxTablesMixin._docx_value_looks_like_period_label(value)
        )
        if explicit_year_cells >= max(1, len(non_empty) // 2):
            return True
        if period_label_cells >= max(2, int(math.ceil(len(non_empty) * 0.6))):
            return True
        if len(non_empty) == 1 and period_or_header_terms > 0:
            return True
        return period_or_header_terms > 0

    @staticmethod
    def _docx_detect_header_rows(
        key_grid: Sequence[Sequence[int | None]],
        text_by_key: Mapping[int, str],
    ) -> list[int]:
        row_count = len(key_grid)
        if row_count <= 1:
            return []

        def _row_features(row_keys: Sequence[int | None]) -> dict[str, float]:
            values = [
                str(text_by_key.get(key, "")).strip()
                for key in row_keys
                if key is not None and str(text_by_key.get(key, "")).strip()
            ]
            if not values:
                return {"non_empty": 0.0, "alpha_ratio": 0.0, "numeric_ratio": 0.0}
            alpha_count = sum(1 for value in values if re.search(r"[A-Za-z\u0600-\u06FF]", value))
            numeric_count = sum(1 for value in values if _column_numeric_signal(value))
            total = max(1, len(values))
            return {
                "non_empty": float(len(values)),
                "alpha_ratio": float(alpha_count) / total,
                "numeric_ratio": float(numeric_count) / total,
            }

        probe_rows = min(4, row_count)
        header_rows: list[int] = []
        for row_idx in range(probe_rows):
            values = [
                str(text_by_key.get(key, "")).strip()
                for key in key_grid[row_idx]
                if key is not None and str(text_by_key.get(key, "")).strip()
            ]
            stats = _row_features(key_grid[row_idx])
            header_band = IngestionDocxTablesMixin._docx_row_looks_like_header_band(values)
            if stats["non_empty"] <= 0:
                if row_idx == 0:
                    continue
                break
            looks_header = stats["alpha_ratio"] >= 0.5 and stats["numeric_ratio"] <= 0.5
            if row_idx == 0:
                if looks_header or header_band or stats["numeric_ratio"] < 0.8:
                    header_rows.append(row_idx)
                continue
            multi_value_header = looks_header and len(values) >= 2
            if (multi_value_header or header_band) and header_rows:
                header_rows.append(row_idx)
                continue
            break

        if len(header_rows) >= row_count:
            header_rows = header_rows[: max(1, row_count - 1)]
        if header_rows:
            next_idx = header_rows[-1] + 1
            if next_idx < row_count - 1 and next_idx not in header_rows:
                next_values = [
                    str(text_by_key.get(key, "")).strip()
                    for key in key_grid[next_idx]
                    if key is not None and str(text_by_key.get(key, "")).strip()
                ]
                if (
                    len(next_values) == 1
                    and IngestionDocxTablesMixin._docx_value_looks_like_period_label(next_values[0])
                    and not IngestionDocxTablesMixin._docx_row_has_financial_data_signal(next_values)
                ):
                    following_stats = _row_features(key_grid[next_idx + 1])
                    if following_stats["non_empty"] >= 1:
                        header_rows.append(next_idx)
        return header_rows

    @staticmethod
    def _docx_detect_section_rows(
        key_grid: Sequence[Sequence[int | None]],
        text_by_key: Mapping[int, str],
        header_row_set: set[int],
    ) -> list[int]:
        row_count = len(key_grid)
        if row_count <= 2:
            return []

        def _non_empty_values(row_idx: int) -> list[str]:
            return [
                str(text_by_key.get(key, "")).strip()
                for key in key_grid[row_idx]
                if key is not None and str(text_by_key.get(key, "")).strip()
            ]

        section_rows: list[int] = []
        for row_idx in range(row_count):
            if row_idx in header_row_set:
                continue
            values = _non_empty_values(row_idx)
            repeated_label, _ = IngestionDocxTablesMixin._docx_repeated_row_label_info(values)
            if len(values) != 1 and not repeated_label:
                continue
            label = repeated_label or values[0]
            normalized = label.strip().lower()
            if not normalized:
                continue
            if re.search(r"\d", normalized) and not IngestionDocxTablesMixin._docx_value_looks_like_period_label(label):
                continue
            if normalized in {"total", "subtotal", "totals"}:
                continue
            if len(re.findall(r"\w+", label)) > 10 or len(label) > 80:
                continue
            prev_counts = [
                len(_non_empty_values(candidate))
                for candidate in range(max(0, row_idx - 2), row_idx)
                if candidate not in header_row_set
            ]
            next_counts = [
                len(_non_empty_values(candidate))
                for candidate in range(row_idx + 1, min(row_count, row_idx + 3))
                if candidate not in header_row_set
            ]
            if not next_counts:
                continue
            prev_supports_section = max(prev_counts, default=0) >= 2 or not prev_counts
            if prev_supports_section and max(next_counts, default=0) >= 2:
                section_rows.append(row_idx)
        return section_rows
