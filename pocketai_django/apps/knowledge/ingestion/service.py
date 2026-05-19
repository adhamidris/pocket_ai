from __future__ import annotations

from collections import Counter, deque
import base64
import csv
import gzip
import hashlib
import json
import logging
import math
import mimetypes
import io
import os
import random
import re
import shutil
import statistics
import time
import uuid
import unicodedata
from urllib.parse import urlencode
from datetime import datetime, timedelta, timezone as datetime_timezone
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional, Sequence

from django.conf import settings
from django.db import connection
from django.db import transaction
from django.db.models import Case, IntegerField, Prefetch, Q, Value, When
from django.utils import timezone
from django.utils.text import slugify
from core.otel import otel_trace

from apps.accounts.models import (
    KnowledgeIngestionJobStatus,
    KnowledgeIngestionJobType,
    KnowledgeSourceType,
    KnowledgeStatus,
    KnowledgeVisibility,
    KnowledgeBlockType,
)
from apps.knowledge.models import (
    KnowledgeIngestionJob,
    KnowledgeUpload,
    KnowledgeUploadChunk,
    KnowledgeUploadShadowChunk,
    KnowledgeUploadFile,
    KnowledgeUploadText,
    KnowledgeUploadIssue,
    KnowledgeUploadPage,
    KnowledgeUploadPageBlock,
    KnowledgeUploadTable,
    KnowledgeTableColumn,
    KnowledgeUploadTableCell,
    KnowledgeUploadTableRow,
    KnowledgeEntity,
    KnowledgeAlias,
)
from apps.knowledge.lexicon.learning import TenantLexiconAutoLearningService
from apps.knowledge.documents import DocumentScrapeError, scrape_document_source
from apps.knowledge.datasets.cards import build_dataset_card_segment_payload
from apps.knowledge.datasets.key_index import (
    BloomFilter,
    bloom_spec_for_items,
    normalize_identifier_value,
    resolve_key_index_storage_path,
    write_bloom_filter,
)
from apps.knowledge.ingestion.aliases import (
    ALIAS_KEYWORDS,
    ALIAS_MAX_LENGTH,
    ALIAS_MIN_LENGTH,
    ALIAS_SYMBOL_MIN_LENGTH,
    DATE_TOKEN_PATTERN,
    IDENTIFIER_TOKEN_PATTERN,
    ID_LINE_PATTERN,
)
from apps.knowledge.ingestion.azure_di import AzureDocumentIntelligenceExtractor
from apps.knowledge.ingestion.chunks import IngestionChunksMixin
from apps.knowledge.ingestion.contracts import (
    EnhancedContextDocument,
    ExtractionResult,
    KnowledgeIngestionError,
    PageLayout,
    PageRendererResult,
    PdfSpan,
    TableCellPayload,
    TableRowPayload,
    _union_bbox,
)
from apps.knowledge.tables.docx_tables.extraction import IngestionDocxTablesMixin
from apps.knowledge.ingestion.embeddings import IngestionEmbeddingsMixin
from apps.knowledge.ingestion.entities import IngestionEntitiesMixin
from apps.knowledge.tables.geometry_tools.reconstructor import GeometryTableReconstructor
from apps.knowledge.ingestion.extraction import IngestionExtractionMixin
from apps.knowledge.ingestion.file_formats import IngestionFileFormatsMixin
from apps.knowledge.ingestion.jobs import IngestionJobResult, get_ingestion_queue_health, queue_ingestion_job
from apps.knowledge.ingestion.job_processing import IngestionJobProcessingMixin
from apps.knowledge.ingestion.json_files import IngestionJsonFilesMixin
from apps.knowledge.ingestion.link_extraction import IngestionLinkExtractionMixin
from apps.knowledge.ingestion.ocr import OCRReconciler, create_ocr_reconciler, create_tesseract_ocr_callable
from apps.knowledge.ingestion.page_renderer import PageRenderer
from apps.knowledge.ingestion.persistence import IngestionPersistenceMixin
from apps.knowledge.tables.pdf.geometry import IngestionPdfTableGeometryMixin
from apps.knowledge.tables.pdf.promotion import IngestionPdfTablePromotionMixin
from apps.knowledge.tables.pdf.reconstruction import IngestionPdfTableReconstructionMixin
from apps.knowledge.tables.pdf.routing import IngestionPdfTableRoutingMixin
from apps.knowledge.ingestion.pdfplumber import PdfPlumberTableExtractor
from apps.knowledge.ingestion.signals import (
    COLUMN_ROLE_INFERENCE_VERSION,
    OCR_NORMALIZATION_VERSION,
    TABLE_SCOPE_CONTRACT_VERSION,
    _ARABIC_CHAR_RE,
    _ARABIC_DIACRITICS_RE,
    _SPREADSHEET_CONTROL_CELL_RE,
    _SPREADSHEET_INSTRUCTION_SHEET_RE,
    _SPREADSHEET_INSTRUCTION_TOKEN_RE,
    _SPREADSHEET_MASKED_PLACEHOLDER_RE,
    _SPREADSHEET_PLACEHOLDER_CELL_RE,
    _SPREADSHEET_PURE_NUMBER_RE,
    _SPREADSHEET_RECORD_ID_RE,
    _SPREADSHEET_REFERENCE_SHEET_RE,
    _SPREADSHEET_SUMMARY_ROW_RE,
    _SPREADSHEET_ZERO_LIKE_RE,
    _TABLE_DATE_TIME_LIKE_RE,
    _TABLE_NUMBER_LIKE_RE,
    _TABLE_NUMBER_WITH_UNIT_RE,
    _TABLE_NUMERIC_SIGNAL_TOKEN_RE,
    _TABLE_ROW_VALUE_KEYWORD_RE,
    _column_numeric_signal,
    _infer_contextual_column_indices,
    _log_normalization_summary,
)
from apps.knowledge.tables.selection import IngestionTableSelectionMixin
from apps.knowledge.tables.detection import TableDetector
from apps.knowledge.tables.postprocess.aggregator import IngestionTablePostprocessingMixin
from apps.knowledge.tables.quality import IngestionTableQualityMixin
from apps.knowledge.tables.limits import IngestionTableLimitsMixin
from apps.knowledge.tables.semantic.aggregator import IngestionTableSemanticsMixin
from apps.knowledge.tables.vlm.repair import IngestionTableVlmRepairMixin
from apps.knowledge.tables.tabular.aggregator import IngestionTabularFilesMixin
from apps.knowledge.ingestion.text_utils import IngestionTextUtilsMixin
from apps.rag.embeddings import LocalEmbeddingService, build_embedding_service, EmbeddingProviderError
from apps.accounts.feature_flags import FeatureFlagService
from apps.rag.quality_monitor import QualityMonitor
from apps.rag.rag_logging import structured_log
from apps.rag.table_semantics import normalize_column_name
from core.tenancy import tenant_context
from apps.core.logging_utils import log_start, log_success, log_progress, log_warning, log_error, LogEmoji
from apps.knowledge.tables.normalization import (
    NormalizedSheet,
    SpreadsheetRowInput,
    SheetNormalizationDiagnostics,
    normalize_sheet_rows,
    resolve_normalization_policy,
    sheet_is_allowed,
    summarize_normalization,
)
from apps.knowledge.tables.semantic.column_roles import (
    COLUMN_ROLE_DESCRIPTOR,
    COLUMN_ROLE_NOTE,
    COLUMN_ROLE_QUALIFIER,
    COLUMN_ROLE_SCOPE_DIMENSION,
    column_role_groups,
    column_role_payloads,
    infer_column_roles,
    role_lookup_by_index,
)
from apps.knowledge.tables.semantic.scope_engine import (
    SCOPE_ENGINE_VERSION,
    SCOPE_REASON_ABSTAIN,
    build_scope_table_profile,
    canonical_scope_reason,
    infer_scope_for_row,
)

logger = logging.getLogger(__name__)
TRACER = otel_trace.get_tracer(__name__)


try:  # pragma: no cover - dependency failure should be surfaced at runtime
    import fitz  # type: ignore[attr-defined]  # PyMuPDF
except ImportError:  # pragma: no cover - optional dependency
    fitz = None  # type: ignore

try:  # pragma: no cover - dependency failure should be surfaced at runtime
    from pypdf import PdfReader
except ImportError:  # pragma: no cover - fallback handled via runtime check
    PdfReader = None  # type: ignore

try:  # pragma: no cover - dependency failure should be surfaced at runtime
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover - fallback handled via runtime check
    DocxDocument = None  # type: ignore

try:  # pragma: no cover - dependency failure should be surfaced at runtime
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - fallback handled via runtime check
    load_workbook = None  # type: ignore

try:  # pragma: no cover - dependency failure should be surfaced at runtime
    import xlrd
except ImportError:  # pragma: no cover - fallback handled via runtime check
    xlrd = None  # type: ignore



class KnowledgeIngestionService(
    IngestionJobProcessingMixin,
    IngestionExtractionMixin,
    IngestionDocxTablesMixin,
    IngestionPdfTableGeometryMixin,
    IngestionPdfTableRoutingMixin,
    IngestionPdfTableReconstructionMixin,
    IngestionTableSelectionMixin,
    IngestionTableVlmRepairMixin,
    IngestionLinkExtractionMixin,
    IngestionPersistenceMixin,
    IngestionPdfTablePromotionMixin,
    IngestionChunksMixin,
    IngestionEntitiesMixin,
    IngestionEmbeddingsMixin,
    IngestionTableQualityMixin,
    IngestionTablePostprocessingMixin,
    IngestionTabularFilesMixin,
    IngestionJsonFilesMixin,
    IngestionTableLimitsMixin,
    IngestionTableSemanticsMixin,
    IngestionTextUtilsMixin,
    IngestionFileFormatsMixin,
):
    """
    Pulled-text ingestion pipeline for PDF/DOCX/TXT uploads and external links.

    Designed to run inside a management command or async worker. Fetches queued jobs,
    extracts text, and persists normalized content so the orchestrator and dashboard
    can serve full document context.
    """

    def __init__(self, *, media_root: Path | None = None, enable_ocr: bool = True):
        root = media_root or getattr(settings, "MEDIA_ROOT", None)
        if not root:
            raise RuntimeError("MEDIA_ROOT must be configured for ingestion.")
        self.media_root = Path(root).resolve()
        self.embedding_service = build_embedding_service()
        self._fallback_embedding_service: LocalEmbeddingService | None = None
        self.ingest_inline_chunk_limit = max(0, int(getattr(settings, "INGEST_SYNC_EMBED_CHUNK_LIMIT", 200)))
        self.embedding_batch_size = max(16, int(getattr(settings, "INGEST_EMBED_BATCH_SIZE", 64)))
        self.embedding_job_payload_size = max(self.embedding_batch_size * 4, 256)
        self.ingest_concurrency_limit = max(0, int(getattr(settings, "INGEST_MAX_ACTIVE_JOBS_PER_BUSINESS", 0)))
        self.embedding_backlog_threshold = max(0, int(getattr(settings, "INGEST_EMBEDDING_BACKLOG_THRESHOLD", 500)))
        self.embedding_prewarm_limit = max(0, int(getattr(settings, "INGEST_EMBED_PREWARM_CHUNK_LIMIT", 32)))
        self._fallback_embedding_attempted = False
        
        ocr_render_dpi = int(getattr(settings, "RAG_OCR_RENDER_DPI", 200) or 200)

        # NEW: Create OCR reconciler with Tesseract support
        self.ocr_reconciler = create_ocr_reconciler(enable_ocr=enable_ocr, render_dpi=ocr_render_dpi)
        
        self.page_renderer = PageRenderer(pymupdf_module=fitz)
        self.table_detector = TableDetector()
        self.pdfplumber_enabled = bool(getattr(settings, "RAG_PDFPLUMBER_ENABLED", True))
        self.pdf_table_extractor = str(
            getattr(settings, "RAG_PDF_TABLE_EXTRACTOR", "auto") or "auto"
        ).lower()
        self.table_selection_mode = str(
            getattr(settings, "RAG_TABLE_SELECTION_MODE", "scored_promotion_v2") or "scored_promotion_v2"
        ).strip().lower()
        self.pdfplumber_table_settings = self._normalize_pdfplumber_settings(
            getattr(settings, "RAG_PDFPLUMBER_TABLE_SETTINGS", None)
        )
        self.azure_di_enabled = bool(getattr(settings, "RAG_AZURE_DI_ENABLED", True))
        self.azure_di_endpoint = getattr(settings, "AZURE_DOCUMENT_INTELLIGENCE_ENDPOINT", None)
        self.azure_di_key = getattr(settings, "AZURE_DOCUMENT_INTELLIGENCE_KEY", None)
        self.azure_di_model = getattr(settings, "AZURE_DOCUMENT_INTELLIGENCE_MODEL", "prebuilt-layout")
        self.azure_di_api_version = getattr(settings, "AZURE_DOCUMENT_INTELLIGENCE_API_VERSION", "2023-07-31")
        self.azure_di_base_path = getattr(settings, "AZURE_DOCUMENT_INTELLIGENCE_BASE_PATH", "formrecognizer")
        self.azure_di_locale = str(getattr(settings, "AZURE_DOCUMENT_INTELLIGENCE_LOCALE", "") or "").strip()
        self.azure_di_timeout_seconds = float(getattr(settings, "RAG_AZURE_DI_TIMEOUT_SECONDS", 60.0))
        self.azure_di_poll_interval_seconds = float(getattr(settings, "RAG_AZURE_DI_POLL_INTERVAL_SECONDS", 1.5))
        self.azure_di_max_polls = int(getattr(settings, "RAG_AZURE_DI_MAX_POLLS", 40))
        self.azure_di_request_max_attempts = int(getattr(settings, "RAG_AZURE_DI_REQUEST_MAX_ATTEMPTS", 3))
        self.azure_di_poll_request_max_attempts = int(
            getattr(settings, "RAG_AZURE_DI_POLL_REQUEST_MAX_ATTEMPTS", 3)
        )
        self.azure_di_retry_backoff_base_seconds = float(
            getattr(settings, "RAG_AZURE_DI_RETRY_BACKOFF_BASE_SECONDS", 1.0)
        )
        self.azure_di_retry_backoff_max_seconds = float(
            getattr(settings, "RAG_AZURE_DI_RETRY_BACKOFF_MAX_SECONDS", 8.0)
        )
        self.azure_di_max_retry_after_seconds = float(
            getattr(settings, "RAG_AZURE_DI_MAX_RETRY_AFTER_SECONDS", 30.0)
        )
        self.table_vlm_enabled = bool(getattr(settings, "RAG_TABLE_VLM_ENABLED", True))
        self.table_vlm_model = str(getattr(settings, "RAG_TABLE_VLM_MODEL", "gpt-4o") or "gpt-4o").strip()
        self.table_vlm_confidence_threshold = float(
            getattr(settings, "RAG_TABLE_VLM_CONFIDENCE_THRESHOLD", 0.6)
        )
        self.table_vlm_max_repairs = int(getattr(settings, "RAG_TABLE_VLM_MAX_REPAIRS_PER_UPLOAD", 3))
        self.table_vlm_guardrails_enabled = bool(getattr(settings, "RAG_TABLE_VLM_GUARDRAILS_ENABLED", True))
        self.table_vlm_guardrail_min_row_recall = float(
            getattr(settings, "RAG_TABLE_VLM_GUARDRAIL_MIN_ROW_RECALL", 0.99)
        )
        self.table_vlm_guardrail_hard_row_recall_floor = float(
            getattr(settings, "RAG_TABLE_VLM_GUARDRAIL_HARD_ROW_RECALL_FLOOR", 0.75)
        )
        self.table_vlm_guardrail_min_order_ratio = float(
            getattr(settings, "RAG_TABLE_VLM_GUARDRAIL_MIN_ORDER_RATIO", 0.7)
        )
        self.table_vlm_guardrail_min_schema_recall = float(
            getattr(settings, "RAG_TABLE_VLM_GUARDRAIL_MIN_SCHEMA_RECALL", 0.9)
        )
        self.table_vlm_guardrail_min_cell_recall = float(
            getattr(settings, "RAG_TABLE_VLM_GUARDRAIL_MIN_CELL_RECALL", 0.9)
        )
        self.table_schema_chunking = bool(getattr(settings, "RAG_TABLE_SCHEMA_CHUNKING", True))
        self.table_parent_max_rows = max(1, int(getattr(settings, "RAG_TABLE_PARENT_MAX_ROWS", 200)))
        self.table_parent_max_chars = max(2000, int(getattr(settings, "RAG_TABLE_PARENT_MAX_CHARS", 16000)))
        self.table_child_max_rows = max(0, int(getattr(settings, "RAG_TABLE_CHILD_MAX_ROWS", 500)))
        self.table_summary_enabled = bool(getattr(settings, "RAG_TABLE_SUMMARY_ENABLED", True))
        self.table_summary_max_row_labels = max(0, int(getattr(settings, "RAG_TABLE_SUMMARY_MAX_ROW_LABELS", 50)))
        self.table_header_propagation_enabled = bool(
            getattr(settings, "RAG_TABLE_HEADER_PROPAGATION_ENABLED", True)
        )
        self.table_header_propagation_min_overlap = float(
            getattr(settings, "RAG_TABLE_HEADER_PROPAGATION_MIN_OVERLAP", 0.45)
        )
        self.table_dedupe_enabled = bool(getattr(settings, "RAG_TABLE_DEDUPE_ENABLED", True))
        self.table_dedupe_min_overlap = float(getattr(settings, "RAG_TABLE_DEDUPE_MIN_OVERLAP", 0.6))
        self.pdf_structural_acceptance_enabled = bool(
            getattr(settings, "RAG_PDF_STRUCTURAL_ACCEPTANCE_ENABLED", True)
        )
        self.pdf_collapsed_matrix_reconstruction_enabled = bool(
            getattr(settings, "RAG_PDF_COLLAPSED_MATRIX_RECONSTRUCTION_ENABLED", True)
        )
        self.table_postprocess_row_limit = max(
            5, int(getattr(settings, "RAG_TABLE_POSTPROCESS_ROW_LIMIT", 40))
        )
        self.table_row_signal_min_pairs = max(
            1,
            int(getattr(settings, "RAG_TABLE_ROW_SIGNAL_MIN_PAIRS", 2)),
        )
        self.table_row_signal_min_score = float(
            getattr(settings, "RAG_TABLE_ROW_SIGNAL_MIN_SCORE", 1.6)
        )
        if self.table_row_signal_min_score < 0.0:
            self.table_row_signal_min_score = 0.0
        self.ocr_normalization_enabled = bool(getattr(settings, "RAG_OCR_NORMALIZATION_ENABLED", True))
        self.ocr_word_replacements = self._compile_ocr_replacements(
            getattr(settings, "RAG_OCR_NORMALIZATION_REPLACEMENTS", None)
        )
        self.ocr_percent_fix_enabled = bool(getattr(settings, "RAG_OCR_PERCENT_FIX_ENABLED", True))
        self.ocr_percent_space_fix_enabled = bool(getattr(settings, "RAG_OCR_PERCENT_SPACE_FIX_ENABLED", True))
        self.ocr_percent_sanity_max = float(getattr(settings, "RAG_OCR_PERCENT_SANITY_MAX", 100.0))
        self.ocr_currency_spacing_enabled = bool(getattr(settings, "RAG_OCR_CURRENCY_SPACING_ENABLED", True))
        self.default_json_entity_limit = max(
            1,
            int(getattr(settings, "INGEST_MAX_JSON_ENTITIES_DEFAULT", 200)),
        )
        candidate_cap = int(getattr(settings, "INGEST_MAX_JSON_ENTITY_CANDIDATES", 0)) or (
            self.default_json_entity_limit * 4
        )
        self.max_json_entity_candidates = max(self.default_json_entity_limit, candidate_cap)
        self.default_table_max_rows = max(1, int(getattr(settings, "TABLE_MAX_ROWS_DEFAULT", 5000)))
        self.default_table_max_columns = max(0, int(getattr(settings, "TABLE_MAX_COLUMNS_DEFAULT", 0) or 0))
        self.alias_warning_threshold = int(getattr(settings, "INGEST_ALIAS_WARNING_THRESHOLD", 2000))
        self.chunk_quality_min_tokens = max(1, int(getattr(settings, "RAG_CHUNK_MIN_TOKENS", 20)))
        self.chunk_quality_min_unique_ratio = float(getattr(settings, "RAG_CHUNK_MIN_UNIQUE_RATIO", 0.35))
        if not (0.0 <= self.chunk_quality_min_unique_ratio <= 1.0):
            self.chunk_quality_min_unique_ratio = 0.35
        self.chunk_quality_low_score = float(getattr(settings, "RAG_CHUNK_LOW_QUALITY_SCORE", 0.45))
        if not (0.0 <= self.chunk_quality_low_score <= 1.0):
            self.chunk_quality_low_score = 0.45
        self.native_pdf_text_density_floor = float(
            getattr(settings, "RAG_NATIVE_PDF_TEXT_DENSITY_FLOOR", 0.00005)
        )
        if self.native_pdf_text_density_floor < 0.0:
            self.native_pdf_text_density_floor = 0.00005
        self.native_pdf_pdfplumber_score_boost = float(
            getattr(settings, "RAG_NATIVE_PDF_PDFPLUMBER_SCORE_BOOST", 1.35)
        )
        if self.native_pdf_pdfplumber_score_boost < 1.0:
            self.native_pdf_pdfplumber_score_boost = 1.0
        self.native_pdf_low_readability_penalty = float(
            getattr(settings, "RAG_NATIVE_PDF_LOW_READABILITY_PENALTY", 0.7)
        )
        if not (0.1 <= self.native_pdf_low_readability_penalty <= 1.0):
            self.native_pdf_low_readability_penalty = 0.7
        self.page_block_flat_text_preference_margin = float(
            getattr(settings, "RAG_PAGE_BLOCK_FLAT_TEXT_PREFERENCE_MARGIN", 0.12)
        )
        if self.page_block_flat_text_preference_margin < 0.0:
            self.page_block_flat_text_preference_margin = 0.12
        self.page_block_fragmentation_min_blocks = max(
            4,
            int(getattr(settings, "RAG_PAGE_BLOCK_FRAGMENTATION_MIN_BLOCKS", 8)),
        )
        self.page_block_fragmentation_short_token_limit = max(
            2,
            int(getattr(settings, "RAG_PAGE_BLOCK_FRAGMENTATION_SHORT_TOKEN_LIMIT", 6)),
        )
        self.page_block_fragmentation_short_ratio = float(
            getattr(settings, "RAG_PAGE_BLOCK_FRAGMENTATION_SHORT_RATIO", 0.45)
        )
        if not (0.0 <= self.page_block_fragmentation_short_ratio <= 1.0):
            self.page_block_fragmentation_short_ratio = 0.45
        self.evidence_grouping_enabled = bool(getattr(settings, "RAG_EVIDENCE_GROUPING_ENABLED", True))
        self.evidence_text_link_max_chars = max(
            200,
            int(getattr(settings, "RAG_EVIDENCE_TEXT_LINK_MAX_CHARS", 650)),
        )
        self.evidence_text_link_max_lines = max(
            1,
            int(getattr(settings, "RAG_EVIDENCE_TEXT_LINK_MAX_LINES", 4)),
        )
        self.pdf_table_text_overlap_filter_enabled = bool(
            getattr(settings, "RAG_PDF_TABLE_TEXT_OVERLAP_FILTER_ENABLED", True)
        )
        self.pdf_table_promotion_gate_enabled = bool(
            getattr(settings, "RAG_PDF_TABLE_PROMOTION_GATE_ENABLED", True)
        )
        self.pdf_table_recurring_scaffold_min_repeats = max(
            2,
            int(getattr(settings, "RAG_PDF_TABLE_RECURRING_SCAFFOLD_MIN_REPEATS", 5)),
        )
        self.pdf_table_paragraph_long_cell_words = max(
            6,
            int(getattr(settings, "RAG_PDF_TABLE_PARAGRAPH_LONG_CELL_WORDS", 12)),
        )
        self.pdf_table_paragraph_long_cell_ratio = float(
            getattr(settings, "RAG_PDF_TABLE_PARAGRAPH_LONG_CELL_RATIO", 0.35)
        )
        if not (0.0 <= self.pdf_table_paragraph_long_cell_ratio <= 1.0):
            self.pdf_table_paragraph_long_cell_ratio = 0.35
        self.pdf_table_paragraph_min_rows = max(
            2,
            int(getattr(settings, "RAG_PDF_TABLE_PARAGRAPH_MIN_ROWS", 3)),
        )
        self.pdf_table_leading_blank_row_limit = max(
            1,
            int(getattr(settings, "RAG_PDF_TABLE_LEADING_BLANK_ROW_LIMIT", 2)),
        )
        self.pdf_page_chrome_suppression_enabled = bool(
            getattr(settings, "RAG_PDF_PAGE_CHROME_SUPPRESSION_ENABLED", True)
        )
        self.pdf_page_chrome_min_repeats = max(
            2,
            int(getattr(settings, "RAG_PDF_PAGE_CHROME_MIN_REPEATS", 3)),
        )
        self.pdf_page_chrome_top_ratio = float(
            getattr(settings, "RAG_PDF_PAGE_CHROME_TOP_RATIO", 0.16)
        )
        if not (0.0 <= self.pdf_page_chrome_top_ratio <= 1.0):
            self.pdf_page_chrome_top_ratio = 0.16
        self.pdf_page_chrome_bottom_ratio = float(
            getattr(settings, "RAG_PDF_PAGE_CHROME_BOTTOM_RATIO", 0.12)
        )
        if not (0.0 <= self.pdf_page_chrome_bottom_ratio <= 1.0):
            self.pdf_page_chrome_bottom_ratio = 0.12
        self.pdf_page_chrome_max_words = max(
            4,
            int(getattr(settings, "RAG_PDF_PAGE_CHROME_MAX_WORDS", 24)),
        )
        self.pdf_table_micro_fragment_min_columns = max(
            4,
            int(getattr(settings, "RAG_PDF_TABLE_MICRO_FRAGMENT_MIN_COLUMNS", 8)),
        )
        self.pdf_table_micro_fragment_short_cell_ratio = float(
            getattr(settings, "RAG_PDF_TABLE_MICRO_FRAGMENT_SHORT_CELL_RATIO", 0.45)
        )
        if not (0.0 <= self.pdf_table_micro_fragment_short_cell_ratio <= 1.0):
            self.pdf_table_micro_fragment_short_cell_ratio = 0.45
        self.pdf_table_bridge_max_rows = max(
            1,
            int(getattr(settings, "RAG_PDF_TABLE_BRIDGE_MAX_ROWS", 4)),
        )
        self.pdf_table_bridge_long_cell_ratio = float(
            getattr(settings, "RAG_PDF_TABLE_BRIDGE_LONG_CELL_RATIO", 0.5)
        )
        if not (0.0 <= self.pdf_table_bridge_long_cell_ratio <= 1.0):
            self.pdf_table_bridge_long_cell_ratio = 0.5
        self.pdf_table_text_overlap_min_ratio = float(
            getattr(settings, "RAG_PDF_TABLE_TEXT_OVERLAP_MIN_RATIO", 0.55)
        )
        if not (0.0 <= self.pdf_table_text_overlap_min_ratio <= 1.0):
            self.pdf_table_text_overlap_min_ratio = 0.55
        self.pdf_table_region_merge_x_margin_ratio = float(
            getattr(settings, "RAG_PDF_TABLE_REGION_MERGE_X_MARGIN_RATIO", 0.012)
        )
        if self.pdf_table_region_merge_x_margin_ratio < 0.0:
            self.pdf_table_region_merge_x_margin_ratio = 0.0
        self.pdf_table_region_merge_y_margin_ratio = float(
            getattr(settings, "RAG_PDF_TABLE_REGION_MERGE_Y_MARGIN_RATIO", 0.008)
        )
        if self.pdf_table_region_merge_y_margin_ratio < 0.0:
            self.pdf_table_region_merge_y_margin_ratio = 0.0
        self.pdf_table_residual_overlap_min_ratio = float(
            getattr(settings, "RAG_PDF_TABLE_RESIDUAL_OVERLAP_MIN_RATIO", 0.08)
        )
        if not (0.0 <= self.pdf_table_residual_overlap_min_ratio <= 1.0):
            self.pdf_table_residual_overlap_min_ratio = 0.08
        self.pdf_table_residual_near_region_ratio = float(
            getattr(settings, "RAG_PDF_TABLE_RESIDUAL_NEAR_REGION_RATIO", 0.012)
        )
        if self.pdf_table_residual_near_region_ratio < 0.0:
            self.pdf_table_residual_near_region_ratio = 0.0
        self.table_residual_max_per_region = max(
            1,
            int(getattr(settings, "RAG_TABLE_RESIDUAL_MAX_PER_REGION", 1)),
        )
        self.table_residual_equivalence_min_overlap = float(
            getattr(settings, "RAG_TABLE_RESIDUAL_EQUIV_MIN_OVERLAP", 0.65)
        )
        if not (0.0 <= self.table_residual_equivalence_min_overlap <= 1.0):
            self.table_residual_equivalence_min_overlap = 0.65
        self.table_residual_equivalence_min_shared_tokens = max(
            1,
            int(getattr(settings, "RAG_TABLE_RESIDUAL_EQUIV_MIN_SHARED_TOKENS", 5)),
        )
        self.table_residual_compact_max_chars = max(
            200,
            int(getattr(settings, "RAG_TABLE_RESIDUAL_COMPACT_MAX_CHARS", 700)),
        )
        self.table_annotation_enabled = bool(getattr(settings, "RAG_TABLE_ANNOTATION_ENABLED", True))
        self.table_annotation_max_chars = max(
            200,
            int(getattr(settings, "RAG_TABLE_ANNOTATION_MAX_CHARS", 1200)),
        )
        self.table_annotation_max_per_table = max(
            1,
            int(getattr(settings, "RAG_TABLE_ANNOTATION_MAX_PER_TABLE", 1)),
        )
        self.canonical_chunk_schema_version = max(
            1,
            int(getattr(settings, "RAG_CANONICAL_CHUNK_SCHEMA_VERSION", 1)),
        )
        self.chunk_quality_heading_max_lines = max(
            1,
            int(getattr(settings, "RAG_CHUNK_HEADING_MAX_LINES", 3)),
        )
        self.chunk_quality_heading_max_tokens = max(
            1,
            int(getattr(settings, "RAG_CHUNK_HEADING_MAX_TOKENS", 12)),
        )
        self.dataset_mode_enabled = bool(getattr(settings, "DATASET_MODE_ENABLED", True))
        self.dataset_row_threshold = max(
            1,
            int(
                getattr(
                    settings,
                    "DATASET_MODE_ROW_THRESHOLD",
                    getattr(settings, "RAG_TABLE_LARGE_ROW_LIMIT", 20000),
                )
            ),
        )
        self.dataset_preview_rows = max(5, int(getattr(settings, "DATASET_MODE_PREVIEW_ROWS", 200)))
        self.dataset_sample_rows = max(1, int(getattr(settings, "DATASET_MODE_SAMPLE_ROWS", 20)))
        self.dataset_storage_format = str(getattr(settings, "DATASET_STORAGE_FORMAT", "csv_gz") or "csv_gz").strip()
        if self.dataset_storage_format not in {"csv_gz"}:
            self.dataset_storage_format = "csv_gz"
        self.job_lease_seconds = max(60, int(getattr(settings, "INGEST_JOB_LEASE_SECONDS", 1800)))
        self.job_retry_base_seconds = max(1.0, float(getattr(settings, "INGEST_JOB_RETRY_BASE_SECONDS", 5.0)))
        self.job_retry_max_seconds = max(
            self.job_retry_base_seconds,
            float(getattr(settings, "INGEST_JOB_RETRY_MAX_SECONDS", 300.0)),
        )
        self.job_retry_jitter_seconds = max(0.0, float(getattr(settings, "INGEST_JOB_RETRY_JITTER_SECONDS", 2.0)))
        self.ingest_job_max_attempts = max(1, int(getattr(settings, "INGEST_JOB_MAX_ATTEMPTS", 3)))
        self.embed_job_max_attempts = max(1, int(getattr(settings, "INGEST_EMBED_JOB_MAX_ATTEMPTS", 5)))
        self.requeue_stale_jobs = bool(getattr(settings, "INGEST_JOB_REQUEUE_STALE_ENABLED", True))
        self.tenant_lexicon_auto_learning_enabled = bool(
            getattr(settings, "RAG_TENANT_LEXICON_AUTO_LEARN_ENABLED", True)
        )
        self._tenant_lexicon_auto_learning_service: TenantLexiconAutoLearningService | None = None

__all__ = [
    "queue_ingestion_job",
    "KnowledgeIngestionService",
    "KnowledgeIngestionError",
    "IngestionJobResult",
]
