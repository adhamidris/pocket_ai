from __future__ import annotations

import json
import shutil
import tempfile
import uuid
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from django.conf import settings
from django.test import SimpleTestCase, TestCase, override_settings

from apps.accounts.models import (
    BusinessProfile,
    KnowledgeBlockType,
    KnowledgeSourceType,
    KnowledgeStatus,
    RegistrationSession,
    User,
)
from apps.knowledge.models import (
    KnowledgeAlias,
    KnowledgeEntity,
    KnowledgeUpload,
    KnowledgeUploadChunk,
    KnowledgeUploadFile,
    KnowledgeUploadIssue,
    KnowledgeUploadTable,
    KnowledgeUploadTableRow,
    KnowledgeUploadText,
)
from apps.knowledge.knowledge_ingestion import (
    GeometryTableReconstructor,
    KnowledgeIngestionService,
    PageBlockPayload,
    PageLayout,
    PdfSpan,
    TableCellPayload,
    TablePayload,
    TableRowPayload,
)
from core.tenancy import tenant_context

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None


class KnowledgeIngestionAliasTests(TestCase):
    def test_collect_aliases_detects_short_identifiers(self):
        aliases, sources = KnowledgeIngestionService._collect_aliases_from_record(
            record={"slug": "trip-1"},
            flattened={"slug": "trip-1", "code": "A-9"},
            attributes={"slug": "trip-1"},
            entity_name="Trip 1",
        )
        self.assertIn("trip-1", aliases)
        self.assertTrue(any(source.startswith("record_slug") for source in sources))


class KnowledgeIngestionChunkingTests(SimpleTestCase):
    def test_chunk_text_overlap_does_not_split_words(self) -> None:
        # Regression: overlap should not start mid-token (e.g., "Free" -> "ee"), which creates
        # noisy fragments that pollute retrieval for tabular PDFs.
        text = "AAAAA Free\nBBBBB\nCCCCC"
        segments = KnowledgeIngestionService._chunk_text(text, chunk_chars=12, overlap=2)
        self.assertGreaterEqual(len(segments), 2)
        first_line = segments[1].splitlines()[0]
        self.assertNotEqual(first_line, "ee")
        self.assertTrue(segments[1].startswith("Free") or segments[1].startswith("BBBBB"))

    def test_numeric_signal_detection_is_domain_agnostic(self) -> None:
        self.assertTrue(
            KnowledgeIngestionService._has_numeric_table_signal("Latency P95 320 ms P99 870 ms")
        )
        self.assertTrue(
            KnowledgeIngestionService._has_numeric_table_signal("Seat capacity 120 units")
        )
        self.assertTrue(
            KnowledgeIngestionService._has_numeric_table_signal("Window 2026-02-08 14:30")
        )
        self.assertFalse(
            KnowledgeIngestionService._has_numeric_table_signal("General policy overview and notes")
        )

    def test_table_row_chunks_emit_applies_to_contract(self) -> None:
        class _Manager:
            def __init__(self, items):
                self._items = list(items)

            def all(self):
                return list(self._items)

        class _Cell:
            def __init__(self, cell_id: str, column_index: int, raw_text: str, column_key: str = ""):
                self.id = cell_id
                self.column_index = column_index
                self.column_key = column_key
                self.raw_text = raw_text

        class _Row:
            def __init__(self):
                self.row_index = 3
                self.metadata = {
                    "row_type": "data",
                    "table_scope_contract_version": "v2",
                    "observed_value_columns": ["Service", "Wealth"],
                    "qualifier_columns": ["Service"],
                    "scope_dimension_columns": ["Prime", "Plus", "Wealth", "Exclusive Wealth", "Private"],
                    "inferred_scope_columns": ["Prime", "Plus", "Wealth", "Exclusive Wealth", "Private"],
                    "scope_reason": "inferred_sparse_expansion",
                    "scope_confidence": 0.72,
                    "scope_value": "1% (Min USD 2)",
                }
                self.cells = _Manager(
                    [
                        _Cell("c-1", 0, "Traveler cheques (sell) in foreign currency"),
                        _Cell("c-2", 4, "1% (Min USD 2)", "wealth"),
                    ]
                )

        class _Table:
            def __init__(self):
                self.title = "CIB Teller Fees"
                self.order_index = 1
                self.section_heading = "Over the counter"
                self.rows = _Manager([_Row()])

        service = KnowledgeIngestionService(enable_ocr=False)
        payloads = service._table_row_chunk_payloads(
            table=_Table(),
            column_map=[
                ("Service", "service", 0),
                ("Tariff", "tariff", 1),
                ("Prime", "prime", 2),
                ("Plus", "plus", 3),
                ("Wealth", "wealth", 4),
                ("Exclusive Wealth", "exclusive_wealth", 5),
                ("Private", "private", 6),
            ],
            raw_schema=["service", "tariff", "prime", "plus", "wealth", "exclusive_wealth", "private"],
            privacy_rules={},
            base_metadata={"is_table_chunk": True, "table_id": "table-1"},
            max_rows=10,
        )

        self.assertEqual(len(payloads), 1)
        payload = payloads[0]
        text = str(payload.get("text") or "")
        metadata = payload.get("metadata") or {}

        self.assertIn("[Scope] Prime, Plus, Wealth, Exclusive Wealth, Private", text)
        self.assertEqual(
            metadata.get("table_row_inferred_scope_columns"),
            ["Prime", "Plus", "Wealth", "Exclusive Wealth", "Private"],
        )
        self.assertNotIn("table_row_applies_to_columns", metadata)
        self.assertNotIn("table_row_applicability_mode", metadata)
        self.assertEqual(metadata.get("table_row_fee_value"), "1% (Min USD 2)")

    def test_prefers_flat_text_fallback_when_page_blocks_are_fragmented(self) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        pages = [
            PageLayout(
                page_number=1,
                width=800.0,
                height=1000.0,
                rotation=0,
                text_density=0.3,
                has_ocr_content=False,
                content_type="application/pdf",
                blocks=[
                    PageBlockPayload(
                        block_type=KnowledgeBlockType.PARAGRAPH,
                        order_index=0,
                        text="Issuance Fees",
                        bbox={"x0": 0.0, "y0": 0.0, "x1": 120.0, "y1": 20.0},
                        metadata={},
                    ),
                    PageBlockPayload(
                        block_type=KnowledgeBlockType.PARAGRAPH,
                        order_index=1,
                        text="EGP 500",
                        bbox={"x0": 130.0, "y0": 0.0, "x1": 220.0, "y1": 20.0},
                        metadata={},
                    ),
                    PageBlockPayload(
                        block_type=KnowledgeBlockType.PARAGRAPH,
                        order_index=2,
                        text="Replacement Fees",
                        bbox={"x0": 0.0, "y0": 22.0, "x1": 140.0, "y1": 40.0},
                        metadata={},
                    ),
                    PageBlockPayload(
                        block_type=KnowledgeBlockType.PARAGRAPH,
                        order_index=3,
                        text="Free",
                        bbox={"x0": 150.0, "y0": 22.0, "x1": 220.0, "y1": 40.0},
                        metadata={},
                    ),
                    PageBlockPayload(
                        block_type=KnowledgeBlockType.PARAGRAPH,
                        order_index=4,
                        text="Grace Period",
                        bbox={"x0": 0.0, "y0": 44.0, "x1": 120.0, "y1": 64.0},
                        metadata={},
                    ),
                    PageBlockPayload(
                        block_type=KnowledgeBlockType.PARAGRAPH,
                        order_index=5,
                        text="55 days",
                        bbox={"x0": 130.0, "y0": 44.0, "x1": 210.0, "y1": 64.0},
                        metadata={},
                    ),
                    PageBlockPayload(
                        block_type=KnowledgeBlockType.PARAGRAPH,
                        order_index=6,
                        text="Cash Withdrawal",
                        bbox={"x0": 0.0, "y0": 66.0, "x1": 140.0, "y1": 84.0},
                        metadata={},
                    ),
                    PageBlockPayload(
                        block_type=KnowledgeBlockType.PARAGRAPH,
                        order_index=7,
                        text="4% min. EGP 40",
                        bbox={"x0": 150.0, "y0": 66.0, "x1": 260.0, "y1": 84.0},
                        metadata={},
                    ),
                ],
                metadata={},
            )
        ]
        page_segments = service._build_text_segments_from_blocks(pages, chunk_chars=220, overlap=0)
        flat_segments = service._build_flat_text_segment_payloads(
            "Issuance Fees EGP 500 Replacement Fees Free Grace Period 55 days Cash Withdrawal 4% min. EGP 40",
            alias_hygiene=False,
        )

        self.assertTrue(page_segments)
        self.assertTrue(flat_segments)
        self.assertTrue(
            service._should_prefer_flat_text_segments(
                pages=pages,
                page_segments=page_segments,
                flat_segments=flat_segments,
            )
        )
        decision = service._text_chunk_source_decision(
            pages=pages,
            page_segments=page_segments,
            flat_segments=flat_segments,
        )
        self.assertEqual(decision.get("selected_source"), "flat_text")

    def test_table_row_chunks_derive_scope_with_structural_context_columns(self) -> None:
        class _Manager:
            def __init__(self, items):
                self._items = list(items)

            def all(self):
                return list(self._items)

        class _Cell:
            def __init__(self, cell_id: str, column_index: int, raw_text: str, column_key: str = ""):
                self.id = cell_id
                self.column_index = column_index
                self.raw_text = raw_text
                self.column_key = column_key

        class _Row:
            def __init__(self, row_index: int, metadata: dict[str, object], cells: list[_Cell]):
                self.row_index = row_index
                self.metadata = metadata
                self.cells = _Manager(cells)

        class _Table:
            def __init__(self, rows: list[_Row]):
                self.title = "Generic matrix"
                self.order_index = 1
                self.section_heading = ""
                self.rows = _Manager(rows)

        header = _Row(
            0,
            {"row_type": "header"},
            [
                _Cell("h-1", 0, "Descriptor A"),
                _Cell("h-2", 1, "Descriptor B"),
                _Cell("h-3", 2, "Band 1"),
                _Cell("h-4", 3, "Band 2"),
                _Cell("h-5", 4, "Band 3"),
            ],
        )
        row_one = _Row(
            1,
            {"row_type": "data"},
            [
                _Cell("r1-1", 0, "Route settlement fees"),
                _Cell("r1-2", 1, "Category for local currency payouts"),
                _Cell("r1-3", 2, "1.5%"),
                _Cell("r1-4", 3, "1.5%"),
                _Cell("r1-5", 4, "1.5%"),
            ],
        )
        row_two = _Row(
            2,
            {
                "row_type": "data",
                "table_scope_contract_version": "v2",
                "observed_value_columns": ["Descriptor A", "Descriptor B", "Band 2"],
                "qualifier_columns": ["Descriptor A", "Descriptor B"],
                "scope_dimension_columns": ["Band 1", "Band 2", "Band 3"],
                "inferred_scope_columns": ["Band 2"],
                "scope_reason": "single_populated_scope_column",
                "scope_confidence": 0.82,
                "scope_value": "2.0%",
            },
            [
                _Cell("r2-1", 0, "Remote processing surcharge"),
                _Cell("r2-2", 1, "Category for foreign currency payouts"),
                _Cell("r2-4", 3, "2.0%"),
            ],
        )

        service = KnowledgeIngestionService(enable_ocr=False)
        payloads = service._table_row_chunk_payloads(
            table=_Table([header, row_one, row_two]),
            column_map=[
                ("Descriptor A", "descriptor_a", 0),
                ("Descriptor B", "descriptor_b", 1),
                ("Band 1", "band_1", 2),
                ("Band 2", "band_2", 3),
                ("Band 3", "band_3", 4),
            ],
            raw_schema=["descriptor_a", "descriptor_b", "band_1", "band_2", "band_3"],
            privacy_rules={},
            base_metadata={"is_table_chunk": True, "table_id": "table-2"},
            max_rows=10,
        )

        self.assertEqual(len(payloads), 2)
        row_two_payload = next(
            payload for payload in payloads if (payload.get("metadata") or {}).get("table_row_index") == 2
        )
        inferred_scope = (row_two_payload.get("metadata") or {}).get("table_row_inferred_scope_columns") or []
        self.assertIn("Band 2", inferred_scope)
        self.assertNotIn("Descriptor A", inferred_scope)
        self.assertNotIn("Descriptor B", inferred_scope)

    def test_table_row_chunks_mark_structural_context_rows(self) -> None:
        class _Manager:
            def __init__(self, items):
                self._items = list(items)

            def all(self):
                return list(self._items)

        class _Cell:
            def __init__(self, cell_id: str, column_index: int, raw_text: str, column_key: str = ""):
                self.id = cell_id
                self.column_index = column_index
                self.raw_text = raw_text
                self.column_key = column_key

        class _Row:
            def __init__(self, row_index: int, metadata: dict[str, object], cells: list[_Cell]):
                self.row_index = row_index
                self.metadata = metadata
                self.cells = _Manager(cells)

        class _Table:
            def __init__(self, rows: list[_Row]):
                self.title = "Loan fees"
                self.order_index = 1
                self.section_heading = ""
                self.rows = _Manager(rows)

        header = _Row(
            0,
            {"row_type": "header"},
            [
                _Cell("h-1", 0, "service"),
                _Cell("h-2", 1, "fees_charges"),
                _Cell("h-3", 2, "Prime"),
                _Cell("h-4", 3, "Plus"),
                _Cell("h-5", 4, "Wealth"),
                _Cell("h-6", 5, "Private"),
            ],
        )
        structural = _Row(
            1,
            {
                "row_type": "data",
                "scope_dimension_columns": ["Prime", "Plus", "Wealth", "Private"],
                "observed_value_columns": [
                    "service",
                    "fees_charges",
                    "Prime",
                    "Plus",
                    "Wealth",
                    "Private",
                ],
            },
            [
                _Cell("s-1", 0, "Administration Fees on the total Loan Amount up to 8 years", "service"),
                _Cell("s-2", 1, "Segment/Product", "fees_charges"),
                _Cell("s-3", 2, "Prime", "Prime"),
                _Cell("s-4", 3, "Plus", "Plus"),
                _Cell("s-5", 4, "Wealth", "Wealth"),
                _Cell("s-6", 5, "Private", "Private"),
            ],
        )

        service = KnowledgeIngestionService(enable_ocr=False)
        payloads = service._table_row_chunk_payloads(
            table=_Table([header, structural]),
            column_map=[
                ("service", "service", 0),
                ("fees_charges", "fees_charges", 1),
                ("Prime", "prime", 2),
                ("Plus", "plus", 3),
                ("Wealth", "wealth", 4),
                ("Private", "private", 5),
            ],
            raw_schema=["service", "fees_charges", "prime", "plus", "wealth", "private"],
            privacy_rules={},
            base_metadata={"is_table_chunk": True, "table_id": "table-3"},
            max_rows=10,
        )

        self.assertEqual(len(payloads), 1)
        metadata = payloads[0].get("metadata") or {}
        self.assertTrue(metadata.get("table_row_is_structural_context"))
        self.assertEqual(metadata.get("table_row_structural_scope_echo_count"), 4)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_table_overlap_blocks_become_residual_segments_for_coverage(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        page = PageLayout(
            page_number=1,
            width=600.0,
            height=800.0,
            rotation=0,
            text_density=0.1,
            has_ocr_content=False,
            content_type="application/pdf",
            blocks=[
                PageBlockPayload(
                    block_type=KnowledgeBlockType.PARAGRAPH,
                    order_index=0,
                    text="Annual fee EGP 40",
                    bbox={"x0": 100.0, "y0": 100.0, "x1": 300.0, "y1": 140.0},
                    metadata={},
                ),
                PageBlockPayload(
                    block_type=KnowledgeBlockType.PARAGRAPH,
                    order_index=1,
                    text="General terms apply outside the fee table.",
                    bbox={"x0": 20.0, "y0": 720.0, "x1": 280.0, "y1": 760.0},
                    metadata={},
                ),
            ],
            metadata={},
        )
        table = TablePayload(
            order_index=1,
            title="Fees",
            section_heading="",
            page_number=1,
            bbox={"x0": 90.0, "y0": 90.0, "x1": 560.0, "y1": 610.0},
        )

        annotated_pages, diagnostics = service._annotate_pdf_blocks_with_table_overlap([page], [table])

        self.assertEqual(diagnostics.get("suppressed_text_blocks"), 1)
        self.assertEqual(diagnostics.get("overlapping_text_blocks"), 1)
        first_meta = annotated_pages[0].blocks[0].metadata
        second_meta = annotated_pages[0].blocks[1].metadata
        self.assertTrue(first_meta.get("table_overlap_candidate"))
        self.assertEqual(first_meta.get("table_overlap_candidate_reason"), "table_overlap")
        self.assertFalse(first_meta.get("suppress_text_chunk", False))
        self.assertFalse(second_meta.get("suppress_text_chunk", False))

        segments = service._build_text_segments_from_blocks(annotated_pages)
        combined = "\n".join(str(segment.get("text") or "") for segment in segments)
        self.assertIn("Annual fee EGP 40", combined)
        self.assertIn("General terms apply", combined)
        residual_segments = [seg for seg in segments if (seg.get("metadata") or {}).get("table_residual")]
        self.assertFalse(residual_segments)
        narrative_segments = [
            seg
            for seg in segments
            if (seg.get("metadata") or {}).get("content_source") == "page_blocks"
        ]
        annual_fee_segments = [seg for seg in narrative_segments if "Annual fee EGP 40" in str(seg.get("text") or "")]
        self.assertTrue(annual_fee_segments)
        self.assertTrue((annual_fee_segments[0].get("metadata") or {}).get("table_adjacent"))

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_non_overlapping_blocks_remain_chunked(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        page = PageLayout(
            page_number=1,
            width=600.0,
            height=800.0,
            rotation=0,
            text_density=0.1,
            has_ocr_content=False,
            content_type="application/pdf",
            blocks=[
                PageBlockPayload(
                    block_type=KnowledgeBlockType.PARAGRAPH,
                    order_index=0,
                    text="Frequently asked questions for branch services.",
                    bbox={"x0": 25.0, "y0": 680.0, "x1": 420.0, "y1": 740.0},
                    metadata={},
                ),
            ],
            metadata={},
        )
        table = TablePayload(
            order_index=1,
            title="Fees",
            section_heading="",
            page_number=1,
            bbox={"x0": 90.0, "y0": 90.0, "x1": 560.0, "y1": 610.0},
        )

        annotated_pages, diagnostics = service._annotate_pdf_blocks_with_table_overlap([page], [table])
        self.assertEqual(diagnostics.get("suppressed_text_blocks"), 0)
        self.assertFalse(annotated_pages[0].blocks[0].metadata.get("suppress_text_chunk", False))

        segments = service._build_text_segments_from_blocks(annotated_pages)
        self.assertTrue(segments)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_restores_consumed_blocks_when_owning_table_is_suppressed(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        page = PageLayout(
            page_number=1,
            width=600.0,
            height=800.0,
            rotation=0,
            text_density=0.1,
            has_ocr_content=False,
            content_type="application/pdf",
            blocks=[
                PageBlockPayload(
                    block_type=KnowledgeBlockType.PARAGRAPH,
                    order_index=1,
                    text="International Delivery Shipment Fees\nUSD 30",
                    bbox={"x0": 80.0, "y0": 120.0, "x1": 340.0, "y1": 180.0},
                    metadata={
                        "anchor": "p1-b1",
                        "region_role": "text",
                        "canonical_consumed_by_table": True,
                        "canonical_consumed_reason": "cell_completion_attachment",
                        "canonical_consumed_table_order_index": 2,
                        "canonical_consumed_table_page_number": 1,
                        "canonical_consumed_row_index": 1,
                        "canonical_consumed_column_index": 1,
                    },
                )
            ],
            metadata={},
        )
        surviving_table = TablePayload(
            order_index=1,
            title="Customer Service Fees & Charges",
            section_heading="Customer Service Fees & Charges",
            page_number=1,
            bbox={"x0": 60.0, "y0": 200.0, "x1": 560.0, "y1": 720.0},
            column_schema=["service", "tariff"],
            data_dictionary={},
            metadata={"promotion_decision": "keep"},
            rows=[],
        )

        restored_pages, stats = service._restore_consumed_blocks_for_suppressed_tables(
            [page],
            [surviving_table],
        )

        self.assertEqual(stats.get("restored_blocks"), 1)
        restored_meta = restored_pages[0].blocks[0].metadata or {}
        self.assertFalse(restored_meta.get("canonical_consumed_by_table"))
        self.assertTrue(restored_meta.get("canonical_restored_after_table_suppression"))

        segments = service._build_text_segments_from_blocks(restored_pages, chunk_chars=400, overlap=0)
        rendered = "\n".join(str(segment.get("text") or "") for segment in segments)
        self.assertIn("International Delivery Shipment Fees", rendered)
        self.assertIn("USD 30", rendered)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_keeps_consumed_blocks_hidden_when_owning_table_survives(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        page = PageLayout(
            page_number=1,
            width=600.0,
            height=800.0,
            rotation=0,
            text_density=0.1,
            has_ocr_content=False,
            content_type="application/pdf",
            blocks=[
                PageBlockPayload(
                    block_type=KnowledgeBlockType.PARAGRAPH,
                    order_index=1,
                    text="International Delivery Shipment Fees\nUSD 30",
                    bbox={"x0": 80.0, "y0": 120.0, "x1": 340.0, "y1": 180.0},
                    metadata={
                        "anchor": "p1-b1",
                        "region_role": "text",
                        "canonical_consumed_by_table": True,
                        "canonical_consumed_reason": "cell_completion_attachment",
                        "canonical_consumed_table_order_index": 2,
                        "canonical_consumed_table_page_number": 1,
                    },
                )
            ],
            metadata={},
        )
        surviving_table = TablePayload(
            order_index=2,
            title="International Delivery Tariffs",
            section_heading="International Delivery Tariffs",
            page_number=1,
            bbox={"x0": 60.0, "y0": 80.0, "x1": 420.0, "y1": 220.0},
            column_schema=["service", "value"],
            data_dictionary={},
            metadata={"promotion_decision": "keep"},
            rows=[],
        )

        restored_pages, stats = service._restore_consumed_blocks_for_suppressed_tables(
            [page],
            [surviving_table],
        )

        self.assertEqual(stats.get("restored_blocks"), 0)
        self.assertTrue(restored_pages[0].blocks[0].metadata.get("canonical_consumed_by_table"))

        segments = service._build_text_segments_from_blocks(restored_pages, chunk_chars=400, overlap=0)
        self.assertEqual(segments, [])

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_table_regions_merge_fragmented_boxes_without_merging_far_tables(
        self,
        _build_embeddings,
    ) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        page = PageLayout(
            page_number=1,
            width=600.0,
            height=800.0,
            rotation=0,
            text_density=0.1,
            has_ocr_content=False,
            content_type="application/pdf",
            blocks=[],
            metadata={},
        )
        tables = [
            TablePayload(
                order_index=1,
                title="Fees (fragment 1)",
                section_heading="",
                page_number=1,
                bbox={"x0": 48.0, "y0": 100.0, "x1": 560.0, "y1": 220.0},
            ),
            TablePayload(
                order_index=2,
                title="Fees (fragment 2)",
                section_heading="",
                page_number=1,
                bbox={"x0": 52.0, "y0": 224.0, "x1": 558.0, "y1": 360.0},
            ),
            TablePayload(
                order_index=3,
                title="Separate mini table",
                section_heading="",
                page_number=1,
                bbox={"x0": 60.0, "y0": 620.0, "x1": 260.0, "y1": 700.0},
            ),
        ]

        regions = service._table_regions_by_page(tables, pages=[page]).get(1, [])
        self.assertEqual(len(regions), 2)
        merged_top = regions[0]
        self.assertLessEqual(float(merged_top["x0"]), 48.0)
        self.assertGreaterEqual(float(merged_top["x1"]), 558.0)
        self.assertLessEqual(float(merged_top["y0"]), 100.0)
        self.assertGreaterEqual(float(merged_top["y1"]), 360.0)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_table_residual_blocks_are_tagged_as_fallback_segments(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        page = PageLayout(
            page_number=1,
            width=600.0,
            height=800.0,
            rotation=0,
            text_density=0.1,
            has_ocr_content=False,
            content_type="application/pdf",
            blocks=[
                PageBlockPayload(
                    block_type=KnowledgeBlockType.PARAGRAPH,
                    order_index=0,
                    text="Cash deposit with same day value date (T+3 customers) 0,3% minimum EGP 100",
                    # 2px below the table region to exercise near-region residual tagging.
                    bbox={"x0": 90.0, "y0": 612.0, "x1": 560.0, "y1": 640.0},
                    metadata={},
                ),
                PageBlockPayload(
                    block_type=KnowledgeBlockType.PARAGRAPH,
                    order_index=1,
                    text="General branch notes and disclaimers for customers.",
                    bbox={"x0": 40.0, "y0": 700.0, "x1": 460.0, "y1": 740.0},
                    metadata={},
                ),
            ],
            metadata={},
        )
        table = TablePayload(
            order_index=1,
            title="Fees",
            section_heading="",
            page_number=1,
            bbox={"x0": 80.0, "y0": 80.0, "x1": 570.0, "y1": 610.0},
        )

        annotated_pages, diagnostics = service._annotate_pdf_blocks_with_table_overlap([page], [table])
        self.assertEqual(diagnostics.get("table_residual_blocks"), 1)

        residual_meta = annotated_pages[0].blocks[0].metadata
        generic_meta = annotated_pages[0].blocks[1].metadata
        self.assertTrue(residual_meta.get("table_residual_candidate"))
        self.assertEqual(residual_meta.get("search_tier"), "fallback")
        self.assertFalse(residual_meta.get("suppress_text_chunk", False))
        self.assertFalse(generic_meta.get("table_residual_candidate", False))

        segments = service._build_text_segments_from_blocks(annotated_pages)
        residual_segments = [seg for seg in segments if (seg.get("metadata") or {}).get("table_residual")]
        self.assertTrue(residual_segments)
        self.assertEqual((residual_segments[0].get("metadata") or {}).get("search_tier"), "fallback")
        self.assertEqual((residual_segments[0].get("metadata") or {}).get("content_source"), "table_residual")

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_overlap_candidates_remain_narrative_page_segments_per_block(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        page = PageLayout(
            page_number=1,
            width=600.0,
            height=800.0,
            rotation=0,
            text_density=0.1,
            has_ocr_content=False,
            content_type="application/pdf",
            blocks=[
                PageBlockPayload(
                    block_type=KnowledgeBlockType.PARAGRAPH,
                    order_index=0,
                    text="Cash deposit with same day value date 0,2% minimum EGP 100.",
                    bbox={"x0": 120.0, "y0": 120.0, "x1": 500.0, "y1": 150.0},
                    metadata={},
                ),
                PageBlockPayload(
                    block_type=KnowledgeBlockType.PARAGRAPH,
                    order_index=1,
                    text="In case of weekend submissions, one extra working day applies.",
                    bbox={"x0": 120.0, "y0": 170.0, "x1": 560.0, "y1": 210.0},
                    metadata={},
                ),
            ],
            metadata={},
        )
        table = TablePayload(
            order_index=1,
            title="Fees",
            section_heading="",
            page_number=1,
            bbox={"x0": 90.0, "y0": 90.0, "x1": 570.0, "y1": 610.0},
        )

        annotated_pages, _ = service._annotate_pdf_blocks_with_table_overlap([page], [table])
        segments = service._build_text_segments_from_blocks(annotated_pages)
        residual_segments = [seg for seg in segments if (seg.get("metadata") or {}).get("table_residual")]
        self.assertFalse(residual_segments)
        narrative_segments = [
            seg
            for seg in segments
            if (seg.get("metadata") or {}).get("content_source") == "page_blocks"
        ]
        self.assertEqual(len(narrative_segments), 1)
        self.assertTrue((narrative_segments[0].get("metadata") or {}).get("table_adjacent"))
        rendered = "\n".join(seg.get("text") or "" for seg in narrative_segments)
        self.assertIn("same day value date", rendered)
        self.assertIn("extra working day applies", rendered)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_note_like_residual_is_not_dropped_by_row_subset_overlap(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        segment_payloads = [
            {
                "text": "[Table] Fees\n[Row] 18\nService: Cash deposit with same day value date\nTariff: 0,2% (With minimum EGP 100 or Equivalent and with no maximum)",
                "metadata": {
                    "is_table_chunk": True,
                    "index_type": "table",
                    "content_source": "table_row",
                    "table_page_number": 1,
                    "table_row_index": 18,
                },
            },
            {
                "text": "In case cash deposits are made after 2:00 pm, Saturdays or public holidays, an extra working day will be counted to the applied value date according to account currency.",
                "metadata": {
                    "index_type": "text",
                    "content_source": "table_residual",
                    "region_role": "table_residual",
                    "table_residual": True,
                    "table_residual_region_key": "p1-r0",
                    "page_numbers": [1],
                    "table_overlap_ratio_max": 0.88,
                },
            },
        ]

        reconciled, stats = service._reconcile_table_residual_segments(segment_payloads)
        residual_reconciled = [
            payload for payload in reconciled if (payload.get("metadata") or {}).get("table_residual")
        ]

        self.assertEqual(stats.get("input_residual_segments"), 1)
        self.assertEqual(stats.get("dropped_equivalent_segments"), 0)
        self.assertEqual(stats.get("kept_residual_segments"), 1)
        self.assertEqual(len(residual_reconciled), 1)
        residual_meta = residual_reconciled[0].get("metadata") or {}
        self.assertEqual(residual_meta.get("table_residual_segment_kind"), "note_like")
        self.assertTrue(residual_meta.get("table_residual_equivalence_checked"))

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_cell_like_residual_segments_are_dropped(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        segment_payloads = [
            {
                "text": "[Table] Fees\n[Row] 1\nService: Cash Withdrawal/Deposit over the counter\nPrime: EGP 40",
                "metadata": {
                    "is_table_chunk": True,
                    "index_type": "table",
                    "content_source": "table_row",
                    "table_page_number": 1,
                    "table_row_index": 1,
                },
            },
            {
                "text": "Free",
                "metadata": {
                    "index_type": "text",
                    "content_source": "table_residual",
                    "region_role": "table_residual",
                    "table_residual": True,
                    "table_residual_region_key": "p1-r0",
                    "page_numbers": [1],
                    "table_overlap_ratio_max": 0.9,
                },
            },
        ]

        reconciled, stats = service._reconcile_table_residual_segments(segment_payloads)
        residual_reconciled = [
            payload for payload in reconciled if (payload.get("metadata") or {}).get("table_residual")
        ]

        self.assertEqual(stats.get("input_residual_segments"), 1)
        self.assertEqual(stats.get("dropped_equivalent_segments"), 1)
        self.assertEqual(stats.get("dropped_equivalent_cell_like_segments"), 1)
        self.assertEqual(stats.get("kept_residual_segments"), 0)
        self.assertFalse(residual_reconciled)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_table_residual_reconciliation_dedupes_equivalent_rows_and_preserves_unique_segments(
        self,
        _build_embeddings,
    ) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        segment_payloads = [
            {
                "text": "[Table] Fees\n[Row] 1\nService: Cash deposit with same day value date\nTariff: 0,3% minimum EGP 100",
                "metadata": {
                    "is_table_chunk": True,
                    "index_type": "table",
                    "content_source": "table_row",
                    "table_page_number": 1,
                    "table_row_index": 1,
                },
            },
            {
                "text": "Cash deposit with same day value date Tariff 0,3% minimum EGP 100",
                "metadata": {
                    "index_type": "text",
                    "content_source": "table_residual",
                    "region_role": "table_residual",
                    "table_residual": True,
                    "table_residual_region_key": "p1-r0",
                    "page_numbers": [1],
                    "table_overlap_ratio_max": 0.18,
                },
            },
            {
                "text": "Cash deposit T+3 customers 0,3% minimum EGP 100 equivalent no max",
                "metadata": {
                    "index_type": "text",
                    "content_source": "table_residual",
                    "region_role": "table_residual",
                    "table_residual": True,
                    "table_residual_region_key": "p1-r0",
                    "page_numbers": [1],
                    "table_overlap_ratio_max": 0.11,
                },
            },
            {
                "text": "Cash deposit T+5 customers 0,5% minimum EGP 100 equivalent no max",
                "metadata": {
                    "index_type": "text",
                    "content_source": "table_residual",
                    "region_role": "table_residual",
                    "table_residual": True,
                    "table_residual_region_key": "p1-r0",
                    "page_numbers": [1],
                    "table_overlap_ratio_max": 0.27,
                },
            },
            {
                "text": "Cost per Bag (EGP 70,000) Fee EGP 14",
                "metadata": {
                    "index_type": "text",
                    "content_source": "table_residual",
                    "region_role": "table_residual",
                    "table_residual": True,
                    "table_residual_region_key": "p1-r1",
                    "page_numbers": [1],
                    "table_overlap_ratio_max": 0.22,
                },
            },
        ]

        reconciled, stats = service._reconcile_table_residual_segments(segment_payloads)
        residual_reconciled = [
            payload for payload in reconciled if (payload.get("metadata") or {}).get("table_residual")
        ]

        self.assertEqual(stats.get("input_residual_segments"), 4)
        self.assertEqual(stats.get("dropped_equivalent_segments"), 1)
        self.assertEqual(stats.get("dropped_cap_segments"), 0)
        self.assertEqual(stats.get("kept_residual_segments"), 3)
        self.assertEqual(len(residual_reconciled), 3)
        region_keys = {(payload.get("metadata") or {}).get("table_residual_region_key") for payload in residual_reconciled}
        self.assertEqual(region_keys, {"p1-r0", "p1-r1"})

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_canonicalize_payloads_projects_table_residuals_to_annotations(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        upload = SimpleNamespace(id=uuid.uuid4())
        segment_payloads = [
            {
                "text": "[Table] Fees\n[Columns] Service | Fee\n[Rows] 2 data rows",
                "metadata": {
                    "index_type": "table",
                    "is_table_chunk": True,
                    "content_source": "table_summary",
                    "table_chunk_role": "summary",
                    "table_id": "table-1",
                    "table_title": "Fees",
                    "table_order_index": 1,
                    "table_page_number": 1,
                    "search_tier": "primary",
                },
            },
            {
                "text": "[Table] Fees\n[Row] 1\nService: Cash deposit\nFee: EGP 14",
                "metadata": {
                    "index_type": "table",
                    "is_table_chunk": True,
                    "content_source": "table_row",
                    "table_chunk_role": "row",
                    "table_id": "table-1",
                    "table_row_index": 1,
                    "table_order_index": 1,
                    "table_page_number": 1,
                    "search_tier": "drill_down",
                },
            },
            {
                "text": "In case of cash deposits made after 2:00 pm, one extra working day applies.",
                "metadata": {
                    "index_type": "text",
                    "content_source": "table_residual",
                    "region_role": "table_residual",
                    "table_residual": True,
                    "table_residual_region_key": "p1-r0",
                    "table_overlap_ratio_max": 0.82,
                    "page_numbers": [1],
                },
            },
        ]

        canonicalized, stats = service._canonicalize_segment_payloads(
            upload=upload,
            segment_payloads=segment_payloads,
        )
        content_sources = [(payload.get("metadata") or {}).get("content_source") for payload in canonicalized]
        self.assertNotIn("table_residual", content_sources)
        self.assertIn("table_annotation", content_sources)

        annotation_payload = next(
            payload
            for payload in canonicalized
            if (payload.get("metadata") or {}).get("content_source") == "table_annotation"
        )
        annotation_meta = annotation_payload.get("metadata") or {}
        self.assertEqual(annotation_meta.get("canonical_chunk_kind"), "table_annotation")
        self.assertEqual(annotation_meta.get("table_id"), "table-1")
        self.assertEqual(annotation_meta.get("canonical_parent_anchor_id"), "table:table-1:summary")
        self.assertTrue(str(annotation_meta.get("canonical_anchor_id") or "").startswith("table:table-1:annotation:"))

        self.assertEqual((stats.get("kind_counts") or {}).get("table_annotation"), 1)
        projection_stats = stats.get("table_residual_projection") or {}
        self.assertEqual(projection_stats.get("input_residual_segments"), 1)
        self.assertEqual(projection_stats.get("table_annotation_chunks_created"), 1)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_canonicalize_payloads_assigns_narrative_anchor(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        upload = SimpleNamespace(id=uuid.uuid4())
        segment_payloads = [
            {
                "text": "General branch policy for customer service windows.",
                "metadata": {
                    "index_type": "text",
                    "content_source": "page_blocks",
                    "region_role": "text",
                    "page_numbers": [2],
                    "block_anchors": ["p2-b4"],
                },
            }
        ]

        canonicalized, _ = service._canonicalize_segment_payloads(
            upload=upload,
            segment_payloads=segment_payloads,
        )
        self.assertEqual(len(canonicalized), 1)
        metadata = canonicalized[0].get("metadata") or {}
        self.assertEqual(metadata.get("canonical_chunk_kind"), "narrative_paragraph")
        self.assertEqual(metadata.get("coverage_reason"), "layout_paragraph")
        self.assertTrue(str(metadata.get("canonical_anchor_id") or "").startswith("page:2:paragraph:"))

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_heading_like_table_residual_is_promoted_to_narrative(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        upload = SimpleNamespace(id=uuid.uuid4())
        segment_payloads = [
            {
                "text": "[Table] CIB-Teller EN – Table 1\n[Columns] Service | Tariff | Prime",
                "metadata": {
                    "index_type": "table",
                    "is_table_chunk": True,
                    "content_source": "table_summary",
                    "table_chunk_role": "summary",
                    "table_id": "table-main",
                    "table_title": "CIB-Teller EN – Table 1",
                    "table_order_index": 1,
                    "table_page_number": 1,
                },
            },
            {
                "text": "Over the counter in branch Fees & Charges",
                "metadata": {
                    "index_type": "text",
                    "content_source": "table_residual",
                    "region_role": "table_residual",
                    "table_residual": True,
                    "table_residual_region_key": "p1-r0",
                    "table_overlap_ratio_max": 0.91,
                    "page_numbers": [1],
                },
            },
        ]

        canonicalized, stats = service._canonicalize_segment_payloads(
            upload=upload,
            segment_payloads=segment_payloads,
        )
        narrative_payload = next(
            payload
            for payload in canonicalized
            if (payload.get("metadata") or {}).get("content_source") == "page_blocks"
        )
        narrative_meta = narrative_payload.get("metadata") or {}
        self.assertEqual(narrative_meta.get("canonical_chunk_kind"), "narrative_paragraph")
        self.assertEqual(narrative_meta.get("table_residual_projected"), "narrative")
        self.assertTrue(narrative_meta.get("table_adjacent"))
        self.assertEqual(
            (stats.get("table_residual_projection") or {}).get("narrative_promoted_segments"),
            1,
        )

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    @override_settings(
        RAG_TABLE_ANNOTATION_MAX_CHARS=80,
        RAG_TABLE_ANNOTATION_MAX_PER_TABLE=1,
    )
    def test_annotation_projection_preserves_overflow_chunks_beyond_soft_limit(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        upload = SimpleNamespace(id=uuid.uuid4())
        segment_payloads = [
            {
                "text": "[Table] Fees\n[Columns] Service | Tariff",
                "metadata": {
                    "index_type": "table",
                    "is_table_chunk": True,
                    "content_source": "table_summary",
                    "table_chunk_role": "summary",
                    "table_id": "table-overflow",
                    "table_title": "Fees",
                    "table_order_index": 1,
                    "table_page_number": 1,
                },
            },
            {
                "text": "*In case cash deposits are made after 2:00 pm, one extra working day applies for value date.",
                "metadata": {
                    "index_type": "text",
                    "content_source": "table_residual",
                    "region_role": "table_residual",
                    "table_residual": True,
                    "table_residual_region_key": "p1-r0",
                    "table_overlap_ratio_max": 0.88,
                    "page_numbers": [1],
                },
            },
            {
                "text": "*For public holidays, settlement timing shifts according to account currency and cut-off windows.",
                "metadata": {
                    "index_type": "text",
                    "content_source": "table_residual",
                    "region_role": "table_residual",
                    "table_residual": True,
                    "table_residual_region_key": "p1-r0",
                    "table_overlap_ratio_max": 0.86,
                    "page_numbers": [1],
                },
            },
        ]

        canonicalized, stats = service._canonicalize_segment_payloads(
            upload=upload,
            segment_payloads=segment_payloads,
        )
        annotation_chunks = [
            payload
            for payload in canonicalized
            if (payload.get("metadata") or {}).get("content_source") == "table_annotation"
        ]
        self.assertGreaterEqual(len(annotation_chunks), 2)
        self.assertTrue(
            all((payload.get("metadata") or {}).get("table_annotation_soft_limit_exceeded") for payload in annotation_chunks)
        )
        projection_stats = stats.get("table_residual_projection") or {}
        self.assertEqual(projection_stats.get("soft_limit_exceeded_tables"), 1)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_auto_selection_overrides_fragmented_heuristic_when_non_heuristic_is_viable(
        self,
        _build_embeddings,
    ) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)

        heuristic_tables = [
            TablePayload(
                order_index=i + 1,
                title="EGP",
                section_heading="",
                page_number=1,
                bbox={"x0": 70.0 + (i * 5.0), "y0": 620.0 + (i * 40.0), "x1": 380.0, "y1": 650.0 + (i * 40.0)},
                rows=[
                    TableRowPayload(
                        row_index=1,
                        page_number=1,
                        raw_text=f"row-{i + 1}",
                        metadata={"row_type": "data"},
                    )
                ],
                metadata={"detected_via": "heuristic"},
            )
            for i in range(4)
        ]
        geometry_table = TablePayload(
            order_index=1,
            title="Over the counter in branch Fees & Charges",
            section_heading="",
            page_number=1,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 1196.0, "y1": 989.0},
            rows=[
                TableRowPayload(
                    row_index=i + 1,
                    page_number=1,
                    raw_text=f"data-row-{i + 1}",
                    metadata={"row_type": "data"},
                )
                for i in range(19)
            ],
            metadata={"detected_via": "geometry"},
        )
        candidates = {
            "heuristic": heuristic_tables,
            "geometry": [geometry_table],
        }

        selected, tables, diag = service._select_table_candidates(candidates)

        self.assertEqual(selected, "heuristic")
        self.assertEqual(len(tables), len(heuristic_tables))
        self.assertEqual(diag.get("selection_mode"), "scored_promotion_v2")
        self.assertFalse(diag.get("selector_disabled"))

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_auto_selection_keeps_heuristic_for_normal_multi_table_layout(
        self,
        _build_embeddings,
    ) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)

        heuristic_tables = []
        for i, title in enumerate(["Retail Fees", "Corporate Fees", "Loan Fees"]):
            heuristic_tables.append(
                TablePayload(
                    order_index=i + 1,
                    title=title,
                    section_heading="",
                    page_number=1,
                    bbox={
                        "x0": 40.0,
                        "y0": 80.0 + (i * 180.0),
                        "x1": 560.0,
                        "y1": 220.0 + (i * 180.0),
                    },
                    rows=[
                        TableRowPayload(
                            row_index=row_index + 1,
                            page_number=1,
                            raw_text=f"{title} row {row_index + 1}",
                            metadata={"row_type": "data"},
                        )
                        for row_index in range(3)
                    ],
                    metadata={"detected_via": "heuristic"},
                )
            )

        geometry_table = TablePayload(
            order_index=1,
            title="Document snapshot",
            section_heading="",
            page_number=1,
            bbox={"x0": 60.0, "y0": 60.0, "x1": 360.0, "y1": 180.0},
            rows=[
                TableRowPayload(
                    row_index=row_index + 1,
                    page_number=1,
                    raw_text=f"snapshot row {row_index + 1}",
                    metadata={"row_type": "data"},
                )
                for row_index in range(4)
            ],
            metadata={"detected_via": "geometry"},
        )

        candidates = {
            "heuristic": heuristic_tables,
            "geometry": [geometry_table],
        }
        selected, _, diag = service._select_table_candidates(candidates)

        self.assertEqual(selected, "heuristic")
        self.assertEqual(diag.get("selection_mode"), "scored_promotion_v2")
        self.assertFalse(diag.get("selector_disabled"))

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_auto_selection_region_blend_prefers_keepable_region_candidate_over_document_winner(
        self,
        _build_embeddings,
    ) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)

        azure_page_one = TablePayload(
            order_index=11,
            title="Azure page 1",
            section_heading="",
            page_number=1,
            bbox={"x0": 10.0, "y0": 10.0, "x1": 300.0, "y1": 220.0},
            rows=[TableRowPayload(row_index=1, page_number=1, raw_text="azure-p1", metadata={"row_type": "data"})],
            metadata={"detected_via": "azure:layout"},
        )
        geometry_page_one = TablePayload(
            order_index=12,
            title="Geometry page 1",
            section_heading="",
            page_number=1,
            bbox={"x0": 12.0, "y0": 12.0, "x1": 302.0, "y1": 222.0},
            rows=[TableRowPayload(row_index=1, page_number=1, raw_text="geometry-p1", metadata={"row_type": "data"})],
            metadata={"detected_via": "geometry"},
        )
        azure_page_two = TablePayload(
            order_index=21,
            title="Azure page 2",
            section_heading="",
            page_number=2,
            bbox={"x0": 15.0, "y0": 20.0, "x1": 310.0, "y1": 250.0},
            rows=[TableRowPayload(row_index=1, page_number=2, raw_text="azure-p2", metadata={"row_type": "data"})],
            metadata={"detected_via": "azure:layout"},
        )
        geometry_page_two = TablePayload(
            order_index=22,
            title="Geometry page 2",
            section_heading="",
            page_number=2,
            bbox={"x0": 14.0, "y0": 18.0, "x1": 308.0, "y1": 248.0},
            rows=[TableRowPayload(row_index=1, page_number=2, raw_text="geometry-p2", metadata={"row_type": "data"})],
            metadata={"detected_via": "geometry"},
        )

        candidates = {
            "azure:layout": [azure_page_one, azure_page_two],
            "geometry": [geometry_page_one, geometry_page_two],
        }
        def _score_table_set(tables: list[TablePayload]) -> float:
            if tables is candidates["azure:layout"]:
                return 10.0
            return 5.0

        with mock.patch.object(service, "_score_table_set", side_effect=_score_table_set):
            selected, tables, diag = service._select_table_candidates(candidates)

        self.assertEqual(selected, "azure:layout")
        self.assertEqual([table.order_index for table in tables], [11, 21])
        self.assertEqual(diag.get("selection_mode"), "scored_promotion_v2")
        self.assertFalse(diag.get("selector_disabled"))


class GeometryLogicalRowReconstructionTests(SimpleTestCase):
    def _row(self, row_index: int, values: list[str]) -> TableRowPayload:
        cells = [
            TableCellPayload(
                row_index=row_index,
                column_index=idx,
                column_key=f"column_{idx+1}",
                raw_text=value,
                bbox={"x0": float(idx * 10), "y0": float(row_index * 10), "x1": float((idx * 10) + 8), "y1": float((row_index * 10) + 8)},
                metadata={"span_count": 1},
            )
            for idx, value in enumerate(values)
        ]
        return TableRowPayload(
            row_index=row_index,
            page_number=1,
            bbox={"x0": 0.0, "y0": float(row_index * 10), "x1": 100.0, "y1": float((row_index * 10) + 8)},
            raw_text=" | ".join(values),
            metadata={"row_type": "data"},
            cells=cells,
        )

    def _table(self, rows: list[TableRowPayload]) -> TablePayload:
        header = TableRowPayload(
            row_index=0,
            page_number=1,
            raw_text="Service | Tariff | Prime | Plus | Wealth | Exclusive | Private",
            metadata={"row_type": "header"},
            cells=[
                TableCellPayload(row_index=0, column_index=idx, column_key=f"column_{idx+1}", raw_text=value)
                for idx, value in enumerate(["Service", "Tariff", "Prime", "Plus", "Wealth", "Exclusive", "Private"])
            ],
        )
        return TablePayload(
            order_index=1,
            title="Geometry candidate",
            section_heading="",
            page_number=1,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["service", "tariff", "prime", "plus", "wealth", "exclusive", "private"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[header, *rows],
        )

    def _narrow_row(self, row_index: int, values: list[str]) -> TableRowPayload:
        cells = [
            TableCellPayload(
                row_index=row_index,
                column_index=idx,
                column_key=f"column_{idx+1}",
                raw_text=value,
            )
            for idx, value in enumerate(values)
        ]
        return TableRowPayload(
            row_index=row_index,
            page_number=1,
            raw_text=" | ".join(values),
            metadata={"row_type": "data"},
            cells=cells,
        )

    def _narrow_table(self, rows: list[TableRowPayload], *, order_index: int = 1) -> TablePayload:
        return TablePayload(
            order_index=order_index,
            title=f"Narrow {order_index}",
            section_heading="",
            page_number=1,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["column_1", "column_2"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=rows,
        )

    def test_geometry_logical_row_merge_combines_descriptor_fragments(self) -> None:
        reconstructor = GeometryTableReconstructor()
        rows = [
            self._row(1, ["", "Cheques Drawn", "", "", "", "", ""]),
            self._row(2, ["", "on CIB Branches", "", "", "", "", ""]),
            self._row(3, ["Collection of", "", "Min. USD 10 /", "Min. USD 10 /", "", "50% Discount", "50% Discount"]),
            self._row(4, ["", "", "Max. USD 100", "Max. USD 100", "10 / Max. USD", "+", "+"]),
        ]

        normalized, merged_pairs = reconstructor._normalize_geometry_logical_rows(rows, 7)

        self.assertEqual(merged_pairs, 2)
        self.assertEqual(len(normalized), 2)
        first = [cell.raw_text for cell in normalized[0].cells]
        second = [cell.raw_text for cell in normalized[1].cells]
        self.assertEqual(first[1], "Cheques Drawn on CIB Branches")
        self.assertEqual(second[0], "Collection of")
        self.assertIn("Min. USD 10 / Max. USD 100", second[2])
        self.assertTrue(normalized[0].metadata.get("geometry_logical_row_merged"))

    def test_geometry_logical_row_merge_does_not_collapse_complete_tier_rows(self) -> None:
        reconstructor = GeometryTableReconstructor()
        rows = [
            self._row(1, ["From 1 to 100 transactions", "No fees", "EGP 10"]),
            self._row(2, ["From 101 to 500 transactions", "No fees", "EGP 9"]),
        ]

        normalized, merged_pairs = reconstructor._normalize_geometry_logical_rows(rows, 3)

        self.assertEqual(merged_pairs, 0)
        self.assertEqual(len(normalized), 2)
        self.assertEqual([cell.raw_text for cell in normalized[0].cells], ["From 1 to 100 transactions", "No fees", "EGP 10"])

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_quality_scoring_penalizes_fragmented_geometry_rows(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)

        fragmented = self._table(
            [
                self._row(1, ["", "Cheques Drawn", "", "", "", "", ""]),
                self._row(2, ["", "on CIB Branches", "", "", "", "", ""]),
                self._row(3, ["Collection of", "", "Min. USD 10 /", "Min. USD 10 /", "", "50% Discount", "50% Discount"]),
                self._row(4, ["", "", "Max. USD 100", "Max. USD 100", "10 / Max. USD", "+", "+"]),
                self._row(5, ["Cheques favor of", "FCY Cheques Outside", "+", "+", "100 +", "Correspondent", "Correspondent"]),
                self._row(6, ["CIB's Customers", "CBE Clearing House", "Correspondent Fees", "Correspondent Fees", "Courier Fees", "50% Discount", "50% Discount"]),
            ]
        )
        clean = self._table(
            [
                self._row(1, ["Cheques Drawn on CIB Branches", "", "", "", "Free", "", ""]),
                self._row(2, ["LCY Cheques Inside CBE Clearing House (ATM Deposit: waived 50%)", "", "EGP 20", "EGP 20", "EGP 20", "Free", "Free"]),
                self._row(3, ["Collection of Cheques favor of CIB's Customers (Normal Collection)", "LCY Checks Outside CBE Clearing House (ATM Deposit: waived 50% from Commission)", "0.2% Min.EGP 20 / Max. EGP 400 + Correspondent Fees + Courier Fees", "0.2% Min.EGP 20 / Max. EGP 400 + Correspondent Fees + Courier Fees", "0.2% Min. EGP 20/Max. EGP 400 + Correspondent Fees + Courier Fees", "50% Discount + Correspondent Fees + Courier Fees", "50% Discount + Correspondent Fees + Courier Fees"]),
            ]
        )

        fragmented_quality = service._assess_table_quality(fragmented)
        clean_quality = service._assess_table_quality(clean)

        self.assertTrue(fragmented_quality["signals"].get("fragmented_logical_rows"))
        self.assertLess(fragmented_quality["quality_score"], clean_quality["quality_score"])

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_quality_scoring_flags_paragraph_like_tables(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        paragraph_like = self._narrow_table(
            [
                self._narrow_row(1, ["Agency Code: 51___", "Contract Number: ________"]),
                self._narrow_row(
                    2,
                    [
                        "Department Code: 3660___ __________ DDSOO consideration herein. this contract. services for two consecutive weeks.",
                        "Amendment Number: ________",
                    ],
                ),
                self._narrow_row(
                    3,
                    [
                        "This agreement may be amended pursuant to agreement of the parties and all related statutory approvals and notice periods applicable under state law.",
                        "The contractor shall comply with all laws, rules, ordinances and reporting obligations described in the agreement.",
                    ],
                ),
            ],
            order_index=9,
        )

        quality = service._assess_table_quality(paragraph_like)

        self.assertTrue(quality["signals"].get("paragraph_like_table"))
        self.assertGreaterEqual(int(quality["signals"].get("paragraph_like_rows") or 0), 2)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_quality_scoring_flags_low_structure_pdf_tables(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        low_structure = TablePayload(
            order_index=10,
            title="Admin header",
            section_heading="",
            page_number=2,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["column_1", "column_2", "column_3"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=2,
                    raw_text="DIVISION OF ACCOUNTS | Rev. | Page 2 of 2",
                    metadata={"row_type": "header"},
                    cells=[
                        TableCellPayload(row_index=0, column_index=0, column_key="column_1", raw_text="DIVISION OF ACCOUNTS"),
                        TableCellPayload(row_index=0, column_index=1, column_key="column_2", raw_text="Rev."),
                        TableCellPayload(row_index=0, column_index=2, column_key="column_3", raw_text="Page 2 of 2"),
                    ],
                ),
                self._row(1, ["STANDARD INVOICE", "Page", "of 2"]),
                self._row(2, ["QUANTITY", "AMOUNT - SUBTOTAL - $", ""]),
                self._row(
                    3,
                    [
                        "Submit invoice to billing address shown on contract immediately upon completing shipment of all items per contract.",
                        "",
                        "",
                    ],
                ),
            ],
        )

        quality = service._assess_table_quality(low_structure)

        self.assertTrue(quality["signals"].get("low_structure_table"))
        self.assertLessEqual(float(quality["signals"].get("structured_row_ratio") or 1.0), 0.45)
        self.assertGreaterEqual(float(quality["signals"].get("scaffold_row_ratio") or 0.0), 0.5)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_table_promotion_gate_suppresses_recurring_scaffold_but_keeps_useful_table(
        self,
        _build_embeddings,
    ) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        scaffold_one = self._narrow_table(
            [
                self._narrow_row(1, ["Agency Code: 51___", "Contract Number: ________"]),
                self._narrow_row(2, ["Department Code: 3660___ __________ DDSOO", "Amendment Number: ________"]),
                self._narrow_row(3, ["Placeholder body", "Placeholder body"]),
            ],
            order_index=1,
        )
        scaffold_two = self._narrow_table(
            [
                self._narrow_row(1, ["Agency Code: 51___", "Contract Number: ________"]),
                self._narrow_row(2, ["Department Code: 3660___ __________ DDSOO", "Amendment Number: ________"]),
                self._narrow_row(3, ["Another body", "Another body"]),
            ],
            order_index=2,
        )
        scaffold_three = self._narrow_table(
            [
                self._narrow_row(1, ["Agency Code: 51___", "Contract Number: ________"]),
                self._narrow_row(2, ["Department Code: 3660___ __________ DDSOO", "Amendment Number: ________"]),
                self._narrow_row(3, ["Yet another body", "Yet another body"]),
            ],
            order_index=3,
        )
        scaffold_four = self._narrow_table(
            [
                self._narrow_row(1, ["Agency Code: 51___", "Contract Number: ________"]),
                self._narrow_row(2, ["Department Code: 3660___ __________ DDSOO", "Amendment Number: ________"]),
                self._narrow_row(3, ["More repeated body", "More repeated body"]),
            ],
            order_index=4,
        )
        scaffold_five = self._narrow_table(
            [
                self._narrow_row(1, ["Agency Code: 51___", "Contract Number: ________"]),
                self._narrow_row(2, ["Department Code: 3660___ __________ DDSOO", "Amendment Number: ________"]),
                self._narrow_row(3, ["Final repeated body", "Final repeated body"]),
            ],
            order_index=5,
        )
        useful_toc = TablePayload(
            order_index=6,
            title="TOC",
            section_heading="",
            page_number=2,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["section", "title", "page"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=2,
                    raw_text="1.01 | THE PREMISES | 5",
                    metadata={"row_type": "header"},
                    cells=[
                        TableCellPayload(row_index=0, column_index=0, column_key="section", raw_text="Section"),
                        TableCellPayload(row_index=0, column_index=1, column_key="title", raw_text="Title"),
                        TableCellPayload(row_index=0, column_index=2, column_key="page", raw_text="Page"),
                    ],
                ),
                self._row(1, ["1.01", "THE PREMISES (OCT 2024)", "5"]),
                self._row(2, ["1.02", "EXPRESS APPURTENANT RIGHTS", "5"]),
            ],
        )

        kept, meta, issues = service._apply_pdf_table_promotion_gate(
            [scaffold_one, scaffold_two, scaffold_three, scaffold_four, scaffold_five, useful_toc]
        )

        self.assertEqual(len(kept), 1)
        self.assertEqual(kept[0].order_index, 6)
        self.assertEqual(meta.get("suppressed_tables"), 5)
        self.assertEqual(meta.get("class_counts", {}).get("recurring_scaffold"), 5)
        self.assertTrue(any(issue.code == "pdf_table_suppressed" for issue in issues))

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_page_blocks_trim_repeated_page_chrome_but_keep_body_text(
        self,
        _build_embeddings,
    ) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)

        def _page(page_number: int, body: str) -> PageLayout:
            return PageLayout(
                page_number=page_number,
                width=800.0,
                height=1000.0,
                rotation=0,
                text_density=0.4,
                has_ocr_content=False,
                content_type="text",
                blocks=[
                    PageBlockPayload(
                        block_type=KnowledgeBlockType.PARAGRAPH,
                        order_index=0,
                        text=(
                            f"Contract Number: ________ __________ DDSOO Amendment Number:________ "
                            f"B-{page_number} {body}"
                        ),
                        bbox={"x0": 0.0, "y0": 0.0, "x1": 700.0, "y1": 40.0},
                        metadata={"anchor": f"p{page_number}-b0", "page_region": "header"},
                    )
                ],
            )

        pages = [
            _page(1, "The contractor shall comply with all reporting obligations described herein."),
            _page(2, "The parties may amend the agreement subject to notice and approval requirements."),
            _page(3, "The state reserves all remedies available under law and contract."),
        ]

        segments = service._build_text_segments_from_blocks(pages, chunk_chars=400, overlap=0)

        self.assertEqual(len(segments), 3)
        rendered = "\n".join(segment["text"] for segment in segments)
        self.assertNotIn("Contract Number:", rendered)
        self.assertNotIn("DDSOO Amendment Number", rendered)
        self.assertIn("reporting obligations described herein", rendered)
        self.assertIn("state reserves all remedies available", rendered)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_page_blocks_trim_short_repeated_page_chrome_prefixes(
        self,
        _build_embeddings,
    ) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)

        def _page(page_number: int, body: str) -> PageLayout:
            return PageLayout(
                page_number=page_number,
                width=800.0,
                height=1000.0,
                rotation=0,
                text_density=0.4,
                has_ocr_content=False,
                content_type="text",
                blocks=[
                    PageBlockPayload(
                        block_type=KnowledgeBlockType.PARAGRAPH,
                        order_index=0,
                        text=f"Contract Number: ________ B-{page_number} {body}",
                        bbox={"x0": 0.0, "y0": 0.0, "x1": 700.0, "y1": 40.0},
                        metadata={"anchor": f"p{page_number}-b0", "page_region": "header"},
                    )
                ],
            )

        pages = [
            _page(1, "The contractor shall comply with all reporting obligations described herein."),
            _page(2, "The parties may amend the agreement subject to notice and approval requirements."),
            _page(3, "The state reserves all remedies available under law and contract."),
        ]

        segments = service._build_text_segments_from_blocks(pages, chunk_chars=400, overlap=0)

        self.assertEqual(len(segments), 3)
        rendered = "\n".join(segment["text"] for segment in segments)
        self.assertNotIn("Contract Number:", rendered)
        self.assertIn("reporting obligations described herein", rendered)
        self.assertIn("state reserves all remedies available", rendered)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_table_promotion_gate_suppresses_micro_fragment_and_bridge_tables(
        self,
        _build_embeddings,
    ) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)

        micro_fragment = TablePayload(
            order_index=7,
            title="Micro Fragment",
            section_heading="",
            page_number=1,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=[f"column_{idx}" for idx in range(8)],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                self._row(1, ["This", "ex", "a", "m", "ple", "mar", "ketin", "g"]),
                self._row(2, ["J", "&", "K", "Auto", "Re", "pa", "ir", "1"]),
                self._row(3, ["Dri", "vers", "on", "the", "near", "by", "high", "way"]),
            ],
        )

        bridge_table = self._narrow_table(
            [
                self._narrow_row(
                    1,
                    [
                        "At minimum, three professional references with email addresses and telephone numbers.",
                        "Attachment C",
                    ],
                ),
                self._narrow_row(
                    2,
                    [
                        "The response must clearly address evaluation criteria, costs, qualifications, and implementation.",
                        "Attachment D",
                    ],
                ),
            ],
            order_index=8,
        )

        useful_toc = TablePayload(
            order_index=9,
            title="TOC",
            section_heading="",
            page_number=2,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["section", "title", "page"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=2,
                    raw_text="1.01 | THE PREMISES | 5",
                    metadata={"row_type": "header"},
                    cells=[
                        TableCellPayload(row_index=0, column_index=0, column_key="section", raw_text="Section"),
                        TableCellPayload(row_index=0, column_index=1, column_key="title", raw_text="Title"),
                        TableCellPayload(row_index=0, column_index=2, column_key="page", raw_text="Page"),
                    ],
                ),
                self._row(1, ["1.01", "THE PREMISES (OCT 2024)", "5"]),
                self._row(2, ["1.02", "EXPRESS APPURTENANT RIGHTS", "5"]),
                self._row(3, ["1.03", "RENT AND OTHER CONSIDERATION", "5"]),
            ],
        )

        kept, meta, _issues = service._apply_pdf_table_promotion_gate(
            [micro_fragment, bridge_table, useful_toc]
        )

        self.assertEqual([table.order_index for table in kept], [9])
        self.assertEqual(meta.get("suppressed_tables"), 2)
        self.assertEqual(meta.get("class_counts", {}).get("strong_table"), 1)
        self.assertGreaterEqual(meta.get("class_counts", {}).get("layout_fragment", 0), 2)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_table_promotion_gate_suppresses_low_structure_pseudo_tables(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)

        narrative_pseudo = TablePayload(
            order_index=10,
            title="Narrative split",
            section_heading="",
            page_number=3,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["column_1", "column_2", "column_3"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                self._row(
                    1,
                    [
                        "Acquire a solution meeting the parameters, conditions, and mandatory requirements presented in the document.",
                        "AND",
                        "one electronic copy",
                    ],
                ),
                self._row(2, ["(PDF or Word) on a flash drive with any supplementary materials to: Purchasing Manager", "", ""]),
                self._row(3, ["12th Floor City Hall", "", ""]),
                self._row(4, ["455 N. Main", "", ""]),
            ],
        )
        admin_header = TablePayload(
            order_index=11,
            title="Admin header",
            section_heading="",
            page_number=2,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["column_1", "column_2", "column_3"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=2,
                    raw_text="DIVISION OF ACCOUNTS | Rev. | Page 2 of 2",
                    metadata={"row_type": "header"},
                    cells=[
                        TableCellPayload(row_index=0, column_index=0, column_key="column_1", raw_text="DIVISION OF ACCOUNTS"),
                        TableCellPayload(row_index=0, column_index=1, column_key="column_2", raw_text="Rev."),
                        TableCellPayload(row_index=0, column_index=2, column_key="column_3", raw_text="Page 2 of 2"),
                    ],
                ),
                self._row(1, ["STANDARD INVOICE", "Page", "of 2"]),
                self._row(2, ["QUANTITY", "AMOUNT - SUBTOTAL - $", ""]),
                self._row(
                    3,
                    [
                        "Submit invoice to billing address shown on contract immediately upon completing shipment of all items per contract.",
                        "",
                        "",
                    ],
                ),
            ],
        )
        useful_toc = TablePayload(
            order_index=12,
            title="TOC",
            section_heading="",
            page_number=2,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["section", "title", "page"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=2,
                    raw_text="Section | Title | Page",
                    metadata={"row_type": "header"},
                    cells=[
                        TableCellPayload(row_index=0, column_index=0, column_key="section", raw_text="Section"),
                        TableCellPayload(row_index=0, column_index=1, column_key="title", raw_text="Title"),
                        TableCellPayload(row_index=0, column_index=2, column_key="page", raw_text="Page"),
                    ],
                ),
                self._row(1, ["1.01", "THE PREMISES", "5"]),
                self._row(2, ["1.02", "EXPRESS APPURTENANT RIGHTS", "5"]),
                self._row(3, ["1.03", "RENT AND OTHER CONSIDERATION", "5"]),
            ],
        )

        kept, meta, _issues = service._apply_pdf_table_promotion_gate(
            [narrative_pseudo, admin_header, useful_toc]
        )

        self.assertEqual([table.order_index for table in kept], [12])
        self.assertEqual(meta.get("suppressed_tables"), 2)
        self.assertGreaterEqual(meta.get("class_counts", {}).get("layout_fragment", 0), 2)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_table_promotion_gate_suppresses_compact_banner_pseudo_tables(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)

        compact_banner = TablePayload(
            order_index=12,
            title="Invoice banner",
            section_heading="",
            page_number=1,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["column_1", "column_2", "column_3"],
            data_dictionary={},
            metadata={"detected_via": "azure_di"},
            rows=[
                self._row(0, ["KENTUCKY TRANSPORTATION CABINET TC 31-519", "10/2015 Rev.", "TC 31-519"]),
                self._row(
                    1,
                    ["DIVISION OF ACCOUNTS STANDARD INVOICE Page 1 of 2", "Rev.", "10/2015"],
                ),
                self._row(2, ["STANDARD INVOICE", "Page", "1 of 2"]),
            ],
        )
        useful_toc = TablePayload(
            order_index=13,
            title="TOC",
            section_heading="",
            page_number=2,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["section", "title", "page"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=2,
                    raw_text="Section | Title | Page",
                    metadata={"row_type": "header"},
                    cells=[
                        TableCellPayload(row_index=0, column_index=0, column_key="section", raw_text="Section"),
                        TableCellPayload(row_index=0, column_index=1, column_key="title", raw_text="Title"),
                        TableCellPayload(row_index=0, column_index=2, column_key="page", raw_text="Page"),
                    ],
                ),
                self._row(1, ["1.01", "THE PREMISES", "5"]),
                self._row(2, ["1.02", "EXPRESS APPURTENANT RIGHTS", "5"]),
                self._row(3, ["1.03", "RENT AND OTHER CONSIDERATION", "5"]),
            ],
        )

        compact_quality = service._assess_table_quality(compact_banner)
        self.assertTrue(compact_quality["signals"].get("compact_banner_table"))

        kept, meta, _issues = service._apply_pdf_table_promotion_gate([compact_banner, useful_toc])

        self.assertEqual([table.order_index for table in kept], [13])
        self.assertEqual(meta.get("suppressed_tables"), 1)
        self.assertGreaterEqual(meta.get("class_counts", {}).get("layout_fragment", 0), 1)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_table_promotion_gate_requires_positive_keep_evidence(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)

        contact_banner = TablePayload(
            order_index=13,
            title="Contact banner",
            section_heading="",
            page_number=1,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["column_1", "column_2", "column_3"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=1,
                    raw_text="Purchasing Department | th | Floor",
                    metadata={"row_type": "header"},
                    cells=[
                        TableCellPayload(row_index=0, column_index=0, column_key="column_1", raw_text="Purchasing Department City of Wichita, Kansas City Hall, 455 N Main 12th Floor"),
                        TableCellPayload(row_index=0, column_index=1, column_key="column_2", raw_text="th"),
                        TableCellPayload(row_index=0, column_index=2, column_key="column_3", raw_text="Floor"),
                    ],
                ),
                self._row(1, ["316-268-4636 https://ep.wichita.gov Melinda Walker - Purchasing Manager", "Proposers shall not", ""]),
            ],
        )
        attachment_block = TablePayload(
            order_index=14,
            title="Attachment block",
            section_heading="",
            page_number=7,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["column_1", "column_2"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=7,
                    raw_text="Long attachment instruction | Attachment A",
                    metadata={"row_type": "header"},
                    cells=[
                        TableCellPayload(
                            row_index=0,
                            column_index=0,
                            column_key="column_1",
                            raw_text=(
                                "2. The names of the staff members who will be available for work on the contract, "
                                "including a listing of their work experience (Attachment A). 3. The firm's relevant "
                                "experience, notably experience working with government agencies (Attachment B)."
                            ),
                        ),
                        TableCellPayload(row_index=0, column_index=1, column_key="column_2", raw_text="(Attachment A)."),
                    ],
                ),
                self._row(1, ["At minimum,", "three (3)"]),
                self._row(2, ["provided in this RFP.", "information for each requirement as listed herein."]),
            ],
        )
        useful_grid = TablePayload(
            order_index=15,
            title="Invoice lines",
            section_heading="",
            page_number=1,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["item", "description", "quantity", "unit", "unit_price", "amount"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=1,
                    raw_text="ITEM # | DESCRIPTION | QUANTITY | UNIT | UNIT PRICE | AMOUNT",
                    metadata={"row_type": "header"},
                    cells=[
                        TableCellPayload(row_index=0, column_index=0, column_key="item", raw_text="ITEM #"),
                        TableCellPayload(row_index=0, column_index=1, column_key="description", raw_text="DESCRIPTION"),
                        TableCellPayload(row_index=0, column_index=2, column_key="quantity", raw_text="QUANTITY"),
                        TableCellPayload(row_index=0, column_index=3, column_key="unit", raw_text="UNIT"),
                        TableCellPayload(row_index=0, column_index=4, column_key="unit_price", raw_text="UNIT PRICE"),
                        TableCellPayload(row_index=0, column_index=5, column_key="amount", raw_text="AMOUNT"),
                    ],
                ),
                self._row(1, ["1", "Road salt", "10", "EA", "$4.50", "$45.00"]),
                self._row(2, ["2", "Safety cones", "6", "EA", "$12.00", "$72.00"]),
            ],
        )

        kept, meta, _issues = service._apply_pdf_table_promotion_gate(
            [contact_banner, attachment_block, useful_grid]
        )

        self.assertEqual([table.order_index for table in kept], [15])
        self.assertEqual(meta.get("suppressed_tables"), 2)
        self.assertGreaterEqual(
            (meta.get("class_counts", {}).get("weak_table", 0))
            + (meta.get("class_counts", {}).get("layout_fragment", 0)),
            2,
        )

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_table_promotion_gate_recovers_best_bounded_fallback_when_all_candidates_suppress(
        self,
        _build_embeddings,
    ) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        candidate_one = self._narrow_table(
            [
                self._narrow_row(1, ["Issuance fees", "EGP 450"]),
                self._narrow_row(2, ["Replacement fees", "Free"]),
                self._narrow_row(3, ["Grace period", "55 days"]),
            ],
            order_index=20,
        )
        candidate_two = self._narrow_table(
            [
                self._narrow_row(1, ["Placeholder", "________"]),
                self._narrow_row(2, ["Template", "________"]),
            ],
            order_index=21,
        )

        assessments = {
            20: {"quality_score": 0.72, "signals": {}},
            21: {"quality_score": 0.31, "signals": {}},
        }

        def _classify(table, assessment, recurrence_stats=None):
            if table.order_index == 20:
                return "weak_table", "suppress", ["insufficient_keep_evidence"]
            return "layout_fragment", "suppress", ["compact_banner_table"]

        with (
            mock.patch.object(service, "_assess_table_quality", side_effect=lambda table: assessments[table.order_index]),
            mock.patch.object(service, "_classify_pdf_table_candidate", side_effect=_classify),
        ):
            kept, meta, issues = service._apply_pdf_table_promotion_gate([candidate_one, candidate_two])

        self.assertEqual([table.order_index for table in kept], [20])
        kept_meta = kept[0].metadata or {}
        self.assertEqual(kept_meta.get("promotion_decision"), "keep_fallback")
        self.assertTrue(kept_meta.get("promotion_fallback"))
        self.assertEqual((meta.get("fallback_kept_table") or {}).get("order_index"), 20)
        self.assertEqual(meta.get("suppressed_tables"), 1)
        self.assertFalse(any(issue.table_order_index == 20 for issue in issues))

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_table_promotion_gate_keeps_collapsed_factual_pdfplumber_table(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        collapsed_table = TablePayload(
            order_index=22,
            title="Collapsed factual",
            section_heading="",
            page_number=1,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["column_1"],
            data_dictionary={},
            metadata={"detected_via": "pdfplumber:lines_text"},
            rows=[
                self._narrow_row(1, ["Issuance and Renewal Fees EGP 500 EGP 250 EGP 300 EGP 350"]),
                self._narrow_row(2, ["Replacement Fees Free Free Free Free"]),
                self._narrow_row(3, ["Supplementary Fees EGP 25 Free Free Free"]),
                self._narrow_row(4, ["Late Payment Fees EGP 150 EGP 150 EGP 150 EGP 150"]),
            ],
        )

        assessment = {
            "quality_score": 0.58,
            "signals": {
                "effective_column_count": 1,
                "structured_row_ratio": 1.0,
                "placeholder_cell_ratio": 0.0,
                "header_confidence": 0.55,
                "paragraph_like_table": True,
            },
        }

        with mock.patch.object(service, "_assess_table_quality", return_value=assessment):
            kept, meta, issues = service._apply_pdf_table_promotion_gate([collapsed_table])

        self.assertEqual([table.order_index for table in kept], [22])
        self.assertEqual((kept[0].metadata or {}).get("promotion_class"), "collapsed_factual_table")
        self.assertEqual(meta.get("suppressed_tables"), 0)
        self.assertFalse(issues)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_table_promotion_gate_keeps_coherent_pdfplumber_native_text_matrix(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        matrix_table = TablePayload(
            order_index=23,
            title="Native text matrix",
            section_heading="",
            page_number=1,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["card_type", "white", "classic", "gold", "cash_back"],
            data_dictionary={},
            metadata={"detected_via": "pdfplumber:lines"},
            rows=[
                self._row(1, ["Issuance and Renewal Fees", "EGP 500", "EGP 250", "EGP 300", "EGP 350"]),
                self._row(2, ["Replacement Fees", "Free", "Free", "Free", "Free"]),
                self._row(3, ["Supplementary Fees", "EGP 25", "Free", "Free", "Free"]),
                self._row(4, ["Interest Rate", "3.99%", "3.99%", "3.99%", "3.99%"]),
                self._row(5, ["Over-Limit Fees", "EGP 150", "EGP 150", "EGP 150", "EGP 150"]),
                self._row(6, ["Late Payment Fees", "EGP 150", "EGP 150", "EGP 150", "EGP 150"]),
            ],
        )

        assessment = {
            "quality_score": 0.62,
            "signals": {
                "effective_column_count": 5,
                "structured_row_ratio": 0.72,
                "placeholder_cell_ratio": 0.0,
                "header_confidence": 0.35,
                "multi_cell_row_ratio": 1.0,
                "value_row_ratio": 0.9,
                "scaffold_row_ratio": 0.0,
            },
        }

        with mock.patch.object(service, "_assess_table_quality", return_value=assessment):
            kept, meta, issues = service._apply_pdf_table_promotion_gate([matrix_table])

        self.assertEqual([table.order_index for table in kept], [23])
        self.assertEqual((kept[0].metadata or {}).get("promotion_class"), "coherent_native_text_table")
        self.assertIn("coherent_pdfplumber_keep", (kept[0].metadata or {}).get("promotion_reasons") or [])
        self.assertEqual(meta.get("suppressed_tables"), 0)
        self.assertFalse(issues)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_table_promotion_gate_keeps_collapsed_uniform_value_matrix(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        swype_like_table = TablePayload(
            order_index=24,
            title="Uniform value matrix",
            section_heading="",
            page_number=1,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["column_1"],
            data_dictionary={},
            metadata={"detected_via": "pdfplumber:lines"},
            rows=[
                self._narrow_row(1, ["Card Type Swype 12 Swype 36"]),
                self._narrow_row(2, ["Issuance and Renewal Fees Free Across all Card Types"]),
                self._narrow_row(3, ["Replacement Fees Free Across all Card Types"]),
                self._narrow_row(4, ["Supplementary Card Fees Free Across all Card Types"]),
                self._narrow_row(5, ["Interest Rate (Installment Payment Plan) - Swype Only 2.67% Across all"]),
                self._narrow_row(6, ["Admin Fees on Installed Transaction- One time 2% of the installed transaction value"]),
                self._narrow_row(7, ["Cancellation Fees for Installed Transactions 5 % of the total outstanding transaction value"]),
                self._narrow_row(8, ["Interest Rate (Non-Installment Transactions ) 3.99% Across All"]),
                self._narrow_row(9, ["Credit Shield Subscription 0.9% monthly on the outstanding balance available with min. EGP 15"]),
                self._narrow_row(10, ["Fixed Card Insurance Fees (Solidarity Fees)****** Free across all Swype card types."]),
                self._narrow_row(11, ["Grace Period 40 Days on Purchase Only"]),
                self._narrow_row(12, ["Over Limit Fees EGP 150"]),
                self._narrow_row(13, ["Late Payments Fees EGP 150"]),
            ],
        )

        assessment = {
            "quality_score": 0.78,
            "signals": {
                "effective_column_count": 1,
                "nonsense_columns": True,
                "structured_row_ratio": 0.67,
                "placeholder_cell_ratio": 0.0,
                "header_confidence": 0.0,
                "multi_cell_row_ratio": 0.0,
                "value_row_ratio": 0.57,
                "scaffold_row_ratio": 1.0,
                "row_consistency": 1.0,
                "cell_fill_ratio": 1.0,
                "long_cell_ratio": 0.19,
                "max_cell_word_count": 13,
            },
        }

        with (
            mock.patch.object(service, "_assess_table_quality", return_value=assessment),
            mock.patch.object(service, "_table_is_collapsed_factual", return_value=False),
        ):
            kept, meta, issues = service._apply_pdf_table_promotion_gate([swype_like_table])

        self.assertEqual([table.order_index for table in kept], [24])
        self.assertEqual((kept[0].metadata or {}).get("promotion_class"), "collapsed_uniform_value_matrix")
        self.assertIn("collapsed_uniform_matrix_keep", (kept[0].metadata or {}).get("promotion_reasons") or [])
        self.assertEqual(meta.get("suppressed_tables"), 0)
        self.assertFalse(issues)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_repair_pdf_table_structure_rebuilds_collapsed_native_text_matrix_from_spans(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)

        collapsed = TablePayload(
            order_index=31,
            title="Collapsed matrix",
            section_heading="",
            page_number=1,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["card_type_white_classic_gold_cash_back"],
            data_dictionary={},
            metadata={"detected_via": "pdfplumber:lines"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=1,
                    raw_text="Card Type White Classic Gold Cash Back",
                    metadata={"row_type": "header"},
                    cells=[TableCellPayload(row_index=0, column_index=0, column_key="card_type_white_classic_gold_cash_back", raw_text="Card Type White Classic Gold Cash Back")],
                ),
                self._narrow_row(1, ["Issuance and Renewal Fees EGP 500 EGP 250 EGP 300 EGP 350"]),
                self._narrow_row(2, ["Replacement Fees Free Free Free Free"]),
                self._narrow_row(3, ["Supplementary Fees EGP 25 Free Free Free"]),
            ],
        )

        page_spans = [[
            PdfSpan(page_number=1, text="White", x0=200.0, y0=90.0, x1=240.0, y1=100.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Classic", x0=280.0, y0=90.0, x1=330.0, y1=100.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Gold", x0=360.0, y0=90.0, x1=395.0, y1=100.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Cash Back", x0=440.0, y0=90.0, x1=500.0, y1=100.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Card Type", x0=34.0, y0=96.0, x1=78.0, y1=106.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="E-Commerce", x0=520.0, y0=96.0, x1=590.0, y1=106.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Issuance and Renewal Fees", x0=34.0, y0=118.0, x1=170.0, y1=128.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="EGP 500", x0=200.0, y0=118.0, x1=240.0, y1=128.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="EGP 250", x0=280.0, y0=118.0, x1=320.0, y1=128.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="EGP 300", x0=360.0, y0=118.0, x1=400.0, y1=128.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="EGP 350", x0=440.0, y0=118.0, x1=480.0, y1=128.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="EGP 500", x0=520.0, y0=118.0, x1=560.0, y1=128.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Replacement Fees", x0=34.0, y0=140.0, x1=140.0, y1=150.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Free", x0=200.0, y0=140.0, x1=228.0, y1=150.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Free", x0=280.0, y0=140.0, x1=308.0, y1=150.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Free", x0=360.0, y0=140.0, x1=388.0, y1=150.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Free", x0=440.0, y0=140.0, x1=468.0, y1=150.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Free", x0=520.0, y0=140.0, x1=548.0, y1=150.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Supplementary Fees", x0=34.0, y0=162.0, x1=150.0, y1=172.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="EGP 25", x0=200.0, y0=162.0, x1=236.0, y1=172.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Free", x0=280.0, y0=162.0, x1=308.0, y1=172.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Free", x0=360.0, y0=162.0, x1=388.0, y1=172.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="Free", x0=440.0, y0=162.0, x1=468.0, y1=172.0, font=None, size=9.0),
            PdfSpan(page_number=1, text="N/A", x0=520.0, y0=162.0, x1=545.0, y1=172.0, font=None, size=9.0),
        ]]

        repaired, meta, issues = service._repair_pdf_table_structure(
            [collapsed],
            candidates={},
            page_spans=page_spans,
            selected_extractor="pdfplumber:lines",
            selection_context={"format_hint": "pdf", "native_text_pdf": True, "pdf_lane": "native_text"},
        )

        self.assertEqual(len(repaired), 1)
        repaired_table = repaired[0]
        self.assertTrue((repaired_table.metadata or {}).get("structure_reconstructed"))
        self.assertEqual((repaired_table.metadata or {}).get("canonical_acceptance"), "accepted")
        self.assertEqual(
            repaired_table.column_schema,
            ["card_type", "white", "classic", "gold", "cash_back", "e_commerce"],
        )
        issuance_row = repaired_table.rows[1]
        self.assertEqual([cell.raw_text for cell in issuance_row.cells], [
            "Issuance and Renewal Fees",
            "EGP 500",
            "EGP 250",
            "EGP 300",
            "EGP 350",
            "EGP 500",
        ])
        self.assertEqual((meta.get("reconstruction") or {}).get("reconstructed_tables"), 1)
        self.assertFalse(any(issue.code == "pdf_table_not_canonical" for issue in issues))

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_repair_pdf_table_structure_drops_unreconstructed_collapsed_matrix_from_canonical_persistence(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)
        collapsed = TablePayload(
            order_index=32,
            title="Unsafe collapsed matrix",
            section_heading="",
            page_number=1,
            bbox={},
            column_schema=["column_1"],
            data_dictionary={},
            metadata={"detected_via": "pdfplumber:lines"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=1,
                    raw_text="Card Type White Classic Gold Cash Back",
                    metadata={"row_type": "header"},
                    cells=[TableCellPayload(row_index=0, column_index=0, column_key="column_1", raw_text="Card Type White Classic Gold Cash Back")],
                ),
                self._narrow_row(1, ["Issuance and Renewal Fees EGP 500 EGP 250 EGP 300 EGP 350"]),
                self._narrow_row(2, ["Replacement Fees Free Free Free Free"]),
                self._narrow_row(3, ["Supplementary Fees EGP 25 Free Free Free"]),
            ],
        )

        repaired, meta, issues = service._repair_pdf_table_structure(
            [collapsed],
            candidates={},
            selected_extractor="pdfplumber:lines",
            selection_context={"format_hint": "pdf", "native_text_pdf": True, "pdf_lane": "native_text"},
        )

        self.assertEqual(repaired, [])
        self.assertEqual((meta.get("acceptance") or {}).get("dropped_tables"), 1)
        self.assertTrue(any(issue.code == "pdf_table_not_canonical" for issue in issues))

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_table_promotion_gate_keeps_fillable_transactional_grid(self, _build_embeddings) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)

        fillable_grid = TablePayload(
            order_index=16,
            title="Invoice grid",
            section_heading="",
            page_number=1,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["column_1", "column_2", "column_3"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=1,
                    raw_text="DIVISION OF ACCOUNTS STANDARD INVOICE Page 1 of 2 | 10/2015 Rev. | 10/2015",
                    metadata={"row_type": "header"},
                    cells=[
                        TableCellPayload(row_index=0, column_index=0, column_key="column_1", raw_text="DIVISION OF ACCOUNTS STANDARD INVOICE Page 1 of 2"),
                        TableCellPayload(row_index=0, column_index=1, column_key="column_2", raw_text="10/2015 Rev."),
                        TableCellPayload(row_index=0, column_index=2, column_key="column_3", raw_text="10/2015"),
                    ],
                ),
                self._row(1, ["STANDARD INVOICE CITY", "Page ZIP", "of 2"]),
                self._row(2, ["QUANTITY TITLE DATE", "", "AMOUNT -"]),
            ],
        )
        admin_followup = TablePayload(
            order_index=17,
            title="Admin continuation",
            section_heading="",
            page_number=2,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["column_1", "column_2", "column_3"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=2,
                    raw_text="DIVISION OF ACCOUNTS | Rev. | Page 2 of 2",
                    metadata={"row_type": "header"},
                    cells=[
                        TableCellPayload(row_index=0, column_index=0, column_key="column_1", raw_text="DIVISION OF ACCOUNTS"),
                        TableCellPayload(row_index=0, column_index=1, column_key="column_2", raw_text="Rev."),
                        TableCellPayload(row_index=0, column_index=2, column_key="column_3", raw_text="Page 2 of 2"),
                    ],
                ),
                self._row(1, ["STANDARD INVOICE 2", "Page", "of 2"]),
                self._row(2, ["QUANTITY", "", "AMOUNT - SUBTOTAL - $"]),
                self._row(3, ["Submit invoice to billing address shown on contract immediately upon completing shipment of all items per agreement.", "", ""]),
            ],
        )

        kept, meta, _issues = service._apply_pdf_table_promotion_gate([fillable_grid, admin_followup])

        self.assertEqual([table.order_index for table in kept], [16])
        self.assertEqual(meta.get("suppressed_tables"), 1)
        self.assertEqual((kept[0].metadata or {}).get("promotion_class"), "strong_table")

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_pdf_table_promotion_gate_suppresses_prefix_recurrent_scaffolds_and_single_row_bridges(
        self,
        _build_embeddings,
    ) -> None:
        service = KnowledgeIngestionService(enable_ocr=False)

        recurring_tables = [
            self._narrow_table(
                [
                    self._narrow_row(
                        1,
                        [f"Agency Code: 51___ Department Code: 3660___ __________ DDSOO {suffix}", "Contract Number: ________"],
                    ),
                    self._narrow_row(2, ["Amendment Number: ________", ""]),
                    self._narrow_row(3, [body, ""]),
                ],
                order_index=index,
            )
            for index, suffix, body in [
                (1, "B.", "Placeholder body"),
                (2, "", "Another body"),
                (3, "F.", "Yet another body"),
                (4, "I.", "More repeated body"),
                (5, "X.", "Final repeated body"),
            ]
        ]

        single_row_bridge = TablePayload(
            order_index=6,
            title="Attachment-style bridge",
            section_heading="",
            page_number=7,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["column_1", "column_2", "column_3"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                self._row(
                    1,
                    [
                        "1. Acquire a solution meeting the parameters, conditions, and mandatory requirements presented in the document. Submit one original and one electronic copy with all supplementary materials.",
                        "AND",
                        "one electronic copy",
                    ],
                ),
            ],
        )

        useful_toc = TablePayload(
            order_index=7,
            title="TOC",
            section_heading="",
            page_number=2,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 100.0, "y1": 100.0},
            column_schema=["section", "title", "page"],
            data_dictionary={},
            metadata={"detected_via": "geometry"},
            rows=[
                TableRowPayload(
                    row_index=0,
                    page_number=2,
                    raw_text="1.01 | THE PREMISES | 5",
                    metadata={"row_type": "header"},
                    cells=[
                        TableCellPayload(row_index=0, column_index=0, column_key="section", raw_text="Section"),
                        TableCellPayload(row_index=0, column_index=1, column_key="title", raw_text="Title"),
                        TableCellPayload(row_index=0, column_index=2, column_key="page", raw_text="Page"),
                    ],
                ),
                self._row(1, ["1.01", "THE PREMISES (OCT 2024)", "5"]),
                self._row(2, ["1.02", "EXPRESS APPURTENANT RIGHTS", "5"]),
                self._row(3, ["1.03", "RENT AND OTHER CONSIDERATION", "5"]),
            ],
        )

        kept, meta, _issues = service._apply_pdf_table_promotion_gate(
            [*recurring_tables, single_row_bridge, useful_toc]
        )

        self.assertEqual([table.order_index for table in kept], [7])
        self.assertEqual(meta.get("suppressed_tables"), 6)
        self.assertGreaterEqual(meta.get("class_counts", {}).get("recurring_scaffold", 0), 5)
        self.assertGreaterEqual(meta.get("class_counts", {}).get("layout_fragment", 0), 1)


class KnowledgeIngestionJsonTests(TestCase):
    def setUp(self) -> None:
        super().setUp()
        self._media_root = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self._media_root)
        self.override = override_settings(MEDIA_ROOT=self._media_root)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.user = User.objects.create(email="ingest@example.com", first_name="Test")
        self.registration = RegistrationSession.objects.create(user=self.user)
        self.business = BusinessProfile.objects.create(
            user=self.user,
            registration_session=self.registration,
            name="Travel Co",
            industry="travel",
            metadata={"ingest_max_json_entities": 2},
        )

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_json_ingestion_persists_entities_and_aliases(self, _build_embeddings):
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Trips",
        )
        storage_path = Path("uploads/sample.json")
        target_path = Path(self._media_root) / storage_path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        payload = {
            "trips": [
                {"slug": "trip-1", "name": "Trip One", "code": "TR-ONE"},
                {"slug": "trip-2", "name": "Trip Two", "code": "TR-TWO"},
                {"slug": "trip-3", "name": "Trip Three", "code": "TR-THREE"},
            ]
        }
        target_path.write_text(json.dumps(payload), encoding="utf-8")
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="sample.json",
            storage_path=str(storage_path),
            content_type="application/json",
            size_bytes=target_path.stat().st_size,
        )

        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        extraction = service._extract_upload(upload)
        service._persist_extraction(upload, extraction)

        upload.refresh_from_db()
        self.assertEqual(KnowledgeEntity.objects.filter(upload=upload).count(), 2)
        alias_count = KnowledgeAlias.objects.filter(entity__upload=upload).count()
        self.assertGreater(alias_count, 0)
        self.assertEqual(upload.ingestion_metadata.get("alias_count"), alias_count)
        self.assertEqual(upload.ingestion_metadata.get("truncated_entities"), 2)


class KnowledgeIngestionTextTests(TestCase):
    def setUp(self) -> None:
        super().setUp()
        self._media_root = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self._media_root)
        self.override = override_settings(MEDIA_ROOT=self._media_root)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.user = User.objects.create(email="text@example.com", first_name="Text")
        self.registration = RegistrationSession.objects.create(user=self.user)
        self.business = BusinessProfile.objects.create(
            user=self.user,
            registration_session=self.registration,
            name="Docs Co",
            industry="support",
        )

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_text_ingestion_persists_chunks(self, _build_embeddings) -> None:
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.TEXT,
            status=KnowledgeStatus.PENDING,
            display_name="Manual snippet",
        )
        KnowledgeUploadText.objects.create(
            upload=upload,
            content="Refund policy:\n- Refunds are allowed within 14 days of purchase.\n",
        )
        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        with tenant_context(self.business.id):
            extraction = service._extract_upload(upload)
            self.assertEqual(extraction.format_hint, "text")
            service._persist_extraction(upload, extraction)

            upload.refresh_from_db()
            self.assertEqual(upload.status, KnowledgeStatus.ACTIVE)
            self.assertGreater(int(upload.chunk_count or 0), 0)
            self.assertTrue(KnowledgeUploadChunk.objects.filter(upload=upload).exists())
            self.assertEqual(upload.ingestion_metadata.get("format"), "text")
            self.assertEqual(upload.ingestion_metadata.get("content_type"), "text/plain")


class KnowledgeIngestionSpreadsheetTests(TestCase):
    def setUp(self) -> None:
        super().setUp()
        self._media_root = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self._media_root)
        self.override = override_settings(MEDIA_ROOT=self._media_root)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.user = User.objects.create(email="spreadsheet@example.com", first_name="Sheet")
        self.registration = RegistrationSession.objects.create(user=self.user)
        self.business = BusinessProfile.objects.create(
            user=self.user,
            registration_session=self.registration,
            name="Data Co",
            industry="finance",
        )

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_csv_ingestion_creates_tables(self, _build_embeddings):
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Plans CSV",
        )
        storage_path = Path("uploads/plans.csv")
        target_path = Path(self._media_root) / storage_path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        target_path.write_text("plan,price\nBasic,10\nPro,25\n", encoding="utf-8")
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="plans.csv",
            storage_path=str(storage_path),
            content_type="text/csv",
            size_bytes=target_path.stat().st_size,
        )
        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        extraction = service._extract_upload(upload)
        service._persist_extraction(upload, extraction)
        upload.refresh_from_db()
        tables = KnowledgeUploadTable.objects.filter(upload=upload)
        self.assertEqual(tables.count(), 1)
        rows = KnowledgeUploadTableRow.objects.filter(table__upload=upload)
        self.assertEqual(rows.count(), 2)
        self.assertEqual(tables.first().column_schema, ["plan", "price"])
        entities = KnowledgeEntity.objects.filter(upload=upload)
        self.assertEqual(entities.count(), 2)
        self.assertTrue(
            KnowledgeAlias.objects.filter(entity__upload=upload, alias_normalized="basic").exists()
        )
        table_stats = upload.ingestion_metadata.get("table_stats") or {}
        self.assertEqual(table_stats.get("total_rows"), 2)
        self.assertEqual(table_stats.get("indexed_rows"), 2)
        self.assertEqual(table_stats.get("row_cap"), 2)
        self.assertEqual(table_stats.get("row_tier"), "small")
        self.assertFalse(table_stats.get("partial_index"))
        self.assertFalse(KnowledgeUploadIssue.objects.filter(upload=upload).exists())

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    @override_settings(RAG_TABLE_CHILD_MAX_ROWS=2, RAG_TABLE_SUMMARY_MAX_ROW_LABELS=5)
    def test_large_table_row_chunks_are_sharded_and_summary_is_emitted_per_shard(self, _build_embeddings):
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Shard CSV",
        )
        storage_path = Path("uploads/shard.csv")
        target_path = Path(self._media_root) / storage_path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        target_path.write_text(
            "plan,price\nBasic,10\nPro,25\nGrowth,40\nScale,55\nEnterprise,99\n",
            encoding="utf-8",
        )
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="shard.csv",
            storage_path=str(storage_path),
            content_type="text/csv",
            size_bytes=target_path.stat().st_size,
        )

        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        with tenant_context(self.business.id):
            extraction = service._extract_upload(upload)
            service._persist_extraction(upload, extraction)
            service.reingest_from_persisted_artifacts(upload)

        row_chunks = list(
            KnowledgeUploadChunk.objects.filter(upload=upload, metadata__table_chunk_role="row").order_by("chunk_index")
        )
        self.assertEqual(len(row_chunks), 5)
        shard_indices = {
            int((chunk.metadata or {}).get("table_row_shard_index"))
            for chunk in row_chunks
            if (chunk.metadata or {}).get("table_row_shard_index") is not None
        }
        self.assertEqual(shard_indices, {0, 1, 2})

        summary_chunks = list(
            KnowledgeUploadChunk.objects.filter(upload=upload, metadata__table_chunk_role="summary").order_by("chunk_index")
        )
        self.assertEqual(len(summary_chunks), 3)
        summary_shards = sorted(
            int((chunk.metadata or {}).get("table_row_shard_index"))
            for chunk in summary_chunks
            if (chunk.metadata or {}).get("table_row_shard_index") is not None
        )
        self.assertEqual(summary_shards, [0, 1, 2])
        self.assertTrue(
            all((chunk.metadata or {}).get("table_shard_count") == 3 for chunk in summary_chunks)
        )

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_reingest_from_persisted_artifacts_rebuilds_canonical_chunk_metadata(self, _build_embeddings):
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Plans CSV",
        )
        storage_path = Path("uploads/plans.csv")
        target_path = Path(self._media_root) / storage_path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        target_path.write_text("plan,price\nBasic,10\nPro,25\n", encoding="utf-8")
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="plans.csv",
            storage_path=str(storage_path),
            content_type="text/csv",
            size_bytes=target_path.stat().st_size,
        )

        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        with tenant_context(self.business.id):
            extraction = service._extract_upload(upload)
            service._persist_extraction(upload, extraction)
            service.reingest_from_persisted_artifacts(upload)
            upload.refresh_from_db()
            chunks = list(KnowledgeUploadChunk.objects.filter(upload=upload).order_by("chunk_index"))
        self.assertEqual(upload.chunk_count, len(chunks))
        self.assertTrue(chunks)
        self.assertTrue(all((chunk.metadata or {}).get("canonical_anchor_id") for chunk in chunks))
        kinds = {(chunk.metadata or {}).get("canonical_chunk_kind") for chunk in chunks}
        self.assertIn("table_row", kinds)
        self.assertIn("table_summary", kinds)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_table_entities_persist_when_entity_chunking_disabled(self, _build_embeddings):
        self.business.metadata = {"features": {"entity_chunking": False}}
        self.business.save(update_fields=["metadata"])

        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Plans CSV",
        )
        storage_path = Path("uploads/plans.csv")
        target_path = Path(self._media_root) / storage_path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        target_path.write_text("plan,price\nBasic,10\nPro,25\n", encoding="utf-8")
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="plans.csv",
            storage_path=str(storage_path),
            content_type="text/csv",
            size_bytes=target_path.stat().st_size,
        )

        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        extraction = service._extract_upload(upload)
        service._persist_extraction(upload, extraction)
        upload.refresh_from_db()

        self.assertGreater(KnowledgeEntity.objects.filter(upload=upload).count(), 0)
        self.assertGreater(KnowledgeAlias.objects.filter(entity__upload=upload).count(), 0)

    def test_persist_entities_clamps_overlong_entity_fields(self):
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Long Entity Spreadsheet",
        )

        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        long_entity_type = "t" * 180
        long_entity_name = "invoice-row-" + ("x" * 400)
        service._persist_entities(
            upload,
            entities=[
                {
                    "entity_index": 0,
                    "entity_type": long_entity_type,
                    "entity_name": long_entity_name,
                    "attributes": {"description": "x" * 512},
                    "columns": ["description"],
                    "table_metadata": {"source": "spreadsheet"},
                    "aliases": [],
                }
            ],
            chunks=[],
        )

        entity = KnowledgeEntity.objects.get(upload=upload)
        self.assertEqual(len(entity.entity_type), 120)
        self.assertEqual(len(entity.entity_name), 255)
        self.assertEqual(len(entity.primary_label), 255)
        self.assertEqual(entity.metadata.get("entity_type_truncated"), long_entity_type)
        self.assertEqual(entity.metadata.get("entity_name_truncated"), long_entity_name)
        self.assertEqual(entity.metadata.get("primary_label_truncated"), long_entity_name)

    def test_detect_format_prefers_csv_over_text_content_type(self):
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Plans CSV",
        )
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="plans.csv",
            storage_path="uploads/plans.csv",
            content_type="text/csv",
            size_bytes=123,
        )
        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        file_detail = upload.file_detail
        self.assertEqual(service._detect_format(file_detail), "csv")

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    @override_settings(
        TABLE_MAX_ROWS_DEFAULT=3,
        RAG_TABLE_SMALL_ROW_LIMIT=2,
        RAG_TABLE_LARGE_ROW_LIMIT=4,
        RAG_TABLE_MAX_HARD_CAP=100,
    )
    def test_large_csv_default_path_keeps_rows_with_pagination_contract(self, _build_embeddings):
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Large CSV",
        )
        storage_path = Path("uploads/large.csv")
        target_path = Path(self._media_root) / storage_path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        rows = ["plan,price"]
        for idx in range(1, 7):
            rows.append(f"Tier{idx},{idx * 10}")
        target_path.write_text("\n".join(rows) + "\n", encoding="utf-8")
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="large.csv",
            storage_path=str(storage_path),
            content_type="text/csv",
            size_bytes=target_path.stat().st_size,
        )
        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        extraction = service._extract_upload(upload)
        service._persist_extraction(upload, extraction)
        upload.refresh_from_db()
        self.assertEqual(KnowledgeUploadTableRow.objects.filter(table__upload=upload).count(), 6)
        table_truncation = upload.ingestion_metadata.get("table_truncation") or {}
        self.assertEqual(table_truncation.get("truncated_rows"), 0)
        table_stats = upload.ingestion_metadata.get("table_stats") or {}
        self.assertFalse(table_stats.get("partial_index"))
        self.assertEqual(table_stats.get("row_cap"), 6)
        self.assertEqual(table_stats.get("row_tier"), "large")
        self.assertFalse(KnowledgeUploadIssue.objects.filter(upload=upload, issue_code="table_rows_truncated").exists())

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    @override_settings(
        TABLE_MAX_ROWS_DEFAULT=3,
        RAG_TABLE_SMALL_ROW_LIMIT=2,
        RAG_TABLE_LARGE_ROW_LIMIT=4,
        RAG_TABLE_MAX_HARD_CAP=5,
    )
    def test_large_csv_hard_cap_still_emits_truncation_issue(self, _build_embeddings):
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Large CSV hard cap",
        )
        storage_path = Path("uploads/large-hard-cap.csv")
        target_path = Path(self._media_root) / storage_path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        rows = ["plan,price"]
        for idx in range(1, 7):
            rows.append(f"Tier{idx},{idx * 10}")
        target_path.write_text("\n".join(rows) + "\n", encoding="utf-8")
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="large-hard-cap.csv",
            storage_path=str(storage_path),
            content_type="text/csv",
            size_bytes=target_path.stat().st_size,
        )

        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        extraction = service._extract_upload(upload)
        service._persist_extraction(upload, extraction)
        upload.refresh_from_db()

        self.assertEqual(KnowledgeUploadTableRow.objects.filter(table__upload=upload).count(), 5)
        table_truncation = upload.ingestion_metadata.get("table_truncation") or {}
        self.assertEqual(table_truncation.get("truncated_rows"), 1)
        table_stats = upload.ingestion_metadata.get("table_stats") or {}
        self.assertTrue(table_stats.get("partial_index"))
        self.assertEqual(table_stats.get("row_cap"), 5)
        self.assertEqual(table_stats.get("row_tier"), "large")
        self.assertTrue(KnowledgeUploadIssue.objects.filter(upload=upload, issue_code="table_rows_truncated").exists())

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_xlsx_ingestion_creates_tables(self, _build_embeddings):
        if Workbook is None:
            self.skipTest("openpyxl not installed")
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Plans XLSX",
        )
        storage_path = Path("uploads/plans.xlsx")
        target_path = Path(self._media_root) / storage_path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Plans"
        sheet.append(["plan", "price"])
        sheet.append(["Starter", 5])
        sheet.append(["Enterprise", 99])
        workbook.save(target_path)
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="plans.xlsx",
            storage_path=str(storage_path),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=target_path.stat().st_size,
        )
        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        extraction = service._extract_upload(upload)
        service._persist_extraction(upload, extraction)
        upload.refresh_from_db()
        tables = KnowledgeUploadTable.objects.filter(upload=upload)
        self.assertEqual(tables.count(), 1)
        rows = KnowledgeUploadTableRow.objects.filter(table__upload=upload)
        self.assertEqual(rows.count(), 2)
        self.assertEqual(tables.first().column_schema, ["plan", "price"])
        entities = KnowledgeEntity.objects.filter(upload=upload)
        self.assertEqual(entities.count(), 2)
        self.assertTrue(
            KnowledgeAlias.objects.filter(entity__upload=upload, alias_normalized="starter").exists()
        )
        table_stats = upload.ingestion_metadata.get("table_stats") or {}
        self.assertEqual(table_stats.get("total_rows"), 2)
        self.assertEqual(table_stats.get("indexed_rows"), 2)
        self.assertEqual(table_stats.get("row_cap"), 2)
        self.assertEqual(table_stats.get("row_tier"), "small")
        self.assertFalse(table_stats.get("partial_index"))

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_xlsx_ingestion_drops_hidden_template_rows_and_scaffold_columns(self, _build_embeddings):
        if Workbook is None:
            self.skipTest("openpyxl not installed")
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Invoice Template XLSX",
        )
        storage_path = Path("uploads/invoice-template.xlsx")
        target_path = Path(self._media_root) / storage_path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Equipment"
        sheet.append(["Code", "Amount", "SANDBOX AREA - rough work only"])
        sheet.append(["E-1", 1250, ""])
        sheet.append([0, 0, 0])
        sheet.row_dimensions[3].hidden = True
        workbook.save(target_path)
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="invoice-template.xlsx",
            storage_path=str(storage_path),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=target_path.stat().st_size,
        )

        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        extraction = service._extract_upload(upload)
        service._persist_extraction(upload, extraction)
        upload.refresh_from_db()

        table = KnowledgeUploadTable.objects.get(upload=upload)
        self.assertEqual(table.column_schema, ["Code", "Amount"])
        rows = list(KnowledgeUploadTableRow.objects.filter(table=table).order_by("row_index"))
        self.assertEqual(len(rows), 1)
        row_cells = list(rows[0].cells.order_by("column_index").values_list("raw_text", flat=True))
        self.assertEqual(row_cells, ["E-1", "1250"])
        self.assertEqual(KnowledgeEntity.objects.filter(upload=upload).count(), 1)
        normalization = upload.ingestion_metadata.get("normalization") or {}
        self.assertEqual((normalization.get("hidden_rows_dropped") or {}).get("total"), 1)
        self.assertEqual((normalization.get("scaffold_columns_trimmed") or {}).get("total"), 1)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_xlsx_entity_promotion_keeps_real_rows_and_skips_template_rows(self, _build_embeddings):
        if Workbook is None:
            self.skipTest("openpyxl not installed")
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Invoice Entity Filtering XLSX",
        )
        storage_path = Path("uploads/invoice-entities.xlsx")
        target_path = Path(self._media_root) / storage_path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()

        instructions = workbook.active
        instructions.title = "Instructions"
        instructions.append(["Template Version", "09/23/25"])
        instructions.append(["Workbook Instructions", "Fill all required fields"])

        subs = workbook.create_sheet("Subs & Vendors")
        subs.append(["Reference", "Name", "Amount", "Approval"])
        subs.append(["S-1", "Alpha Security, Inc.", 19382.36, "Yes"])
        subs.append(["Vendor 2", 0, 0, "Select Yes or No"])

        lists = workbook.create_sheet("Lists")
        lists.sheet_state = "hidden"
        lists.append(["Choices"])
        lists.append(["Select Yes or No"])
        lists.append(["Approved"])

        workbook.save(target_path)
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="invoice-entities.xlsx",
            storage_path=str(storage_path),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=target_path.stat().st_size,
        )

        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        extraction = service._extract_upload(upload)
        service._persist_extraction(upload, extraction)

        entities = list(KnowledgeEntity.objects.filter(upload=upload).values_list("entity_name", flat=True))
        self.assertEqual(entities, ["S-1"])
        self.assertFalse(
            KnowledgeEntity.objects.filter(upload=upload, entity_name__icontains="Instructions").exists()
        )
        self.assertFalse(
            KnowledgeEntity.objects.filter(upload=upload, entity_name__icontains="Vendor 2").exists()
        )
        self.assertFalse(
            KnowledgeEntity.objects.filter(upload=upload, entity_name__icontains="Select Yes or No").exists()
        )

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_xlsx_sheet_roles_shape_representation(self, _build_embeddings):
        if Workbook is None:
            self.skipTest("openpyxl not installed")
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Spreadsheet Sheet Roles XLSX",
        )
        storage_path = Path("uploads/sheet-roles.xlsx")
        target_path = Path(self._media_root) / storage_path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()

        instructions = workbook.active
        instructions.title = "Instructions"
        instructions.append(["Template Version", "09/23/25"])
        instructions.append(["Workbook Instructions", "Fill all required fields"])

        cover = workbook.create_sheet("Payment Cover")
        cover.append(["Field", "Value", "Notes"])
        cover.append(["Organization Name", "Acme Labs", ""])
        cover.append(["Agreement Number", "EPC-19-000", ""])
        cover.append(["Billing Contact", "ops@acme.test", "Primary contact"])

        summary = workbook.create_sheet("Invoice Summary")
        summary.append(["Category", "Amount", "Balance"])
        summary.append(["Direct Labor", 0, 0])
        summary.append(["Grand Totals", 1000, 200])

        equipment = workbook.create_sheet("Equipment")
        equipment.append(["Reference", "Description", "Amount"])
        equipment.append(["Worksheet Specific Instructions", "Document your purchase", ""])
        equipment.append(["E-1", "Heat Pump", 5500])

        lists = workbook.create_sheet("Lists")
        lists.sheet_state = "hidden"
        lists.append(["Choices"])
        lists.append(["Select Yes or No"])

        workbook.save(target_path)
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="sheet-roles.xlsx",
            storage_path=str(storage_path),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=target_path.stat().st_size,
        )

        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        extraction = service._extract_upload(upload)
        service._persist_extraction(upload, extraction)
        upload.refresh_from_db()

        tables = {table.title: table for table in KnowledgeUploadTable.objects.filter(upload=upload)}
        self.assertEqual(tables["Instructions"].metadata.get("sheet_role"), "instructional")
        self.assertEqual(tables["Payment Cover"].metadata.get("sheet_role"), "form_like")
        self.assertEqual(tables["Invoice Summary"].metadata.get("sheet_role"), "summary")
        self.assertEqual(tables["Equipment"].metadata.get("sheet_role"), "transactional")
        if "Lists" in tables:
            self.assertEqual(tables["Lists"].metadata.get("sheet_role"), "reference_hidden")
            self.assertTrue(tables["Lists"].metadata.get("is_decorative"))

        entities = list(KnowledgeEntity.objects.filter(upload=upload).values_list("entity_name", flat=True))
        self.assertIn("Grand Totals", entities)
        self.assertIn("E-1", entities)
        self.assertNotIn("Workbook Instructions", entities)
        self.assertNotIn("Direct Labor", entities)
        self.assertNotIn("Worksheet Specific Instructions", entities)
        self.assertFalse(
            KnowledgeEntity.objects.filter(upload=upload, entity_name__icontains="Select Yes or No").exists()
        )

        chunks = list(KnowledgeUploadChunk.objects.filter(upload=upload))
        compact_chunks = [
            chunk for chunk in chunks if isinstance(chunk.metadata, dict) and chunk.metadata.get("table_entity_compact")
        ]
        self.assertTrue(compact_chunks)
        self.assertTrue(
            any(chunk.metadata.get("sheet_name") == "Payment Cover" for chunk in compact_chunks)
        )
        self.assertTrue(
            any(chunk.metadata.get("sheet_role") == "summary" for chunk in compact_chunks)
        )
        self.assertTrue(
            any(
                isinstance(chunk.metadata, dict)
                and chunk.metadata.get("sheet_name") == "Equipment"
                and not chunk.metadata.get("table_entity_compact")
                and chunk.metadata.get("entity_name") == "E-1"
                for chunk in chunks
            )
        )

        metadata = upload.ingestion_metadata or {}
        role_counts = metadata.get("spreadsheet_sheet_roles") or {}
        self.assertEqual(role_counts.get("instructional"), 1)
        self.assertEqual(role_counts.get("form_like"), 1)
        self.assertEqual(role_counts.get("summary"), 1)
        self.assertEqual(role_counts.get("transactional"), 1)
        if "Lists" in tables:
            self.assertEqual(role_counts.get("reference_hidden"), 1)

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_xlsx_entity_promotion_skips_blank_transactional_slots_with_repeated_boilerplate(self, _build_embeddings):
        if Workbook is None:
            self.skipTest("openpyxl not installed")
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Spreadsheet Transactional Payload XLSX",
        )
        storage_path = Path("uploads/transactional-payload.xlsx")
        target_path = Path(self._media_root) / storage_path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()

        sheet = workbook.active
        sheet.title = "Transactions"
        sheet.append(["Reference", "Agreement", "Organization", "Amount", "Approval", "Share"])
        sheet.append(["R-1", "Agreement 2026", "Org A", 1250, "Yes", 100])
        sheet.append(["Row Slot 2", "Agreement 2026", "Org A", 0, "Select Yes or No", 0])
        sheet.append(["Row Slot 3", "Agreement 2026", "Org A", 0, "Select Yes or No", 0])
        sheet.append(["Row Slot 4", "Agreement 2026", "Org A", 0, "Select Yes or No", 0])

        workbook.save(target_path)
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="transactional-payload.xlsx",
            storage_path=str(storage_path),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=target_path.stat().st_size,
        )

        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        extraction = service._extract_upload(upload)
        service._persist_extraction(upload, extraction)

        entities = list(KnowledgeEntity.objects.filter(upload=upload).values_list("entity_name", flat=True))
        self.assertEqual(entities, ["R-1"])
        self.assertFalse(
            KnowledgeEntity.objects.filter(upload=upload, entity_name__icontains="Row Slot 2").exists()
        )
        self.assertFalse(
            KnowledgeEntity.objects.filter(upload=upload, entity_name__icontains="Select Yes or No").exists()
        )

    @mock.patch("apps.knowledge.knowledge_ingestion.build_embedding_service", return_value=None)
    def test_xlsx_repeated_multifield_rows_without_ids_stay_transactional(self, _build_embeddings):
        if Workbook is None:
            self.skipTest("openpyxl not installed")
        upload = KnowledgeUpload.objects.create(
            business_profile=self.business,
            user=self.user,
            source_type=KnowledgeSourceType.FILE,
            status=KnowledgeStatus.PENDING,
            display_name="Applied Jobs XLSX",
        )
        storage_path = Path("uploads/applied-jobs.xlsx")
        target_path = Path(self._media_root) / storage_path
        target_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()

        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet.append(
            [
                "Job Title",
                "Company Name",
                "Date Applied",
                "Application Deadline",
                "Location",
                "Application Status",
                "Notes",
            ]
        )
        sheet.append(
            [
                "AVP Relationship Manager",
                "Michael Page",
                "2024-04-29",
                "2 Weeks",
                "UAE",
                "Under Review",
                "Summary for Phone Interview",
            ]
        )
        sheet.append(
            [
                "Customer Success Team Leader",
                "Foodics",
                "2024-04-29",
                "",
                "KSA",
                "Under Review",
                "Phone Interview",
            ]
        )
        sheet.append(
            [
                "Customer Service Manager",
                "Nine Star Properties",
                "2024-04-29",
                "",
                "UAE",
                "Applied",
                "Strong fit",
            ]
        )

        workbook.save(target_path)
        KnowledgeUploadFile.objects.create(
            upload=upload,
            filename="applied-jobs.xlsx",
            storage_path=str(storage_path),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=target_path.stat().st_size,
        )

        service = KnowledgeIngestionService(media_root=Path(self._media_root))
        extraction = service._extract_upload(upload)
        service._persist_extraction(upload, extraction)
        upload.refresh_from_db()

        table = KnowledgeUploadTable.objects.get(upload=upload, title="Sheet1")
        self.assertEqual(table.metadata.get("sheet_role"), "transactional")

        entities = list(KnowledgeEntity.objects.filter(upload=upload).values_list("entity_name", flat=True))
        self.assertIn("AVP Relationship Manager", entities)
        self.assertIn("Customer Success Team Leader", entities)
        self.assertIn("Customer Service Manager", entities)

        compact_chunks = list(
            KnowledgeUploadChunk.objects.filter(
                upload=upload,
                metadata__table_entity_compact=True,
            )
        )
        self.assertFalse(compact_chunks)


class DeriveTableTitleTests(SimpleTestCase):
    """Tests for KnowledgeIngestionService._derive_table_title."""

    @staticmethod
    def _make_table(
        title="",
        section_heading="",
        order_index=1,
        column_schema=None,
    ):
        from apps.knowledge.knowledge_ingestion import TablePayload

        return TablePayload(
            order_index=order_index,
            title=title,
            section_heading=section_heading,
            page_number=None,
            column_schema=column_schema or [],
        )

    @staticmethod
    def _make_upload(display_name="", source_name=""):
        from types import SimpleNamespace

        return SimpleNamespace(display_name=display_name, source_name=source_name)

    def test_non_generic_title_preserved(self):
        """Azure DI caption or real section heading is kept as-is."""
        table = self._make_table(title="Revenue by Region")
        upload = self._make_upload(display_name="report.pdf")
        result = KnowledgeIngestionService._derive_table_title(table, upload)
        self.assertEqual(result, "Revenue by Region")

    def test_generic_title_with_doc_name_and_section(self):
        table = self._make_table(title="Table 3", section_heading="Q4 Results")
        upload = self._make_upload(display_name="annual_report.pdf")
        result = KnowledgeIngestionService._derive_table_title(table, upload)
        self.assertEqual(result, "annual_report \u2013 Q4 Results")

    def test_generic_title_with_doc_name_only(self):
        table = self._make_table(title="Table 2", order_index=2)
        upload = self._make_upload(display_name="revenue.csv")
        result = KnowledgeIngestionService._derive_table_title(table, upload)
        self.assertEqual(result, "revenue \u2013 Table 2")

    def test_generic_title_with_section_only(self):
        table = self._make_table(title="Table 1", section_heading="Appendix A")
        upload = self._make_upload()
        result = KnowledgeIngestionService._derive_table_title(table, upload)
        self.assertEqual(result, "Appendix A")

    def test_generic_title_with_column_schema(self):
        table = self._make_table(
            title="Table 5",
            order_index=5,
            column_schema=["Name", "Price", "Quantity", "Total", "Tax"],
        )
        upload = self._make_upload()
        result = KnowledgeIngestionService._derive_table_title(table, upload)
        self.assertEqual(result, "Table 5 (Name, Price, Quantity, Total, ...)")

    def test_final_fallback(self):
        table = self._make_table(title="Table 9", order_index=9)
        upload = self._make_upload()
        result = KnowledgeIngestionService._derive_table_title(table, upload)
        self.assertEqual(result, "Table 9")

    def test_case_insensitive_generic_detection(self):
        table = self._make_table(title="table 7", order_index=7, section_heading="Sales")
        upload = self._make_upload()
        result = KnowledgeIngestionService._derive_table_title(table, upload)
        self.assertEqual(result, "Sales")

    def test_source_name_fallback(self):
        table = self._make_table(title="Table 1", order_index=1)
        upload = self._make_upload(source_name="budget_2024.xlsx")
        result = KnowledgeIngestionService._derive_table_title(table, upload)
        self.assertEqual(result, "budget_2024 \u2013 Table 1")
