from __future__ import annotations

import csv
import gzip
import io
import json
import logging
import math
import re
import shutil
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from django.conf import settings

from apps.accounts.models import KnowledgeBlockType, KnowledgeIssueSeverity
from apps.knowledge.dataset_key_index import (
    BloomFilter,
    bloom_spec_for_items,
    normalize_identifier_value,
    resolve_key_index_storage_path,
    write_bloom_filter,
)
from apps.knowledge.ingestion_aliases import ALIAS_KEYWORDS
from apps.knowledge.ingestion_contracts import (
    ExtractionResult,
    IssuePayload,
    KnowledgeIngestionError,
    PageBlockPayload,
    PageLayout,
    TableCellPayload,
    TablePayload,
    TableRowPayload,
)
from apps.knowledge.ingestion_signals import (
    _log_normalization_summary,
    _SPREADSHEET_INSTRUCTION_SHEET_RE,
    _SPREADSHEET_INSTRUCTION_TOKEN_RE,
    _SPREADSHEET_RECORD_ID_RE,
    _SPREADSHEET_REFERENCE_SHEET_RE,
    _SPREADSHEET_SUMMARY_ROW_RE,
    _SPREADSHEET_ZERO_LIKE_RE,
    _TABLE_NUMBER_LIKE_RE,
)
from apps.knowledge.models import KnowledgeUpload, KnowledgeUploadFile
from apps.knowledge.table_normalization import (
    NormalizedSheet,
    SheetNormalizationDiagnostics,
    SpreadsheetRowInput,
    normalize_sheet_rows,
    resolve_normalization_policy,
    sheet_is_allowed,
    summarize_normalization,
)

logger = logging.getLogger(__name__)


try:  # pragma: no cover - dependency failure should be surfaced at runtime
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - fallback handled via runtime check
    load_workbook = None  # type: ignore

try:  # pragma: no cover - dependency failure should be surfaced at runtime
    import xlrd
except ImportError:  # pragma: no cover - fallback handled via runtime check
    xlrd = None  # type: ignore


class IngestionTabularFilesMixin:


    def _should_use_dataset_mode(self, *, estimated_rows: int | None) -> bool:
        if not self.dataset_mode_enabled:
            return False
        if estimated_rows is None:
            return False
        return int(estimated_rows) >= int(self.dataset_row_threshold)

    def _dataset_storage_directory(self, upload: KnowledgeUpload) -> tuple[Path, Path]:
        rel_dir = Path("datasets") / str(upload.business_profile_id) / str(upload.id)
        abs_dir = (self.media_root / rel_dir).resolve()
        abs_dir.relative_to(self.media_root)
        return abs_dir, rel_dir

    def _reset_dataset_storage(self, upload: KnowledgeUpload) -> tuple[Path, Path]:
        abs_dir, rel_dir = self._dataset_storage_directory(upload)
        if abs_dir.exists():
            try:
                shutil.rmtree(abs_dir)
            except OSError as exc:
                logger.warning("dataset.cleanup_failed upload=%s error=%s", upload.id, exc)
        abs_dir.mkdir(parents=True, exist_ok=True)
        return abs_dir, rel_dir

    @staticmethod
    def _stringify_dataset_cell(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, bool):
            return "true" if value else "false"
        if hasattr(value, "isoformat"):
            try:
                return value.isoformat()  # datetime/date-like
            except Exception:
                pass
        text = str(value)
        if "\x00" in text:
            text = text.replace("\x00", " ")
        return text

    @staticmethod
    def _estimate_delimited_row_count(path: Path, *, sample_bytes: int = 250_000) -> int | None:
        try:
            file_size = path.stat().st_size
        except OSError:
            return None
        if file_size <= 0:
            return 0
        try:
            with path.open("rb") as handle:
                sample = handle.read(sample_bytes)
        except OSError:
            return None
        lines = sample.count(b"\n")
        if lines <= 1:
            return 1 if file_size else 0
        avg = len(sample) / float(lines)
        if avg <= 0:
            return None
        return max(0, int(file_size / avg))

    def _suggest_dataset_key_columns(
        self,
        *,
        column_schema: Sequence[str],
        sample_rows: Sequence[Sequence[str]],
        max_candidates: int = 8,
    ) -> list[dict[str, Any]]:
        if not column_schema or not sample_rows:
            return []
        candidates: list[dict[str, Any]] = []
        sample_count = len(sample_rows)
        min_samples = min(10, sample_count)

        for idx, column in enumerate(column_schema):
            name = str(column or "").strip()
            canonical = self._canonical_column_name(name)
            if not canonical:
                continue
            values = [
                str(row[idx]).strip()
                for row in sample_rows
                if idx < len(row) and str(row[idx]).strip()
            ]
            non_empty = len(values)
            if non_empty < max(3, min_samples):
                continue
            unique = len(set(values))
            unique_ratio = unique / float(non_empty) if non_empty else 0.0
            keyword_match = any(key in canonical for key in ALIAS_KEYWORDS) or canonical.endswith("_id")
            score = (1.0 if keyword_match else 0.0) + unique_ratio
            candidates.append(
                {
                    "column": name,
                    "normalized": canonical,
                    "non_empty": non_empty,
                    "unique": unique,
                    "unique_ratio": round(unique_ratio, 4),
                    "keyword_match": bool(keyword_match),
                    "score": round(score, 4),
                }
            )

        candidates.sort(key=lambda item: (item.get("score", 0), item.get("non_empty", 0)), reverse=True)
        return candidates[: max_candidates]

    def _dataset_key_index_columns(
        self,
        *,
        upload: KnowledgeUpload,
        sheet_name: str | None,
        column_schema: Sequence[str],
        suggested_keys: Sequence[Mapping[str, Any]] | None,
        rules: Mapping[str, Any] | None,
    ) -> list[dict[str, Any]]:
        """
        Decide which columns should receive a dataset key index (Bloom filter).

        Preference order:
        1) Suggested key columns from sampling.
        2) Fallback heuristics on column names.
        """

        max_columns = int(getattr(settings, "DATASET_KEY_INDEX_MAX_COLUMNS", 4) or 4)
        max_columns = max(0, min(20, max_columns))
        if max_columns <= 0:
            return []

        allow_sensitive = str(getattr(settings, "DATASET_KEY_INDEX_ALLOW_SENSITIVE", "false")).lower() in {
            "1",
            "true",
            "yes",
        }
        min_suggested_score = float(getattr(settings, "DATASET_KEY_INDEX_SUGGESTED_MIN_SCORE", 0.9) or 0.9)
        min_suggested_score = max(0.0, min(10.0, min_suggested_score))

        schema_canon: list[str] = [self._canonical_column_name(col) for col in column_schema]
        index_map = {canon: idx for idx, canon in enumerate(schema_canon) if canon}
        name_map = {canon: col for canon, col in zip(schema_canon, column_schema) if canon and col}

        def _eligible(column_name: str) -> bool:
            if not column_name:
                return False
            if allow_sensitive or not rules:
                return True
            return not self._column_is_sensitive(column_name, rules)

        chosen: list[dict[str, Any]] = []
        seen: set[str] = set()

        def _add(*, column: str, source: str, **extra: Any) -> None:
            canon = self._canonical_column_name(column)
            if not canon or canon in seen:
                return
            idx = index_map.get(canon)
            if idx is None:
                return
            actual = name_map.get(canon) or column
            if not _eligible(actual):
                return
            seen.add(canon)
            chosen.append(
                {
                    "column": actual,
                    "column_index": idx,
                    "source": source,
                    **extra,
                }
            )

        if suggested_keys:
            for entry in suggested_keys:
                col = str(entry.get("column") or "").strip()
                if not col:
                    continue
                try:
                    score = float(entry.get("score") or 0.0)
                except (TypeError, ValueError):
                    score = 0.0
                if score < min_suggested_score:
                    continue
                _add(column=col, source="suggested", suggested_score=score, suggested_keyword_match=bool(entry.get("keyword_match")))
                if len(chosen) >= max_columns:
                    break

        if len(chosen) >= max_columns:
            return chosen[:max_columns]

        heuristic_terms = ("invoice", "order", "ticket", "serial", "reference", "ref", "email", "phone", "mobile", "code", "sku")
        for col in column_schema:
            canon = self._canonical_column_name(col)
            if not canon or canon in seen:
                continue
            if not any(term in canon for term in heuristic_terms):
                continue
            _add(column=str(col or "").strip(), source="heuristic")
            if len(chosen) >= max_columns:
                break

        return chosen[:max_columns]

    def _build_dataset_card_segment_payload(
        self,
        *,
        upload: KnowledgeUpload,
        ingestion_metadata: Mapping[str, Any] | None = None,
    ) -> dict[str, Any] | None:
        return build_dataset_card_segment_payload(upload=upload, ingestion_metadata=ingestion_metadata)

    def _build_dataset_preview_table(
        self,
        *,
        order_index: int,
        sheet_name: str,
        sheet_index: int | None,
        column_schema: list[str],
        preview_rows: Sequence[tuple[int, Sequence[str]]],
        file_detail: KnowledgeUploadFile,
        source_label: str,
        content_type: str,
        rules: Mapping[str, Any] | None,
    ) -> tuple[TablePayload, PageLayout]:
        visible_mask = [
            not self._column_is_sensitive(column, rules)
            for column in column_schema
        ] if rules else [True] * len(column_schema)
        rows: list[TableRowPayload] = []
        for row_index, values in preview_rows:
            attributes = {column_schema[i]: (values[i] if i < len(values) else "") for i in range(len(column_schema))}
            if rules and self._row_is_internal(attributes, rules):
                continue
            cells: list[TableCellPayload] = []
            row_values: list[str] = []
            for col_idx, column_key in enumerate(column_schema):
                if col_idx >= len(values):
                    cell_value = ""
                else:
                    cell_value = str(values[col_idx] or "")
                if visible_mask[col_idx]:
                    cells.append(
                        TableCellPayload(
                            row_index=row_index,
                            column_index=col_idx,
                            column_key=column_key,
                            raw_text=cell_value,
                        )
                    )
                    row_values.append(cell_value)
            rows.append(
                TableRowPayload(
                    row_index=row_index,
                    page_number=sheet_index,
                    raw_text="\t".join(row_values),
                    metadata={"source": source_label, "sheet": sheet_name},
                    cells=cells,
                )
            )

        table = TablePayload(
            order_index=order_index,
            title=sheet_name,
            section_heading=sheet_name,
            page_number=sheet_index,
            column_schema=column_schema,
            metadata={
                "source": source_label,
                "sheet_name": sheet_name,
                "filename": file_detail.filename,
            },
            rows=rows,
        )
        preview = self._table_preview_text([table], rules=rules)
        page_layout = PageLayout(
            page_number=sheet_index or order_index,
            width=612,
            height=792,
            rotation=0,
            text_density=len(preview.strip()) / float(612 * 792),
            has_ocr_content=False,
            content_type=content_type,
            blocks=[
                PageBlockPayload(
                    block_type=KnowledgeBlockType.TABLE,
                    order_index=order_index,
                    text=preview,
                    metadata={"source": source_label, "sheet": sheet_name},
                )
            ],
            metadata={"sheet_name": sheet_name, "dataset_mode": True},
        )
        return table, page_layout

    def _extract_csv(
        self,
        path: Path,
        *,
        format_hint: str,
        file_detail: KnowledgeUploadFile,
        upload: KnowledgeUpload | None = None,
    ) -> ExtractionResult:
        estimated_rows = self._estimate_delimited_row_count(path)
        if upload and self._should_use_dataset_mode(estimated_rows=estimated_rows):
            return self._extract_delimited_dataset(
                path,
                format_hint=format_hint,
                file_detail=file_detail,
                upload=upload,
                estimated_rows=estimated_rows,
            )

        raw_text = self._extract_text_file(path)
        normalized = raw_text.lstrip("\ufeff")
        if not normalized.strip():
            raise KnowledgeIngestionError("CSV document did not contain any rows.")
        delimiter = "\t" if format_hint == "tsv" else ","
        sample = normalized[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=[",", ";", "\t", "|"])
            delimiter = dialect.delimiter or delimiter
        except Exception:
            # Keep default delimiter
            pass
        reader = csv.reader(io.StringIO(normalized), delimiter=delimiter)
        parsed_rows = [list(row) for row in reader]
        if not parsed_rows:
            raise KnowledgeIngestionError("CSV document did not contain any usable rows.")

        policy = resolve_normalization_policy(upload)
        sheet_label = (file_detail.filename or "CSV").strip() or "CSV"
        normalized_sheet = normalize_sheet_rows(parsed_rows, sheet_name=sheet_label, policy=policy)
        diagnostics = [normalized_sheet.diagnostics]
        if normalized_sheet.diagnostics.skipped:
            raise KnowledgeIngestionError("CSV document did not contain any usable rows.")

        column_schema = normalized_sheet.column_schema
        table_rows: list[TableRowPayload] = []
        for row_idx, values in enumerate(normalized_sheet.rows, start=1):
            cells: list[TableCellPayload] = []
            formatted_cells: list[str] = []
            for col_idx, column_key in enumerate(column_schema):
                value = values[col_idx] if col_idx < len(values) else ""
                formatted_cells.append(value)
                cells.append(
                    TableCellPayload(
                        row_index=row_idx,
                        column_index=col_idx,
                        column_key=column_key,
                        raw_text=value,
                    )
                )
            table_rows.append(
                TableRowPayload(
                    row_index=row_idx,
                    page_number=None,
                    raw_text="\t".join(formatted_cells),
                    metadata={"source": format_hint or "csv", "line_number": row_idx + 1},
                    cells=cells,
                )
            )
        table = TablePayload(
            order_index=1,
            title=file_detail.filename or "CSV Table",
            section_heading="",
            page_number=None,
            column_schema=column_schema,
            metadata={
                "source": "csv",
                "delimiter": delimiter,
                "filename": file_detail.filename,
            },
            rows=table_rows,
        )
        rules = self._table_privacy_rules(upload)
        ingest_config = self._table_ingest_config(upload)
        tables, table_metrics, limit_issues, table_summary = self._apply_table_limits([table], upload=upload, config=ingest_config)
        if not tables:
            raise KnowledgeIngestionError("CSV document did not contain rows within configured limits.")
        table_entities = self._table_row_entities(
            tables,
            business_profile=getattr(upload, "business_profile", None),
            upload=upload,
        )
        preview_text = self._table_preview_text(tables, rules=rules)
        table_stats = self._table_stats_summary(
            total_rows=table_summary["total_rows"],
            indexed_rows=table_summary["indexed_rows"],
            row_cap=table_summary.get("row_cap_hint"),
            source_row_count=self._integration_row_count(upload),
            table_count=len(tables),
            partial_tables=table_summary.get("partial_tables", 0),
            row_tier=table_summary.get("row_tier_hint"),
        )
        normalization_summary = summarize_normalization(policy, diagnostics)
        issues = limit_issues
        page = PageLayout(
            page_number=1,
            width=612,
            height=792,
            rotation=0,
            text_density=len(preview_text.strip()) / float(612 * 792),
            has_ocr_content=False,
            content_type="text/csv" if delimiter == "," else "text/tab-separated-values",
            blocks=[
                PageBlockPayload(
                    block_type=KnowledgeBlockType.TABLE,
                    order_index=0,
                    text=preview_text or normalized[:2000],
                    metadata={"source": "csv"},
                )
            ],
            metadata={"table_count": 1, "filename": file_detail.filename},
        )
        metadata = {
            "format": format_hint,
            "filename": file_detail.filename,
            "content_type": file_detail.content_type or "",
            "storage_path": file_detail.storage_path,
            "table_count": len(tables),
            "table_truncation": table_metrics,
            "table_stats": table_stats,
        }
        if normalization_summary:
            metadata["normalization"] = normalization_summary
        _log_normalization_summary(upload, "csv", normalization_summary)
        return ExtractionResult(
            text=preview_text or normalized,
            format_hint=format_hint,
            metadata=metadata,
            pages=[page],
            tables=tables,
            issues=issues,
            entities=table_entities,
        )

    def _extract_delimited_dataset(
        self,
        path: Path,
        *,
        format_hint: str,
        file_detail: KnowledgeUploadFile,
        upload: KnowledgeUpload,
        estimated_rows: int | None,
    ) -> ExtractionResult:
        if not upload:
            raise KnowledgeIngestionError("Dataset mode requires an upload record.")

        dataset_root, dataset_root_rel = self._reset_dataset_storage(upload)
        timestamp = timezone.now().strftime("%Y%m%d%H%M%S")
        display_name = (file_detail.filename or "Dataset").strip() or "Dataset"
        safe_label = slugify(Path(display_name).stem) or "dataset"
        dataset_filename = f"{timestamp}_{safe_label}.csv.gz"
        dataset_path = dataset_root / dataset_filename
        dataset_rel_path = (dataset_root_rel / dataset_filename).as_posix()

        rules = self._table_privacy_rules(upload)
        preview_row_cap = int(self.dataset_preview_rows)
        sample_row_cap = int(self.dataset_sample_rows)

        try:
            with path.open("rb") as handle:
                sample_bytes = handle.read(64 * 1024)
        except OSError as exc:
            raise KnowledgeIngestionError(f"Unable to read dataset sample: {exc}") from exc

        sample_text = sample_bytes.decode("utf-8", errors="ignore")
        delimiter = "\t" if format_hint == "tsv" else ","
        try:
            dialect = csv.Sniffer().sniff(sample_text[:4096], delimiters=[",", ";", "\t", "|"])
            delimiter = getattr(dialect, "delimiter", delimiter) or delimiter
        except Exception:
            dialect = csv.excel

        preview_rows: list[tuple[int, list[str]]] = []
        sample_rows: list[list[str]] = []
        sample_visible_rows: list[dict[str, str]] = []
        row_count = 0
        internal_skipped = 0
        sample_target = max(sample_row_cap, 25)

        key_index_enabled = str(getattr(settings, "DATASET_KEY_INDEX_ENABLED", "true")).lower() in {"1", "true", "yes"}
        key_index_max_bytes = int(getattr(settings, "DATASET_KEY_INDEX_MAX_BYTES", 2_000_000) or 2_000_000)
        key_index_max_bytes = max(4096, min(25_000_000, key_index_max_bytes))
        key_indexes: list[dict[str, Any]] = []

        with path.open("rb") as raw_in:
            reader_stream = io.TextIOWrapper(raw_in, encoding="utf-8", errors="ignore", newline="")
            reader = csv.reader(reader_stream, delimiter=delimiter)
            try:
                raw_header = next(reader)
            except StopIteration:
                raise KnowledgeIngestionError("Dataset did not contain a header row.")

            if raw_header and isinstance(raw_header[0], str):
                raw_header[0] = raw_header[0].lstrip("\ufeff")

            column_schema: list[str] = []
            for idx, value in enumerate(raw_header):
                label = str(value or "").strip()
                column_schema.append(label or f"column_{idx + 1}")
            if not column_schema:
                raise KnowledgeIngestionError("Dataset did not contain any columns.")

            visible_mask = [
                not self._column_is_sensitive(column, rules)
                for column in column_schema
            ] if rules else [True] * len(column_schema)

            dataset_path.parent.mkdir(parents=True, exist_ok=True)
            with gzip.open(dataset_path, "wt", encoding="utf-8", newline="") as out_handle:
                writer = csv.writer(out_handle)
                writer.writerow(column_schema)

                for row in reader:
                    values = [str(item or "") for item in row]
                    if len(values) < len(column_schema):
                        values.extend([""] * (len(column_schema) - len(values)))
                    elif len(values) > len(column_schema):
                        values = values[: len(column_schema)]

                    if not any(value.strip() for value in values):
                        continue

                    attributes = {column_schema[i]: values[i] for i in range(len(column_schema))}
                    if rules and self._row_is_internal(attributes, rules):
                        internal_skipped += 1
                        continue

                    row_count += 1
                    writer.writerow(values)

                    if len(preview_rows) < preview_row_cap:
                        preview_rows.append((row_count, values))
                    if len(sample_rows) < sample_target:
                        sample_rows.append(values)
                    if len(sample_visible_rows) < sample_row_cap:
                        row_preview = {
                            column_schema[i]: values[i]
                            for i in range(len(column_schema))
                            if visible_mask[i]
                        }
                        sample_visible_rows.append(row_preview)

                    if key_index_enabled and not key_indexes and len(sample_rows) >= sample_target:
                        suggested_preview = self._suggest_dataset_key_columns(column_schema=column_schema, sample_rows=sample_rows)
                        chosen_columns = self._dataset_key_index_columns(
                            upload=upload,
                            sheet_name=None,
                            column_schema=column_schema,
                            suggested_keys=suggested_preview,
                            rules=rules,
                        )
                        expected_items = int(estimated_rows or row_count or 1)
                        spec = bloom_spec_for_items(expected_items)
                        for entry in chosen_columns:
                            bits = spec.bits
                            byte_len = (bits + 7) // 8
                            if byte_len > key_index_max_bytes:
                                bits = key_index_max_bytes * 8
                            bloom = BloomFilter(bits=bits, hashes=spec.hashes)
                            key_indexes.append(
                                {
                                    "column": entry.get("column"),
                                    "column_index": int(entry.get("column_index")),
                                    "source": entry.get("source"),
                                    "bloom": bloom,
                                    "value_count": 0,
                                }
                            )
                        if key_indexes:
                            for sample in sample_rows:
                                for info in key_indexes:
                                    idx = int(info["column_index"])
                                    if idx >= len(sample):
                                        continue
                                    normalized = normalize_identifier_value(sample[idx])
                                    if not normalized:
                                        continue
                                    info["bloom"].add(normalized)
                                    info["value_count"] += 1
                    elif key_indexes:
                        for info in key_indexes:
                            idx = int(info["column_index"])
                            if idx >= len(values):
                                continue
                            normalized = normalize_identifier_value(values[idx])
                            if not normalized:
                                continue
                            info["bloom"].add(normalized)
                            info["value_count"] += 1

        try:
            dataset_size = dataset_path.stat().st_size
        except OSError:
            dataset_size = 0

        suggested_keys = self._suggest_dataset_key_columns(column_schema=column_schema, sample_rows=sample_rows)
        key_index_payloads: list[dict[str, Any]] = []
        if key_indexes:
            for info in key_indexes:
                column_name = str(info.get("column") or "").strip()
                if not column_name:
                    continue
                index_rel_path = resolve_key_index_storage_path(dataset_rel_path=dataset_rel_path, column_name=column_name)
                abs_index_path = (self.media_root / Path(index_rel_path)).resolve()
                try:
                    abs_index_path.relative_to(self.media_root)
                except ValueError:
                    continue
                bloom = info.get("bloom")
                if not isinstance(bloom, BloomFilter):
                    continue
                try:
                    written_bytes = write_bloom_filter(abs_index_path, bloom)
                except OSError:
                    continue
                key_index_payloads.append(
                    {
                        "column": column_name,
                        "storage_path": index_rel_path,
                        "bits": int(bloom.bits),
                        "hashes": int(bloom.hashes),
                        "bytes": int(written_bytes),
                        "values_indexed": int(info.get("value_count") or 0),
                        "source": info.get("source"),
                    }
                )

        table, page = self._build_dataset_preview_table(
            order_index=1,
            sheet_name=display_name,
            sheet_index=None,
            column_schema=column_schema,
            preview_rows=preview_rows,
            file_detail=file_detail,
            source_label="dataset_csv",
            content_type="text/csv",
            rules=rules,
        )
        tables = [table]
        preview_text = self._table_preview_text(tables, rules=rules)

        issues: list[IssuePayload] = []
        issues.append(
            IssuePayload(
                code="dataset_mode_enabled",
                severity=KnowledgeIssueSeverity.INFO.value,
                description="Large dataset stored in dataset mode (file-backed). Only a small preview is indexed into Postgres.",
                details={
                    "estimated_rows": estimated_rows,
                    "rows_stored": row_count,
                    "preview_rows_indexed": len(table.rows),
                    "internal_rows_skipped": internal_skipped,
                    "dataset_storage_format": self.dataset_storage_format,
                },
            )
        )

        dataset_metadata = {
            "enabled": True,
            "storage_format": self.dataset_storage_format,
            "storage_path": dataset_rel_path,
            "size_bytes": dataset_size,
            "row_count": row_count,
            "estimated_row_count": estimated_rows,
            "column_schema": column_schema,
            "sample_rows": sample_visible_rows,
            "suggested_key_columns": suggested_keys,
            "delimiter": delimiter,
            "key_indexes": key_index_payloads,
        }

        table_stats = self._table_stats_summary(
            total_rows=row_count,
            indexed_rows=len(table.rows),
            row_cap=len(table.rows),
            source_row_count=self._integration_row_count(upload),
            table_count=1,
            partial_tables=1 if row_count > len(table.rows) else 0,
            row_tier="large" if estimated_rows and estimated_rows >= self.dataset_row_threshold else "medium",
        )
        metadata = {
            "format": format_hint,
            "filename": file_detail.filename,
            "content_type": file_detail.content_type or "",
            "storage_path": file_detail.storage_path,
            "table_count": 1,
            "table_truncation": {
                "truncated_tables": 0,
                "truncated_rows": max(0, row_count - len(table.rows)),
                "truncated_columns": 0,
            },
            "table_stats": table_stats,
            "dataset": dataset_metadata,
        }
        return ExtractionResult(
            text=preview_text,
            format_hint=format_hint,
            metadata=metadata,
            pages=[page],
            tables=tables,
            issues=issues,
            entities=self._table_row_entities(
                tables,
                business_profile=getattr(upload, "business_profile", None),
                upload=upload,
            ),
        )

    def _build_table_from_normalized_sheet(
        self,
        normalized: NormalizedSheet,
        *,
        order_index: int,
        sheet_name: str,
        sheet_index: int | None,
        sheet_hidden: bool,
        sheet_role: str,
        file_detail: KnowledgeUploadFile,
        source_label: str,
        content_type: str,
        rules: Mapping[str, Any] | None,
    ) -> tuple[TablePayload, PageLayout]:
        column_schema = normalized.column_schema
        table_rows: list[TableRowPayload] = []
        for row_idx, values in enumerate(normalized.rows, start=1):
            row_meta = (
                normalized.row_metadata[row_idx - 1]
                if row_idx - 1 < len(normalized.row_metadata)
                else None
            )
            cells: list[TableCellPayload] = []
            formatted: list[str] = []
            for col_idx, column_key in enumerate(column_schema):
                value = values[col_idx] if col_idx < len(values) else ""
                formatted.append(value)
                cells.append(
                    TableCellPayload(
                        row_index=row_idx,
                        column_index=col_idx,
                        column_key=column_key,
                        raw_text=value,
                    )
                )
            table_rows.append(
                TableRowPayload(
                    row_index=row_idx,
                    page_number=sheet_index,
                    raw_text="\t".join(formatted),
                    metadata={
                        "source": source_label,
                        "sheet": sheet_name,
                        "source_row_index": getattr(row_meta, "source_row_index", None),
                        "hidden": bool(getattr(row_meta, "hidden", False)),
                        "row_kind": getattr(row_meta, "row_kind", "data"),
                    },
                    cells=cells,
                )
            )
        table = TablePayload(
            order_index=order_index,
            title=sheet_name,
            section_heading=sheet_name,
            page_number=sheet_index,
            column_schema=column_schema,
            metadata={
                "source": source_label,
                "sheet_name": sheet_name,
                "sheet_hidden": bool(sheet_hidden),
                "sheet_role": sheet_role,
                "is_decorative": sheet_role == "reference_hidden",
                "filename": file_detail.filename,
            },
            rows=table_rows,
        )
        preview = self._table_preview_text([table], rules=rules)
        page_layout = PageLayout(
            page_number=sheet_index or order_index,
            width=612,
            height=792,
            rotation=0,
            text_density=len(preview.strip()) / float(612 * 792),
            has_ocr_content=False,
            content_type=content_type,
            blocks=[
                PageBlockPayload(
                    block_type=KnowledgeBlockType.TABLE,
                    order_index=order_index,
                    text=preview,
                    metadata={
                        "source": source_label,
                        "sheet": sheet_name,
                        "sheet_role": sheet_role,
                        "is_decorative": sheet_role == "reference_hidden",
                    },
                )
            ],
            metadata={"sheet_name": sheet_name, "sheet_role": sheet_role},
        )
        return table, page_layout

    def _extract_xlsx(
        self,
        path: Path,
        *,
        file_detail: KnowledgeUploadFile,
        upload: KnowledgeUpload | None = None,
    ) -> ExtractionResult:
        if load_workbook is None:
            raise KnowledgeIngestionError("XLSX ingestion requires the openpyxl package.")
        try:
            # Spreadsheet-template normalization needs row visibility metadata
            # (e.g. hidden filler rows), which openpyxl does not expose on
            # ReadOnlyWorksheet. Load the workbook normally on the non-dataset
            # XLSX path so normalization can make deterministic keep/drop
            # decisions before indexing.
            workbook = load_workbook(filename=path, read_only=False, data_only=True)
        except Exception as exc:
            raise KnowledgeIngestionError(f"Unable to open XLSX file: {exc}") from exc

        policy = resolve_normalization_policy(upload)
        rules = self._table_privacy_rules(upload)

        allowed_sheets: list[tuple[int, str, Any]] = []
        estimated_rows = 0
        policy_skipped: list[str] = []
        for sheet_idx, sheet in enumerate(workbook.worksheets, start=1):
            sheet_name = sheet.title or f"Sheet {sheet_idx}"
            if not sheet_is_allowed(sheet_name, policy):
                policy_skipped.append(sheet_name)
                continue
            allowed_sheets.append((sheet_idx, sheet_name, sheet))
            try:
                estimated_rows += int(getattr(sheet, "max_row", 0) or 0)
            except (TypeError, ValueError):
                continue

        if upload and self._should_use_dataset_mode(estimated_rows=estimated_rows):
            return self._extract_xlsx_dataset_mode(
                file_detail=file_detail,
                upload=upload,
                sheets=allowed_sheets,
                policy_skipped=policy_skipped,
                estimated_rows=estimated_rows,
            )

        diagnostics: list[SheetNormalizationDiagnostics] = [
            SheetNormalizationDiagnostics(sheet_name=name, skipped=True, skip_reason="policy")
            for name in policy_skipped
        ]
        tables: list[TablePayload] = []
        pages: list[PageLayout] = []
        sheet_role_counts: dict[str, int] = {}
        order_index = 1
        for sheet_idx, sheet_name, sheet in allowed_sheets:
            normalized = normalize_sheet_rows(
                self._iter_xlsx_rows_with_metadata(sheet),
                sheet_name=sheet_name,
                policy=policy,
            )
            diagnostics.append(normalized.diagnostics)
            if normalized.diagnostics.skipped:
                continue
            sheet_hidden = getattr(sheet, "sheet_state", "visible") != "visible"
            sheet_role = self._classify_spreadsheet_sheet_role(
                sheet_name=sheet_name,
                sheet_hidden=sheet_hidden,
                normalized=normalized,
            )
            sheet_role_counts[sheet_role] = sheet_role_counts.get(sheet_role, 0) + 1
            table, page_layout = self._build_table_from_normalized_sheet(
                normalized,
                order_index=order_index,
                sheet_name=sheet_name,
                sheet_index=sheet_idx,
                sheet_hidden=sheet_hidden,
                sheet_role=sheet_role,
                file_detail=file_detail,
                source_label="xlsx",
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                rules=rules,
            )
            tables.append(table)
            pages.append(page_layout)
            order_index += 1

        original_tables = list(tables)
        if not original_tables:
            raise KnowledgeIngestionError("XLSX workbook did not contain any populated sheets.")
        ingest_config = self._table_ingest_config(upload)
        tables, table_metrics, limit_issues, table_summary = self._apply_table_limits(original_tables, upload=upload, config=ingest_config)
        if not tables:
            raise KnowledgeIngestionError("XLSX workbook exceeded configured limits and no rows were indexed.")
        preview_text = self._table_preview_text(tables, rules=rules)
        table_entities = self._table_row_entities(
            tables,
            business_profile=getattr(upload, "business_profile", None),
            upload=upload,
        )
        table_stats = self._table_stats_summary(
            total_rows=table_summary["total_rows"],
            indexed_rows=table_summary["indexed_rows"],
            row_cap=table_summary.get("row_cap_hint"),
            source_row_count=self._integration_row_count(upload),
            table_count=len(tables),
            partial_tables=table_summary.get("partial_tables", 0),
            row_tier=table_summary.get("row_tier_hint"),
        )
        normalization_summary = summarize_normalization(policy, diagnostics)
        metadata = {
            "format": "xlsx",
            "filename": file_detail.filename,
            "content_type": file_detail.content_type or "",
            "storage_path": file_detail.storage_path,
            "table_count": len(tables),
            "table_truncation": table_metrics,
            "table_stats": table_stats,
            "spreadsheet_sheet_roles": sheet_role_counts,
        }
        if normalization_summary:
            metadata["normalization"] = normalization_summary
        _log_normalization_summary(upload, "xlsx", normalization_summary)
        return ExtractionResult(
            text=preview_text,
            format_hint="xlsx",
            metadata=metadata,
            pages=pages,
            tables=tables,
            issues=limit_issues,
            entities=table_entities,
        )

    @staticmethod
    def _iter_xlsx_rows_with_metadata(sheet: Any) -> Iterable[SpreadsheetRowInput]:
        for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            row_dimension = None
            try:
                row_dimension = sheet.row_dimensions.get(row_idx)
            except Exception:
                row_dimension = None
            hidden = bool(getattr(row_dimension, "hidden", False))
            yield SpreadsheetRowInput(values=row, row_index=row_idx, hidden=hidden)

    def _classify_spreadsheet_sheet_role(
        self,
        *,
        sheet_name: str,
        sheet_hidden: bool,
        normalized: NormalizedSheet,
    ) -> str:
        name = str(sheet_name or "").strip()
        if sheet_hidden and _SPREADSHEET_REFERENCE_SHEET_RE.search(name):
            return "reference_hidden"
        if _SPREADSHEET_INSTRUCTION_SHEET_RE.search(name):
            return "instructional"
        if _SPREADSHEET_SUMMARY_ROW_RE.search(name):
            return "summary"
        if sheet_hidden:
            return "reference_hidden"

        rows = normalized.rows or []
        if not rows:
            return "unknown"

        record_id_hits = 0
        transactional_record_rows = 0
        narrative_rows = 0
        summary_hits = 0
        form_like_rows = 0
        for row in rows[:100]:
            values = [str(value or "").strip() for value in row if str(value or "").strip()]
            if not values:
                continue
            descriptor = values[0]
            if _SPREADSHEET_RECORD_ID_RE.fullmatch(descriptor):
                record_id_hits += 1
            row_text = " ".join(values)
            if _SPREADSHEET_SUMMARY_ROW_RE.search(row_text):
                summary_hits += 1
            if len(values) <= 2 and sum(len(value.split()) for value in values) >= 6:
                narrative_rows += 1
            if self._spreadsheet_row_looks_transactional_record(values):
                transactional_record_rows += 1
            if self._spreadsheet_row_looks_form_like(values):
                form_like_rows += 1

        if record_id_hits > 0:
            return "transactional"
        if transactional_record_rows >= max(3, min(12, math.ceil(len(rows) * 0.3))) and transactional_record_rows > form_like_rows:
            return "transactional"
        if summary_hits > 0 and summary_hits >= max(narrative_rows, form_like_rows):
            return "summary"
        if narrative_rows >= max(3, len(rows) // 3):
            return "instructional"
        if form_like_rows >= max(2, min(8, math.ceil(len(rows) * 0.25))):
            return "form_like"
        return "unknown"

    def _spreadsheet_row_looks_guidance(self, values: Sequence[str]) -> bool:
        cleaned = [str(value or "").strip() for value in values if str(value or "").strip()]
        if not cleaned:
            return False

        joined = " ".join(cleaned)
        word_count = len(joined.split())
        numeric_like_count = sum(
            1
            for value in cleaned
            if _SPREADSHEET_RECORD_ID_RE.fullmatch(value) or _TABLE_NUMBER_LIKE_RE.search(value)
        )
        long_text_cells = sum(1 for value in cleaned if len(value.split()) >= 8)
        has_instruction_token = any(_SPREADSHEET_INSTRUCTION_TOKEN_RE.search(value) for value in cleaned)
        if has_instruction_token and numeric_like_count == 0:
            return True
        return word_count >= 18 and numeric_like_count <= 1 and long_text_cells >= 1 and len(cleaned) <= 3

    def _spreadsheet_row_looks_form_like(self, values: Sequence[str]) -> bool:
        cleaned = [str(value or "").strip() for value in values if str(value or "").strip()]
        if len(cleaned) < 2 or len(cleaned) > 6:
            return False
        if any(_SPREADSHEET_RECORD_ID_RE.fullmatch(value) for value in cleaned):
            return False
        if self._spreadsheet_row_looks_guidance(cleaned):
            return False

        text_like_count = 0
        numeric_like_count = 0
        control_like_count = 0
        total_words = 0
        for value in cleaned:
            total_words += len(value.split())
            if self._spreadsheet_value_is_control(value):
                control_like_count += 1
                continue
            if _TABLE_NUMBER_LIKE_RE.search(value):
                numeric_like_count += 1
                continue
            text_like_count += 1

        if control_like_count >= len(cleaned) - 1:
            return False
        if text_like_count >= 1 and numeric_like_count >= 1:
            return True
        if text_like_count >= 2 and total_words <= 24:
            return True
        return False

    def _spreadsheet_row_looks_transactional_record(self, values: Sequence[str]) -> bool:
        cleaned = [str(value or "").strip() for value in values if str(value or "").strip()]
        if len(cleaned) < 5:
            return False
        if any(_SPREADSHEET_RECORD_ID_RE.fullmatch(value) for value in cleaned):
            return True
        if self._spreadsheet_row_looks_guidance(cleaned):
            return False

        zero_like_count = sum(1 for value in cleaned if _SPREADSHEET_ZERO_LIKE_RE.fullmatch(value))
        control_like_count = sum(1 for value in cleaned if self._spreadsheet_value_is_control(value))
        if control_like_count >= 2:
            return False
        if zero_like_count >= max(2, len(cleaned) // 2):
            return False

        substantive_text_count = 0
        numeric_like_count = 0
        long_text_count = 0
        for value in cleaned:
            if _SPREADSHEET_ZERO_LIKE_RE.fullmatch(value) or self._spreadsheet_value_is_control(value):
                continue
            if _TABLE_NUMBER_LIKE_RE.search(value):
                numeric_like_count += 1
                continue
            substantive_text_count += 1
            if len(value.split()) >= 2 or len(value) >= 16:
                long_text_count += 1

        if substantive_text_count < 3:
            return False
        if numeric_like_count >= 1:
            return True
        return long_text_count >= 2 and substantive_text_count >= 4

    def _extract_xlsx_dataset_mode(
        self,
        *,
        file_detail: KnowledgeUploadFile,
        upload: KnowledgeUpload,
        sheets: Sequence[tuple[int, str, Any]],
        policy_skipped: Sequence[str],
        estimated_rows: int,
    ) -> ExtractionResult:
        dataset_root, dataset_root_rel = self._reset_dataset_storage(upload)
        timestamp = timezone.now().strftime("%Y%m%d%H%M%S")
        rules = self._table_privacy_rules(upload)
        preview_row_cap = int(self.dataset_preview_rows)
        sample_row_cap = int(self.dataset_sample_rows)
        sample_target = max(sample_row_cap, 25)

        key_index_enabled = str(getattr(settings, "DATASET_KEY_INDEX_ENABLED", "true")).lower() in {"1", "true", "yes"}
        key_index_max_bytes = int(getattr(settings, "DATASET_KEY_INDEX_MAX_BYTES", 2_000_000) or 2_000_000)
        key_index_max_bytes = max(4096, min(25_000_000, key_index_max_bytes))

        tables: list[TablePayload] = []
        pages: list[PageLayout] = []
        issues: list[IssuePayload] = []
        dataset_sheets: list[dict[str, Any]] = []

        total_rows = 0
        total_preview_rows = 0
        internal_skipped_total = 0

        for order_index, (sheet_idx, sheet_name, sheet) in enumerate(sheets, start=1):
            try:
                max_col = int(getattr(sheet, "max_column", 0) or 0)
            except (TypeError, ValueError):
                max_col = 0

            rows_iter = sheet.iter_rows(values_only=True)
            header_values: list[str] | None = None
            for raw_row in rows_iter:
                candidate = [self._stringify_dataset_cell(val) for val in raw_row]
                if any(str(value).strip() for value in candidate):
                    header_values = candidate
                    break
            if header_values is None:
                continue

            if max_col <= 0:
                max_col = len(header_values)
            if len(header_values) < max_col:
                header_values.extend([""] * (max_col - len(header_values)))

            column_schema: list[str] = []
            for idx in range(max_col):
                label = header_values[idx] if idx < len(header_values) else ""
                column_schema.append(str(label or "").strip() or f"column_{idx + 1}")

            visible_mask = [
                not self._column_is_sensitive(column, rules)
                for column in column_schema
            ] if rules else [True] * len(column_schema)

            sheet_slug = slugify(sheet_name) or f"sheet_{sheet_idx}"
            dataset_filename = f"{timestamp}_sheet{sheet_idx}_{sheet_slug}.csv.gz"
            dataset_path = dataset_root / dataset_filename
            dataset_rel_path = (dataset_root_rel / dataset_filename).as_posix()

            preview_rows: list[tuple[int, list[str]]] = []
            sample_rows: list[list[str]] = []
            sample_visible_rows: list[dict[str, str]] = []
            row_count = 0
            internal_skipped = 0
            expected_items = 0
            try:
                expected_items = int(getattr(sheet, "max_row", 0) or 0)
            except (TypeError, ValueError):
                expected_items = 0
            sheet_key_indexes: list[dict[str, Any]] = []

            dataset_path.parent.mkdir(parents=True, exist_ok=True)
            with gzip.open(dataset_path, "wt", encoding="utf-8", newline="") as out_handle:
                writer = csv.writer(out_handle)
                writer.writerow(column_schema)
                for raw_row in rows_iter:
                    values = [self._stringify_dataset_cell(val) for val in raw_row]
                    if len(values) < len(column_schema):
                        values.extend([""] * (len(column_schema) - len(values)))
                    elif len(values) > len(column_schema):
                        values = values[: len(column_schema)]

                    if not any(str(value).strip() for value in values):
                        continue

                    attributes = {column_schema[i]: str(values[i] or "") for i in range(len(column_schema))}
                    if rules and self._row_is_internal(attributes, rules):
                        internal_skipped += 1
                        continue

                    row_count += 1
                    writer.writerow([str(value or "") for value in values])

                    if len(preview_rows) < preview_row_cap:
                        preview_rows.append((row_count, [str(value or "") for value in values]))
                    if len(sample_rows) < sample_target:
                        sample_rows.append([str(value or "") for value in values])
                    if len(sample_visible_rows) < sample_row_cap:
                        row_preview = {
                            column_schema[i]: str(values[i] or "")
                            for i in range(len(column_schema))
                            if visible_mask[i]
                        }
                        sample_visible_rows.append(row_preview)

                    if key_index_enabled and not sheet_key_indexes and len(sample_rows) >= sample_target:
                        suggested_preview = self._suggest_dataset_key_columns(column_schema=column_schema, sample_rows=sample_rows)
                        chosen_columns = self._dataset_key_index_columns(
                            upload=upload,
                            sheet_name=sheet_name,
                            column_schema=column_schema,
                            suggested_keys=suggested_preview,
                            rules=rules,
                        )
                        spec = bloom_spec_for_items(int(expected_items or estimated_rows or row_count or 1))
                        for entry in chosen_columns:
                            bits = spec.bits
                            byte_len = (bits + 7) // 8
                            if byte_len > key_index_max_bytes:
                                bits = key_index_max_bytes * 8
                            bloom = BloomFilter(bits=bits, hashes=spec.hashes)
                            sheet_key_indexes.append(
                                {
                                    "column": entry.get("column"),
                                    "column_index": int(entry.get("column_index")),
                                    "source": entry.get("source"),
                                    "bloom": bloom,
                                    "value_count": 0,
                                }
                            )
                        if sheet_key_indexes:
                            for sample in sample_rows:
                                for info in sheet_key_indexes:
                                    idx = int(info["column_index"])
                                    if idx >= len(sample):
                                        continue
                                    normalized = normalize_identifier_value(sample[idx])
                                    if not normalized:
                                        continue
                                    info["bloom"].add(normalized)
                                    info["value_count"] += 1
                    elif sheet_key_indexes:
                        for info in sheet_key_indexes:
                            idx = int(info["column_index"])
                            if idx >= len(values):
                                continue
                            normalized = normalize_identifier_value(values[idx])
                            if not normalized:
                                continue
                            info["bloom"].add(normalized)
                            info["value_count"] += 1

            try:
                dataset_size = dataset_path.stat().st_size
            except OSError:
                dataset_size = 0

            suggested_keys = self._suggest_dataset_key_columns(column_schema=column_schema, sample_rows=sample_rows)
            key_index_payloads: list[dict[str, Any]] = []
            if sheet_key_indexes:
                for info in sheet_key_indexes:
                    column_name = str(info.get("column") or "").strip()
                    if not column_name:
                        continue
                    index_rel_path = resolve_key_index_storage_path(dataset_rel_path=dataset_rel_path, column_name=column_name)
                    abs_index_path = (self.media_root / Path(index_rel_path)).resolve()
                    try:
                        abs_index_path.relative_to(self.media_root)
                    except ValueError:
                        continue
                    bloom = info.get("bloom")
                    if not isinstance(bloom, BloomFilter):
                        continue
                    try:
                        written_bytes = write_bloom_filter(abs_index_path, bloom)
                    except OSError:
                        continue
                    key_index_payloads.append(
                        {
                            "column": column_name,
                            "storage_path": index_rel_path,
                            "bits": int(bloom.bits),
                            "hashes": int(bloom.hashes),
                            "bytes": int(written_bytes),
                            "values_indexed": int(info.get("value_count") or 0),
                            "source": info.get("source"),
                        }
                    )
            table, page_layout = self._build_dataset_preview_table(
                order_index=order_index,
                sheet_name=sheet_name,
                sheet_index=sheet_idx,
                column_schema=column_schema,
                preview_rows=preview_rows,
                file_detail=file_detail,
                source_label="dataset_xlsx",
                content_type="text/csv",
                rules=rules,
            )
            tables.append(table)
            pages.append(page_layout)
            total_preview_rows += len(table.rows)
            total_rows += row_count
            internal_skipped_total += internal_skipped

            dataset_sheets.append(
                {
                    "sheet_index": sheet_idx,
                    "sheet_name": sheet_name,
                    "storage_path": dataset_rel_path,
                    "size_bytes": dataset_size,
                    "row_count": row_count,
                    "estimated_row_count": int(getattr(sheet, "max_row", 0) or 0),
                    "column_schema": column_schema,
                    "sample_rows": sample_visible_rows,
                    "suggested_key_columns": suggested_keys,
                    "key_indexes": key_index_payloads,
                    "preview_rows_indexed": len(table.rows),
                    "internal_rows_skipped": internal_skipped,
                }
            )
            if row_count > len(table.rows):
                issues.append(
                    IssuePayload(
                        code="dataset_preview_truncated",
                        severity=KnowledgeIssueSeverity.INFO.value,
                        description="Only a small preview of this dataset sheet was indexed into Postgres.",
                        page_number=sheet_idx,
                        table_order_index=order_index,
                        details={
                            "rows_stored": row_count,
                            "preview_rows_indexed": len(table.rows),
                            "storage_path": dataset_rel_path,
                        },
                    )
                )

        if not tables:
            raise KnowledgeIngestionError("XLSX workbook did not contain any populated sheets.")

        partial_tables = sum(1 for sheet in dataset_sheets if sheet.get("row_count", 0) > sheet.get("preview_rows_indexed", 0))
        table_stats = self._table_stats_summary(
            total_rows=total_rows,
            indexed_rows=total_preview_rows,
            row_cap=preview_row_cap,
            source_row_count=self._integration_row_count(upload),
            table_count=len(tables),
            partial_tables=partial_tables,
            row_tier="large",
        )
        metadata = {
            "format": "xlsx",
            "filename": file_detail.filename,
            "content_type": file_detail.content_type or "",
            "storage_path": file_detail.storage_path,
            "table_count": len(tables),
            "table_truncation": {
                "truncated_tables": 0,
                "truncated_rows": max(0, total_rows - total_preview_rows),
                "truncated_columns": 0,
            },
            "table_stats": table_stats,
            "dataset": {
                "enabled": True,
                "storage_format": self.dataset_storage_format,
                "row_count": total_rows,
                "estimated_row_count": estimated_rows,
                "preview_rows_indexed": total_preview_rows,
                "internal_rows_skipped": internal_skipped_total,
                "policy_skipped_sheets": list(policy_skipped),
                "sheets": dataset_sheets,
            },
        }

        issues.insert(
            0,
            IssuePayload(
                code="dataset_mode_enabled",
                severity=KnowledgeIssueSeverity.INFO.value,
                description="Large spreadsheet stored in dataset mode (file-backed). Only a small preview is indexed into Postgres.",
                details={
                    "estimated_rows": estimated_rows,
                    "rows_stored": total_rows,
                    "preview_rows_indexed": total_preview_rows,
                    "internal_rows_skipped": internal_skipped_total,
                    "sheet_count": len(dataset_sheets),
                },
            ),
        )

        preview_text = self._table_preview_text(tables, rules=rules)
        table_entities = self._table_row_entities(
            tables,
            business_profile=getattr(upload, "business_profile", None),
            upload=upload,
        )
        return ExtractionResult(
            text=preview_text,
            format_hint="xlsx",
            metadata=metadata,
            pages=pages,
            tables=tables,
            issues=issues,
            entities=table_entities,
        )

    def _extract_xls(
        self,
        path: Path,
        *,
        file_detail: KnowledgeUploadFile,
        upload: KnowledgeUpload | None = None,
    ) -> ExtractionResult:
        if xlrd is None:
            raise KnowledgeIngestionError("XLS ingestion requires the xlrd package.")
        try:
            workbook = xlrd.open_workbook(filename=str(path), on_demand=True)
        except Exception as exc:
            raise KnowledgeIngestionError(f"Unable to open XLS file: {exc}") from exc

        policy = resolve_normalization_policy(upload)
        rules = self._table_privacy_rules(upload)

        allowed_sheets: list[tuple[int, str, Any]] = []
        estimated_rows = 0
        policy_skipped: list[str] = []
        for sheet_idx in range(1, (getattr(workbook, "nsheets", 0) or 0) + 1):
            sheet = workbook.sheet_by_index(sheet_idx - 1)
            sheet_name = getattr(sheet, "name", None) or f"Sheet {sheet_idx}"
            if not sheet_is_allowed(sheet_name, policy):
                policy_skipped.append(sheet_name)
                continue
            allowed_sheets.append((sheet_idx, sheet_name, sheet))
            try:
                estimated_rows += int(getattr(sheet, "nrows", 0) or 0)
            except (TypeError, ValueError):
                continue

        if upload and self._should_use_dataset_mode(estimated_rows=estimated_rows):
            return self._extract_xls_dataset_mode(
                workbook=workbook,
                file_detail=file_detail,
                upload=upload,
                sheets=allowed_sheets,
                policy_skipped=policy_skipped,
                estimated_rows=estimated_rows,
            )

        diagnostics: list[SheetNormalizationDiagnostics] = [
            SheetNormalizationDiagnostics(sheet_name=name, skipped=True, skip_reason="policy")
            for name in policy_skipped
        ]
        tables: list[TablePayload] = []
        pages: list[PageLayout] = []
        order_index = 1

        for sheet_idx, sheet_name, sheet in allowed_sheets:
            raw_rows: list[list[Any]] = []
            nrows = getattr(sheet, "nrows", 0) or 0
            ncols = getattr(sheet, "ncols", 0) or 0
            for row_idx in range(nrows):
                row_values: list[Any] = []
                for col_idx in range(ncols):
                    cell_value = sheet.cell_value(row_idx, col_idx)
                    cell_type = sheet.cell_type(row_idx, col_idx)
                    if cell_type == xlrd.XL_CELL_DATE:
                        try:
                            cell_value = xlrd.xldate_as_datetime(cell_value, workbook.datemode)
                        except Exception:
                            cell_value = ""
                    elif cell_type == xlrd.XL_CELL_BOOLEAN:
                        cell_value = bool(cell_value)
                    elif cell_type == xlrd.XL_CELL_ERROR:
                        cell_value = ""
                    row_values.append(cell_value)
                raw_rows.append(row_values)
            normalized = normalize_sheet_rows(
                raw_rows,
                sheet_name=sheet_name,
                policy=policy,
            )
            diagnostics.append(normalized.diagnostics)
            if normalized.diagnostics.skipped:
                continue
            table, page_layout = self._build_table_from_normalized_sheet(
                normalized,
                order_index=order_index,
                sheet_name=sheet_name,
                sheet_index=sheet_idx,
                file_detail=file_detail,
                source_label="xls",
                content_type="application/vnd.ms-excel",
                rules=rules,
            )
            tables.append(table)
            pages.append(page_layout)
            order_index += 1

        original_tables = list(tables)
        if not original_tables:
            raise KnowledgeIngestionError("XLS workbook did not contain any populated sheets.")
        ingest_config = self._table_ingest_config(upload)
        tables, table_metrics, limit_issues, table_summary = self._apply_table_limits(
            original_tables,
            upload=upload,
            config=ingest_config,
        )
        if not tables:
            raise KnowledgeIngestionError("XLS workbook exceeded configured limits and no rows were indexed.")
        preview_text = self._table_preview_text(tables, rules=rules)
        table_entities = self._table_row_entities(
            tables,
            business_profile=getattr(upload, "business_profile", None),
            upload=upload,
        )
        table_stats = self._table_stats_summary(
            total_rows=table_summary["total_rows"],
            indexed_rows=table_summary["indexed_rows"],
            row_cap=table_summary.get("row_cap_hint"),
            source_row_count=self._integration_row_count(upload),
            table_count=len(tables),
            partial_tables=table_summary.get("partial_tables", 0),
            row_tier=table_summary.get("row_tier_hint"),
        )
        normalization_summary = summarize_normalization(policy, diagnostics)
        metadata = {
            "format": "xls",
            "filename": file_detail.filename,
            "content_type": file_detail.content_type or "",
            "storage_path": file_detail.storage_path,
            "table_count": len(tables),
            "table_truncation": table_metrics,
            "table_stats": table_stats,
        }
        if normalization_summary:
            metadata["normalization"] = normalization_summary
        _log_normalization_summary(upload, "xls", normalization_summary)
        return ExtractionResult(
            text=preview_text,
            format_hint="xls",
            metadata=metadata,
            pages=pages,
            tables=tables,
            issues=limit_issues,
            entities=table_entities,
        )

    def _extract_xls_dataset_mode(
        self,
        *,
        workbook: Any,
        file_detail: KnowledgeUploadFile,
        upload: KnowledgeUpload,
        sheets: Sequence[tuple[int, str, Any]],
        policy_skipped: Sequence[str],
        estimated_rows: int,
    ) -> ExtractionResult:
        dataset_root, dataset_root_rel = self._reset_dataset_storage(upload)
        timestamp = timezone.now().strftime("%Y%m%d%H%M%S")
        rules = self._table_privacy_rules(upload)
        preview_row_cap = int(self.dataset_preview_rows)
        sample_row_cap = int(self.dataset_sample_rows)

        tables: list[TablePayload] = []
        pages: list[PageLayout] = []
        issues: list[IssuePayload] = []
        dataset_sheets: list[dict[str, Any]] = []

        total_rows = 0
        total_preview_rows = 0
        internal_skipped_total = 0

        for order_index, (sheet_idx, sheet_name, sheet) in enumerate(sheets, start=1):
            nrows = int(getattr(sheet, "nrows", 0) or 0)
            ncols = int(getattr(sheet, "ncols", 0) or 0)
            if nrows <= 0 or ncols <= 0:
                continue

            header_row_idx: int | None = None
            header_values: list[str] | None = None
            scan_limit = min(nrows, 50)
            for idx in range(scan_limit):
                raw_values = list(getattr(sheet, "row_values")(idx))
                candidate = [self._stringify_dataset_cell(val) for val in raw_values[:ncols]]
                if any(str(value).strip() for value in candidate):
                    header_row_idx = idx
                    header_values = candidate
                    break
            if header_row_idx is None or header_values is None:
                continue
            if len(header_values) < ncols:
                header_values.extend([""] * (ncols - len(header_values)))

            column_schema: list[str] = []
            for idx in range(ncols):
                label = header_values[idx] if idx < len(header_values) else ""
                column_schema.append(str(label or "").strip() or f"column_{idx + 1}")

            visible_mask = [
                not self._column_is_sensitive(column, rules)
                for column in column_schema
            ] if rules else [True] * len(column_schema)

            sheet_slug = slugify(sheet_name) or f"sheet_{sheet_idx}"
            dataset_filename = f"{timestamp}_sheet{sheet_idx}_{sheet_slug}.csv.gz"
            dataset_path = dataset_root / dataset_filename
            dataset_rel_path = (dataset_root_rel / dataset_filename).as_posix()

            preview_rows: list[tuple[int, list[str]]] = []
            sample_rows: list[list[str]] = []
            sample_visible_rows: list[dict[str, str]] = []
            row_count = 0
            internal_skipped = 0

            dataset_path.parent.mkdir(parents=True, exist_ok=True)
            with gzip.open(dataset_path, "wt", encoding="utf-8", newline="") as out_handle:
                writer = csv.writer(out_handle)
                writer.writerow(column_schema)

                for row_idx in range(header_row_idx + 1, nrows):
                    raw_values = list(getattr(sheet, "row_values")(row_idx))
                    raw_types = list(getattr(sheet, "row_types")(row_idx))
                    values: list[str] = []
                    for col_idx in range(ncols):
                        cell_value = raw_values[col_idx] if col_idx < len(raw_values) else ""
                        cell_type = raw_types[col_idx] if col_idx < len(raw_types) else None
                        if cell_type == xlrd.XL_CELL_DATE:
                            try:
                                cell_value = xlrd.xldate_as_datetime(cell_value, workbook.datemode)
                            except Exception:
                                cell_value = ""
                        elif cell_type == xlrd.XL_CELL_BOOLEAN:
                            cell_value = bool(cell_value)
                        elif cell_type == xlrd.XL_CELL_ERROR:
                            cell_value = ""
                        values.append(self._stringify_dataset_cell(cell_value))

                    if not any(str(value).strip() for value in values):
                        continue

                    attributes = {column_schema[i]: str(values[i] or "") for i in range(len(column_schema))}
                    if rules and self._row_is_internal(attributes, rules):
                        internal_skipped += 1
                        continue

                    row_count += 1
                    writer.writerow(values)

                    if len(preview_rows) < preview_row_cap:
                        preview_rows.append((row_count, list(values)))
                    if len(sample_rows) < max(sample_row_cap, 25):
                        sample_rows.append(list(values))
                    if len(sample_visible_rows) < sample_row_cap:
                        row_preview = {
                            column_schema[i]: str(values[i] or "")
                            for i in range(len(column_schema))
                            if visible_mask[i]
                        }
                        sample_visible_rows.append(row_preview)

            try:
                dataset_size = dataset_path.stat().st_size
            except OSError:
                dataset_size = 0

            suggested_keys = self._suggest_dataset_key_columns(column_schema=column_schema, sample_rows=sample_rows)
            table, page_layout = self._build_dataset_preview_table(
                order_index=order_index,
                sheet_name=sheet_name,
                sheet_index=sheet_idx,
                column_schema=column_schema,
                preview_rows=preview_rows,
                file_detail=file_detail,
                source_label="dataset_xls",
                content_type="text/csv",
                rules=rules,
            )
            tables.append(table)
            pages.append(page_layout)
            total_preview_rows += len(table.rows)
            total_rows += row_count
            internal_skipped_total += internal_skipped

            dataset_sheets.append(
                {
                    "sheet_index": sheet_idx,
                    "sheet_name": sheet_name,
                    "storage_path": dataset_rel_path,
                    "size_bytes": dataset_size,
                    "row_count": row_count,
                    "estimated_row_count": nrows,
                    "column_schema": column_schema,
                    "sample_rows": sample_visible_rows,
                    "suggested_key_columns": suggested_keys,
                    "preview_rows_indexed": len(table.rows),
                    "internal_rows_skipped": internal_skipped,
                }
            )
            if row_count > len(table.rows):
                issues.append(
                    IssuePayload(
                        code="dataset_preview_truncated",
                        severity=KnowledgeIssueSeverity.INFO.value,
                        description="Only a small preview of this dataset sheet was indexed into Postgres.",
                        page_number=sheet_idx,
                        table_order_index=order_index,
                        details={
                            "rows_stored": row_count,
                            "preview_rows_indexed": len(table.rows),
                            "storage_path": dataset_rel_path,
                        },
                    )
                )

        if not tables:
            raise KnowledgeIngestionError("XLS workbook did not contain any populated sheets.")

        partial_tables = sum(1 for sheet in dataset_sheets if sheet.get("row_count", 0) > sheet.get("preview_rows_indexed", 0))
        table_stats = self._table_stats_summary(
            total_rows=total_rows,
            indexed_rows=total_preview_rows,
            row_cap=preview_row_cap,
            source_row_count=self._integration_row_count(upload),
            table_count=len(tables),
            partial_tables=partial_tables,
            row_tier="large",
        )
        metadata = {
            "format": "xls",
            "filename": file_detail.filename,
            "content_type": file_detail.content_type or "",
            "storage_path": file_detail.storage_path,
            "table_count": len(tables),
            "table_truncation": {
                "truncated_tables": 0,
                "truncated_rows": max(0, total_rows - total_preview_rows),
                "truncated_columns": 0,
            },
            "table_stats": table_stats,
            "dataset": {
                "enabled": True,
                "storage_format": self.dataset_storage_format,
                "row_count": total_rows,
                "estimated_row_count": estimated_rows,
                "preview_rows_indexed": total_preview_rows,
                "internal_rows_skipped": internal_skipped_total,
                "policy_skipped_sheets": list(policy_skipped),
                "sheets": dataset_sheets,
            },
        }

        issues.insert(
            0,
            IssuePayload(
                code="dataset_mode_enabled",
                severity=KnowledgeIssueSeverity.INFO.value,
                description="Large spreadsheet stored in dataset mode (file-backed). Only a small preview is indexed into Postgres.",
                details={
                    "estimated_rows": estimated_rows,
                    "rows_stored": total_rows,
                    "preview_rows_indexed": total_preview_rows,
                    "internal_rows_skipped": internal_skipped_total,
                    "sheet_count": len(dataset_sheets),
                },
            ),
        )

        preview_text = self._table_preview_text(tables, rules=rules)
        table_entities = self._table_row_entities(
            tables,
            business_profile=getattr(upload, "business_profile", None),
            upload=upload,
        )
        return ExtractionResult(
            text=preview_text,
            format_hint="xls",
            metadata=metadata,
            pages=pages,
            tables=tables,
            issues=issues,
            entities=table_entities,
        )

    def _extract_jsonl_dataset(
        self,
        path: Path,
        *,
        file_detail: KnowledgeUploadFile,
        upload: KnowledgeUpload | None = None,
    ) -> ExtractionResult:
        if upload is None:
            raise KnowledgeIngestionError("JSONL dataset ingestion requires an upload record.")

        dataset_root, dataset_root_rel = self._reset_dataset_storage(upload)
        timestamp = timezone.now().strftime("%Y%m%d%H%M%S")
        display_name = (file_detail.filename or "Dataset").strip() or "Dataset"
        safe_label = slugify(Path(display_name).stem) or "dataset"
        dataset_filename = f"{timestamp}_{safe_label}.jsonl.gz"
        dataset_path = dataset_root / dataset_filename
        dataset_rel_path = (dataset_root_rel / dataset_filename).as_posix()

        rules = self._table_privacy_rules(upload)
        preview_row_cap = int(self.dataset_preview_rows)
        sample_row_cap = int(self.dataset_sample_rows)

        row_count = 0
        sample_lines: list[bytes] = []
        dataset_path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("rb") as raw_in, gzip.open(dataset_path, "wb") as out_handle:
            for line in raw_in:
                if not line:
                    continue
                out_handle.write(line)
                if not line.strip():
                    continue
                row_count += 1
                if len(sample_lines) < max(sample_row_cap, 25):
                    sample_lines.append(line)

        try:
            dataset_size = dataset_path.stat().st_size
        except OSError:
            dataset_size = 0

        records: list[dict[str, Any]] = []
        keys: set[str] = set()
        parse_errors = 0
        for line in sample_lines:
            try:
                decoded = line.decode("utf-8", errors="ignore")
                obj = json.loads(decoded)
            except Exception:
                parse_errors += 1
                continue
            if isinstance(obj, dict):
                records.append(obj)
                for key in obj.keys():
                    if key is None:
                        continue
                    keys.add(str(key))

        if not keys:
            preview_text = "\n".join(line.decode("utf-8", errors="ignore").strip() for line in sample_lines[:5]).strip()
            issues = [
                IssuePayload(
                    code="dataset_mode_enabled",
                    severity=KnowledgeIssueSeverity.INFO.value,
                    description="JSONL dataset stored in dataset mode (file-backed). No structured preview was indexed.",
                    details={
                        "rows_stored": row_count,
                        "sample_parse_errors": parse_errors,
                        "dataset_storage_format": "jsonl_gz",
                    },
                )
            ]
            metadata = {
                "format": "jsonl",
                "filename": file_detail.filename,
                "content_type": file_detail.content_type or "",
                "storage_path": file_detail.storage_path,
                "dataset": {
                    "enabled": True,
                    "storage_format": "jsonl_gz",
                    "storage_path": dataset_rel_path,
                    "size_bytes": dataset_size,
                    "row_count": row_count,
                    "sample_parse_errors": parse_errors,
                },
            }
            return ExtractionResult(
                text=preview_text,
                format_hint="jsonl",
                metadata=metadata,
                pages=[],
                tables=[],
                issues=issues,
                entities=[],
            )

        column_schema = sorted(keys)[:200]
        visible_mask = [
            not self._column_is_sensitive(column, rules)
            for column in column_schema
        ] if rules else [True] * len(column_schema)

        preview_rows: list[tuple[int, list[str]]] = []
        sample_rows: list[list[str]] = []
        sample_visible_rows: list[dict[str, str]] = []

        for idx, record in enumerate(records, start=1):
            values: list[str] = []
            for key in column_schema:
                values.append(self._stringify_dataset_cell(record.get(key)))
            preview_rows.append((idx, values))
            if len(sample_rows) < max(sample_row_cap, 25):
                sample_rows.append(values)
            if len(sample_visible_rows) < sample_row_cap:
                row_preview = {
                    column_schema[i]: values[i]
                    for i in range(len(column_schema))
                    if visible_mask[i]
                }
                sample_visible_rows.append(row_preview)
            if len(preview_rows) >= preview_row_cap:
                break

        suggested_keys = self._suggest_dataset_key_columns(column_schema=column_schema, sample_rows=sample_rows)
        table, page = self._build_dataset_preview_table(
            order_index=1,
            sheet_name=display_name,
            sheet_index=None,
            column_schema=column_schema,
            preview_rows=preview_rows,
            file_detail=file_detail,
            source_label="dataset_jsonl",
            content_type="application/x-ndjson",
            rules=rules,
        )
        tables = [table]
        preview_text = self._table_preview_text(tables, rules=rules)

        issues = [
            IssuePayload(
                code="dataset_mode_enabled",
                severity=KnowledgeIssueSeverity.INFO.value,
                description="JSONL dataset stored in dataset mode (file-backed). Only a small preview is indexed into Postgres.",
                details={
                    "rows_stored": row_count,
                    "preview_rows_indexed": len(table.rows),
                    "sample_parse_errors": parse_errors,
                    "dataset_storage_format": "jsonl_gz",
                },
            )
        ]
        dataset_metadata = {
            "enabled": True,
            "storage_format": "jsonl_gz",
            "storage_path": dataset_rel_path,
            "size_bytes": dataset_size,
            "row_count": row_count,
            "column_schema": column_schema,
            "sample_rows": sample_visible_rows,
            "suggested_key_columns": suggested_keys,
            "sample_parse_errors": parse_errors,
        }
        table_stats = self._table_stats_summary(
            total_rows=row_count,
            indexed_rows=len(table.rows),
            row_cap=len(table.rows),
            source_row_count=self._integration_row_count(upload),
            table_count=1,
            partial_tables=1 if row_count > len(table.rows) else 0,
            row_tier="large" if row_count >= self.dataset_row_threshold else "medium",
        )
        metadata = {
            "format": "jsonl",
            "filename": file_detail.filename,
            "content_type": file_detail.content_type or "",
            "storage_path": file_detail.storage_path,
            "table_count": 1,
            "table_truncation": {
                "truncated_tables": 0,
                "truncated_rows": max(0, row_count - len(table.rows)),
                "truncated_columns": 0,
            },
            "table_stats": table_stats,
            "dataset": dataset_metadata,
        }
        return ExtractionResult(
            text=preview_text,
            format_hint="jsonl",
            metadata=metadata,
            pages=[page],
            tables=tables,
            issues=issues,
            entities=self._table_row_entities(
                tables,
                business_profile=getattr(upload, "business_profile", None),
                upload=upload,
            ),
        )
