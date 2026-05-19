from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Mapping, Sequence

from apps.accounts.models import KnowledgeIssueSeverity
from apps.knowledge.ingestion.contracts import IssuePayload, TableCellPayload, TablePayload, TableRowPayload
from apps.knowledge.ingestion.page_renderer import PageRenderer
from apps.knowledge.ingestion.text_utils import IngestionTextUtilsMixin
from apps.knowledge.tables.docx_tables.headers import IngestionDocxHeaderRowsMixin
from apps.knowledge.tables.docx_tables.schema import IngestionDocxSchemaMixin
from apps.knowledge.tables.docx_tables.sparse_series import IngestionDocxSparseSeriesMixin

try:  # pragma: no cover - dependency failure should be surfaced at runtime
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover - fallback handled via runtime check
    DocxDocument = None  # type: ignore


class IngestionDocxTablesMixin(
    IngestionDocxSchemaMixin,
    IngestionDocxSparseSeriesMixin,
    IngestionDocxHeaderRowsMixin,
):

    def _extract_docx_table_candidates(
        self,
        path: Path,
        *,
        filename: str = "",
    ) -> tuple[list[TablePayload], list[IssuePayload], dict[str, Any]]:
        """
        Parse DOCX tables directly from `document.tables` and emit structured TablePayloads.

        This is the primary DOCX table path. Paragraph extraction remains supplemental.
        """
        if DocxDocument is None:
            return [], [
                IssuePayload(
                    code="docx_tables_missing_dependency",
                    severity=KnowledgeIssueSeverity.WARNING.value,
                    description="python-docx is unavailable; DOCX table extraction skipped.",
                    page_number=1,
                )
            ], {"enabled": False, "reason": "python_docx_missing"}

        try:
            document = DocxDocument(str(path))
        except Exception as exc:
            return [], [
                IssuePayload(
                    code="docx_tables_parse_failed",
                    severity=KnowledgeIssueSeverity.ERROR.value,
                    description=f"DOCX table extraction failed: {exc}",
                    page_number=1,
                )
            ], {"enabled": False, "reason": "parse_failed"}

        heading_map = self._docx_table_heading_map(document)
        tables: list[TablePayload] = []
        issues: list[IssuePayload] = []
        merged_regions_total = 0
        header_rows_total = 0
        compacted_blank_rows_total = 0
        section_rows_total = 0
        series_columns_total = 0

        for order_index, table in enumerate(document.tables, start=1):
            key_grid, text_by_key, span_by_key = self._docx_table_grid(table)
            key_grid, span_by_key, grid_meta = self._docx_compact_table_grid(
                key_grid,
                text_by_key,
                span_by_key,
            )
            provisional_header_rows = self._docx_detect_header_rows(key_grid, text_by_key)
            key_grid, text_by_key, span_by_key, collapse_meta = self._docx_collapse_helper_columns(
                key_grid,
                text_by_key,
                span_by_key,
                provisional_header_rows,
            )
            key_grid, text_by_key, span_by_key, series_meta = self._docx_normalize_sparse_series_columns(
                key_grid,
                text_by_key,
                span_by_key,
            )
            row_count = len(key_grid)
            col_count = max((len(row) for row in key_grid), default=0)
            if row_count <= 0 or col_count <= 0:
                issues.append(
                    IssuePayload(
                        code="docx_table_empty",
                        severity=KnowledgeIssueSeverity.INFO.value,
                        description=f"DOCX table {order_index} had no readable cells.",
                        page_number=1,
                        table_order_index=order_index,
                    )
                )
                continue

            header_rows = self._docx_detect_header_rows(key_grid, text_by_key)
            header_row_set = set(header_rows)
            section_rows = self._docx_detect_section_rows(key_grid, text_by_key, header_row_set)
            section_row_set = set(section_rows)
            if header_rows:
                header_rows_total += len(header_rows)
            compacted_blank_rows_total += int(grid_meta.get("removed_blank_rows") or 0)
            if section_rows:
                section_rows_total += len(section_rows)
            series_columns_total += int(series_meta.get("merged_columns") or 0)

            column_schema: list[str] = []
            for col_idx in range(col_count):
                labels: list[str] = []
                seen_labels: set[str] = set()
                for row_idx in header_rows:
                    key = key_grid[row_idx][col_idx] if col_idx < len(key_grid[row_idx]) else None
                    if key is None:
                        continue
                    label = self._docx_canonicalize_header_label(
                        self._sanitize_text(text_by_key.get(key, "")).strip()
                    )
                    if not label:
                        continue
                    dedupe_key = re.sub(r"\s+", " ", label).strip().lower()
                    if dedupe_key in seen_labels:
                        continue
                    seen_labels.add(dedupe_key)
                    labels.append(label)
                merged_label = " | ".join(labels).strip()
                column_schema.append(self._docx_normalize_column_key(merged_label, col_idx))

            if not any(column_schema):
                column_schema = [f"column_{idx + 1}" for idx in range(col_count)]
            else:
                column_schema = self._docx_refine_grouped_column_schema(
                    column_schema,
                    key_grid,
                    text_by_key,
                    header_rows,
                )

            table_rows: list[TableRowPayload] = []
            merged_regions = 0
            for row_idx, row_keys in enumerate(key_grid):
                row_type = "data"
                if row_idx in header_row_set:
                    row_type = "header"
                elif row_idx in section_row_set:
                    row_type = "section_header"

                initial_row_values = [
                    self._sanitize_text(
                        text_by_key.get(row_keys[col_idx], "")
                        if col_idx < len(row_keys) and row_keys[col_idx] is not None
                        else ""
                    ).strip()
                    for col_idx in range(col_count)
                ]
                repeated_section_label = ""
                repeated_section_col_idx: int | None = None
                if row_type == "section_header":
                    repeated_section_label, repeated_section_col_idx = self._docx_repeated_row_label_info(
                        initial_row_values
                    )

                row_cells: list[TableCellPayload] = []
                row_values: list[str] = []
                for col_idx in range(col_count):
                    key = row_keys[col_idx] if col_idx < len(row_keys) else None
                    raw_text = self._sanitize_text(text_by_key.get(key, "") if key is not None else "").strip()
                    if repeated_section_label and raw_text:
                        canonical_cell = self._docx_canonicalize_header_label(raw_text)
                        if canonical_cell.strip().lower() == repeated_section_label.strip().lower():
                            raw_text = repeated_section_label if col_idx == repeated_section_col_idx else ""
                    row_values.append(raw_text)
                    cell_metadata: dict[str, Any] = {}
                    if key is not None:
                        span = span_by_key.get(key) or {}
                        row_span = int(span.get("row_span") or 1)
                        col_span = int(span.get("column_span") or 1)
                        if row_span > 1 or col_span > 1:
                            is_anchor = (
                                row_idx == int(span.get("row_start") or 0)
                                and col_idx == int(span.get("col_start") or 0)
                            )
                            cell_metadata["row_span"] = row_span
                            cell_metadata["column_span"] = col_span
                            cell_metadata["merged_anchor"] = is_anchor
                            if is_anchor:
                                merged_regions += 1
                            else:
                                cell_metadata["merged_from"] = {
                                    "row_index": int(span.get("row_start") or 0),
                                    "column_index": int(span.get("col_start") or 0),
                                }

                    row_cells.append(
                        TableCellPayload(
                            row_index=row_idx,
                            column_index=col_idx,
                            column_key=column_schema[col_idx] if col_idx < len(column_schema) else f"column_{col_idx + 1}",
                            raw_text=raw_text,
                            normalized_value=self.table_detector._normalize_cell_value(raw_text),
                            metadata=cell_metadata,
                        )
                    )

                row_metadata: dict[str, Any] = {"row_type": row_type}
                if row_idx in header_row_set:
                    row_metadata["header_source"] = "docx_detected"
                    row_metadata["header_level"] = header_rows.index(row_idx) + 1
                elif row_idx in section_row_set:
                    row_metadata["section_source"] = "docx_detected"
                table_rows.append(
                    TableRowPayload(
                        row_index=row_idx,
                        page_number=1,
                        raw_text="\t".join(row_values),
                        metadata=row_metadata,
                        cells=row_cells,
                    )
                )

            merged_regions_total += merged_regions
            section_heading = self._sanitize_text(heading_map.get(order_index, "")).strip()
            title = section_heading or f"Table {order_index}"
            table_metadata: dict[str, Any] = {
                "detected_via": "docx:table_xml",
                "extractor": "docx_table_parser",
                "table_index": order_index,
                "row_count": row_count,
                "column_count": col_count,
                "header_rows": list(header_rows),
                "section_rows": list(section_rows),
                "merged_regions": merged_regions,
                "structure_confidence": 0.95,
            }
            if grid_meta.get("removed_blank_rows"):
                table_metadata["blank_rows_compacted"] = int(grid_meta.get("removed_blank_rows") or 0)
            if collapse_meta.get("removed_columns"):
                table_metadata["helper_columns_removed"] = int(collapse_meta.get("removed_columns") or 0)
            if collapse_meta.get("merged_columns"):
                table_metadata["helper_columns_merged"] = int(collapse_meta.get("merged_columns") or 0)
            if series_meta.get("merged_columns"):
                table_metadata["series_columns_merged"] = int(series_meta.get("merged_columns") or 0)
            if filename:
                table_metadata["filename"] = filename
            if section_heading:
                table_metadata["section_heading"] = section_heading
            style_name = self._sanitize_text(getattr(getattr(table, "style", None), "name", "")).strip()
            if style_name:
                table_metadata["style_name"] = style_name

            tables.append(
                TablePayload(
                    order_index=order_index,
                    title=title,
                    section_heading=section_heading,
                    page_number=1,
                    column_schema=column_schema,
                    data_dictionary={},
                    metadata=table_metadata,
                    rows=table_rows,
                )
            )

        meta: dict[str, Any] = {
            "enabled": True,
            "table_count": len(tables),
            "header_rows_detected": header_rows_total,
            "section_rows_detected": section_rows_total,
            "blank_rows_compacted": compacted_blank_rows_total,
            "merged_regions": merged_regions_total,
            "series_columns_merged": series_columns_total,
        }
        return tables, issues, meta

    def _docx_table_heading_map(self, document: Any) -> dict[int, str]:
        body = getattr(getattr(document, "element", None), "body", None)
        if body is None:
            return {}
        try:
            from docx.text.paragraph import Paragraph as DocxParagraphClass  # type: ignore
        except Exception:
            return {}

        current_heading = ""
        heading_map: dict[int, str] = {}
        table_index = 0
        for child in body.iterchildren():
            child_tag = str(getattr(child, "tag", "") or "")
            if child_tag.endswith("}p"):
                paragraph = DocxParagraphClass(child, document)
                text = self._sanitize_text(getattr(paragraph, "text", "")).strip()
                if text and PageRenderer._looks_like_heading(text):
                    current_heading = text
                continue
            if child_tag.endswith("}tbl"):
                table_index += 1
                if current_heading:
                    heading_map[table_index] = current_heading
        return heading_map

    @staticmethod
    def _docx_table_grid(
        table: Any,
    ) -> tuple[list[list[int | None]], dict[int, str], dict[int, dict[str, int]]]:
        rows = [list(getattr(row, "cells", []) or []) for row in getattr(table, "rows", [])]
        if not rows:
            return [], {}, {}
        column_count = max((len(cells) for cells in rows), default=0)
        if column_count <= 0:
            return [], {}, {}

        key_grid: list[list[int | None]] = []
        text_by_key: dict[int, str] = {}
        positions_by_key: dict[int, list[tuple[int, int]]] = {}

        for row_idx, row_cells in enumerate(rows):
            row_keys: list[int | None] = []
            for col_idx in range(column_count):
                if col_idx >= len(row_cells):
                    row_keys.append(None)
                    continue
                cell = row_cells[col_idx]
                tc = getattr(cell, "_tc", None)
                key = id(tc) if tc is not None else id(cell)
                row_keys.append(key)
                positions_by_key.setdefault(key, []).append((row_idx, col_idx))
                if key not in text_by_key:
                    text_by_key[key] = IngestionTextUtilsMixin._sanitize_text(getattr(cell, "text", "")).strip()
            key_grid.append(row_keys)

        span_by_key: dict[int, dict[str, int]] = {}
        for key, positions in positions_by_key.items():
            row_start = min(pos[0] for pos in positions)
            row_end = max(pos[0] for pos in positions)
            col_start = min(pos[1] for pos in positions)
            col_end = max(pos[1] for pos in positions)
            span_by_key[key] = {
                "row_start": row_start,
                "row_end": row_end,
                "col_start": col_start,
                "col_end": col_end,
                "row_span": (row_end - row_start) + 1,
                "column_span": (col_end - col_start) + 1,
            }
        return key_grid, text_by_key, span_by_key

    @staticmethod
    def _docx_compact_table_grid(
        key_grid: Sequence[Sequence[int | None]],
        text_by_key: Mapping[int, str],
        span_by_key: Mapping[int, Mapping[str, int]],
    ) -> tuple[list[list[int | None]], dict[int, dict[str, int]], dict[str, Any]]:
        if not key_grid:
            return [], {}, {"removed_blank_rows": 0}

        compacted_grid: list[list[int | None]] = []
        positions_by_key: dict[int, list[tuple[int, int]]] = {}
        removed_blank_rows = 0

        for row in key_grid:
            has_readable_value = any(
                key is not None and str(text_by_key.get(key, "")).strip()
                for key in row
            )
            if not has_readable_value:
                removed_blank_rows += 1
                continue
            new_row = list(row)
            new_row_index = len(compacted_grid)
            compacted_grid.append(new_row)
            for col_idx, key in enumerate(new_row):
                if key is None:
                    continue
                positions_by_key.setdefault(key, []).append((new_row_index, col_idx))

        if not compacted_grid:
            return [], {}, {"removed_blank_rows": removed_blank_rows}

        compacted_spans: dict[int, dict[str, int]] = {}
        for key, positions in positions_by_key.items():
            prior_span = span_by_key.get(key) or {}
            row_start = min(pos[0] for pos in positions)
            row_end = max(pos[0] for pos in positions)
            col_start = min(pos[1] for pos in positions)
            col_end = max(pos[1] for pos in positions)
            compacted_spans[key] = {
                "row_start": row_start,
                "row_end": row_end,
                "col_start": col_start,
                "col_end": col_end,
                "row_span": (row_end - row_start) + 1,
                "column_span": int(prior_span.get("column_span") or ((col_end - col_start) + 1)),
            }

        return compacted_grid, compacted_spans, {"removed_blank_rows": removed_blank_rows}
