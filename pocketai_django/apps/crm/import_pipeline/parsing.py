from __future__ import annotations

import csv
import io
from datetime import timedelta
from typing import Any, Iterable, Sequence

from django.core.exceptions import ValidationError
from django.utils import timezone
from openpyxl import load_workbook

from apps.crm.domain.services import log_crm_activity
from apps.crm.import_pipeline.config import (
    MAX_IMPORT_COLUMNS,
    MAX_IMPORT_FILE_BYTES,
    MAX_IMPORT_ROWS,
    RETRY_BACKOFF_MINUTES,
)
from apps.crm.models import CrmActivityActor, CrmActivityType, CrmImportJob, CrmImportJobStatus, CrmImportRowStatus, CrmImportSourceFormat


def _infer_default_mapping(columns: Sequence[str]) -> dict[str, Any]:
    normalized = {str(column): _normalize_column_name(column) for column in columns or []}
    contact_columns: dict[str, Any] = {}
    company_columns: dict[str, Any] = {}
    external_identity: dict[str, Any] | None = None

    contact_targets = {
        "displayname": "display_name",
        "name": "display_name",
        "fullname": "display_name",
        "firstname": "first_name",
        "lastname": "last_name",
        "email": "primary_email",
        "primaryemail": "primary_email",
        "phone": "primary_phone",
        "mobile": "primary_phone",
        "phonenumber": "primary_phone",
        "title": "title",
    }
    company_targets = {
        "company": "name",
        "companyname": "name",
        "account": "name",
        "website": "website",
        "domain": "website",
        "companyphone": "primary_phone",
    }
    external_keys = {"id", "customerid", "contactid", "recordid", "externalid", "companyid", "accountid"}

    for source_column, key in normalized.items():
        if key in contact_targets:
            contact_columns[source_column] = contact_targets[key]
        elif key in company_targets:
            company_columns[source_column] = company_targets[key]
        elif key in external_keys and external_identity is None:
            external_identity = {
                "source_column": source_column,
                "source_system": "csv_import",
                "external_object_type": "contact",
                "external_label": source_column,
            }

    mapping: dict[str, Any] = {}
    if contact_columns or external_identity:
        mapping["contact"] = {"columns": contact_columns}
        if external_identity:
            mapping["contact"]["external_identity"] = external_identity
    if company_columns:
        mapping["company"] = {"columns": company_columns}
    return mapping


def _normalize_column_name(value: str) -> str:
    return "".join(ch for ch in str(value or "").strip().lower() if ch.isalnum())


def _load_rows(handle, *, file_format: str, limit: int | None = None) -> Iterable[dict[str, str]]:
    if file_format == CrmImportSourceFormat.XLSX:
        workbook = load_workbook(handle, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator, [])]
        if len(headers) > MAX_IMPORT_COLUMNS:
            raise ValidationError({"source_file": f"Import exceeds the maximum of {MAX_IMPORT_COLUMNS} columns."})
        for index, row in enumerate(iterator, start=1):
            if limit is None and index > MAX_IMPORT_ROWS:
                raise ValidationError({"source_file": f"Import exceeds the maximum of {MAX_IMPORT_ROWS} rows."})
            if limit is not None and index > limit:
                break
            yield {headers[col_index]: str(value or "").strip() for col_index, value in enumerate(row) if col_index < len(headers) and headers[col_index]}
        return
    if hasattr(handle, "read"):
        data = handle.read()
    else:
        data = handle
    if isinstance(data, bytes):
        text = data.decode("utf-8-sig")
    else:
        text = str(data)
    reader = csv.DictReader(io.StringIO(text))
    headers = [str(field or "").strip() for field in (reader.fieldnames or []) if str(field or "").strip()]
    if len(headers) > MAX_IMPORT_COLUMNS:
        raise ValidationError({"source_file": f"Import exceeds the maximum of {MAX_IMPORT_COLUMNS} columns."})
    for index, row in enumerate(reader, start=1):
        if limit is None and index > MAX_IMPORT_ROWS:
            raise ValidationError({"source_file": f"Import exceeds the maximum of {MAX_IMPORT_ROWS} rows."})
        if limit is not None and index > limit:
            break
        yield {str(key or "").strip(): str(value or "").strip() for key, value in row.items() if str(key or "").strip()}


def _validate_uploaded_file(uploaded_file) -> None:
    if uploaded_file is None:
        raise ValidationError({"source_file": "A source file is required."})
    size = int(getattr(uploaded_file, "size", 0) or 0)
    if size <= 0:
        raise ValidationError({"source_file": "Uploaded file is empty."})
    if size > MAX_IMPORT_FILE_BYTES:
        raise ValidationError({"source_file": f"File exceeds the maximum size of {MAX_IMPORT_FILE_BYTES} bytes."})


def _finalize_job(job: CrmImportJob) -> None:
    summary = {
        "created": job.row_results.filter(status=CrmImportRowStatus.CREATED).count(),
        "updated": job.row_results.filter(status=CrmImportRowStatus.UPDATED).count(),
        "duplicates": job.row_results.filter(status=CrmImportRowStatus.DUPLICATE).count(),
        "failed": job.row_results.filter(status=CrmImportRowStatus.FAILED).count(),
        "skipped": job.row_results.filter(status=CrmImportRowStatus.SKIPPED).count(),
    }
    job.status = CrmImportJobStatus.SUCCEEDED if summary["failed"] == 0 else CrmImportJobStatus.FAILED
    job.finished_at = timezone.now()
    job.lease_expires_at = None
    job.summary = summary
    job.save(update_fields=["status", "finished_at", "lease_expires_at", "summary", "updated_at"])
    log_crm_activity(
        business_profile=job.business_profile,
        actor=job.initiated_by,
        actor_type=CrmActivityActor.IMPORT,
        activity_type=CrmActivityType.IMPORTED if job.status == CrmImportJobStatus.SUCCEEDED else CrmActivityType.OTHER,
        summary="Import job finished" if job.status == CrmImportJobStatus.SUCCEEDED else "Import job failed",
        detail=f"job={job.id} summary={summary}",
    )


def _handle_job_failure(job: CrmImportJob, exc: Exception) -> None:
    now = timezone.now()
    job.refresh_from_db(fields=["attempt_count", "max_attempts", "summary"])
    job.error_detail = str(exc)
    job.lease_expires_at = None
    if job.attempt_count >= job.max_attempts:
        job.status = CrmImportJobStatus.FAILED
        job.finished_at = now
        job.run_after = None
        job.save(update_fields=["status", "finished_at", "run_after", "lease_expires_at", "error_detail", "updated_at"])
        log_crm_activity(
            business_profile=job.business_profile,
            actor=job.initiated_by,
            actor_type=CrmActivityActor.IMPORT,
            activity_type=CrmActivityType.OTHER,
            summary="Import job failed",
            detail=f"job={job.id} error={job.error_detail}",
        )
        return
    job.status = CrmImportJobStatus.QUEUED
    job.run_after = now + timedelta(minutes=RETRY_BACKOFF_MINUTES)
    job.save(update_fields=["status", "run_after", "lease_expires_at", "error_detail", "updated_at"])
    log_crm_activity(
        business_profile=job.business_profile,
        actor=job.initiated_by,
        actor_type=CrmActivityActor.IMPORT,
        activity_type=CrmActivityType.OTHER,
        summary="Import job rescheduled",
        detail=f"job={job.id} error={job.error_detail}",
    )
