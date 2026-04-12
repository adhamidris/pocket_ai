from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import timedelta
from pathlib import Path
from typing import Any, Iterable, Sequence

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import F, Q
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import BusinessProfile, User
from .models import (
    CrmActivityActor,
    CrmActivityType,
    CrmCompany,
    CrmContact,
    CrmDuplicateSuggestion,
    CrmDuplicateSuggestionStatus,
    CrmExternalIdentity,
    CrmFieldDefinition,
    CrmImportJob,
    CrmImportJobStatus,
    CrmImportRowResult,
    CrmImportRowStatus,
    CrmImportSourceFile,
    CrmImportSourceFormat,
    CrmImportTemplate,
    CrmRecordType,
)
from .services import create_company, create_contact, log_crm_activity, update_company, update_contact

MAX_IMPORT_FILE_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 5000
MAX_IMPORT_COLUMNS = 200
PREVIEW_ROW_LIMIT = 20
PREVIEW_SAMPLE_LIMIT = 5
RETRY_BACKOFF_MINUTES = 1


@dataclass(frozen=True)
class ImportPreview:
    columns: Sequence[str]
    sample_rows: Sequence[dict[str, str]]
    total_rows: int


def detect_source_format(filename: str) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".xlsx":
        return CrmImportSourceFormat.XLSX
    if suffix == ".csv":
        return CrmImportSourceFormat.CSV
    raise ValidationError({"source_file": "Only CSV and XLSX imports are supported."})


def store_import_source_file(*, business_profile: BusinessProfile, uploaded_by: User | None, uploaded_file) -> CrmImportSourceFile:
    _validate_uploaded_file(uploaded_file)
    file_format = detect_source_format(getattr(uploaded_file, "name", ""))
    preview = build_preview(uploaded_file, file_format=file_format)
    uploaded_file.seek(0)
    source = CrmImportSourceFile.objects.create(
        business_profile=business_profile,
        uploaded_by=uploaded_by,
        original_filename=getattr(uploaded_file, "name", "crm-import"),
        file_format=file_format,
        file=uploaded_file,
        file_size_bytes=int(getattr(uploaded_file, "size", 0) or 0),
        column_snapshot=list(preview.columns),
        sample_rows=list(preview.sample_rows),
    )
    return source


def build_preview(uploaded_file, *, file_format: str) -> ImportPreview:
    rows = list(_load_rows(uploaded_file, file_format=file_format, limit=PREVIEW_ROW_LIMIT))
    columns = tuple(rows[0].keys()) if rows else tuple()
    return ImportPreview(columns=columns, sample_rows=tuple(rows[:PREVIEW_SAMPLE_LIMIT]), total_rows=len(rows))


def save_import_template(*, business_profile: BusinessProfile, created_by: User | None, name: str, mapping: dict[str, Any]) -> CrmImportTemplate:
    template, _created = CrmImportTemplate.objects.update_or_create(
        business_profile=business_profile,
        name=name.strip(),
        defaults={"mapping": mapping or {}, "created_by": created_by},
    )
    return template


def infer_default_mapping(columns: Sequence[str]) -> dict[str, Any]:
    return _infer_default_mapping(columns)


def enqueue_import_job(*, business_profile: BusinessProfile, source_file: CrmImportSourceFile, initiated_by: User | None, mapping: dict[str, Any], template: CrmImportTemplate | None = None) -> CrmImportJob:
    effective_mapping = mapping or (template.mapping if template else {}) or _infer_default_mapping(source_file.column_snapshot)
    job = CrmImportJob.objects.create(
        business_profile=business_profile,
        source_file=source_file,
        initiated_by=initiated_by,
        template=template,
        mapping=effective_mapping,
        run_after=timezone.now(),
    )
    log_crm_activity(
        business_profile=business_profile,
        actor=initiated_by,
        activity_type=CrmActivityType.OTHER,
        summary="Import job queued",
        detail=f"job={job.id} source={source_file.id}",
    )
    return job


def process_next_job() -> CrmImportJob | None:
    now = timezone.now()
    with transaction.atomic():
        job = (
            CrmImportJob.objects.select_for_update(skip_locked=True)
            .filter(Q(status=CrmImportJobStatus.QUEUED) | Q(status=CrmImportJobStatus.RUNNING, lease_expires_at__lte=now))
            .filter(Q(run_after__isnull=True) | Q(run_after__lte=now))
            .filter(attempt_count__lt=F("max_attempts"))
            .order_by("created_at")
            .first()
        )
        if job is None:
            return None
        job.status = CrmImportJobStatus.RUNNING
        job.started_at = now
        job.lease_expires_at = now + timedelta(minutes=5)
        job.attempt_count += 1
        job.error_detail = ""
        job.save(update_fields=["status", "started_at", "lease_expires_at", "attempt_count", "error_detail", "updated_at"])
    log_crm_activity(
        business_profile=job.business_profile,
        actor=job.initiated_by,
        actor_type=CrmActivityActor.IMPORT,
        activity_type=CrmActivityType.OTHER,
        summary="Import job started",
        detail=f"job={job.id} attempt={job.attempt_count}",
    )
    try:
        _process_job(job)
    except Exception as exc:
        _handle_job_failure(job, exc)
    return job


def _process_job(job: CrmImportJob) -> None:
    source = job.source_file
    existing_results = {
        row_result.row_number: row_result.status
        for row_result in job.row_results.all().only("row_number", "status")
    }
    with source.file.open("rb") as handle:
        for row_number, row in enumerate(_load_rows(handle, file_format=source.file_format), start=1):
            if existing_results.get(row_number) in {
                CrmImportRowStatus.CREATED,
                CrmImportRowStatus.UPDATED,
                CrmImportRowStatus.DUPLICATE,
                CrmImportRowStatus.SKIPPED,
            }:
                continue
            try:
                with transaction.atomic():
                    status, messages, contact_id, company_id, duplicate_of_id = _process_row(job=job, row_number=row_number, row=row)
                    CrmImportRowResult.objects.update_or_create(
                        job=job,
                        row_number=row_number,
                        defaults={
                            "status": status,
                            "raw_data": row,
                            "messages": messages,
                            "contact_id": contact_id,
                            "company_id": company_id,
                            "duplicate_of_id": duplicate_of_id,
                        },
                    )
            except Exception as exc:
                CrmImportRowResult.objects.update_or_create(
                    job=job,
                    row_number=row_number,
                    defaults={
                        "status": CrmImportRowStatus.FAILED,
                        "raw_data": row,
                        "messages": [str(exc)],
                        "contact_id": None,
                        "company_id": None,
                        "duplicate_of_id": None,
                    },
                )
        _finalize_job(job)


def _process_row(*, job: CrmImportJob, row_number: int, row: dict[str, Any]) -> tuple[str, list[str], Any, Any, Any]:
    mapping = job.mapping or {}
    business = job.business_profile
    actor = job.initiated_by
    contact_payload = _extract_payload(row, mapping.get("contact") or {})
    company_payload = _extract_payload(row, mapping.get("company") or {})
    if not contact_payload and not company_payload:
        return CrmImportRowStatus.SKIPPED, ["No mapped values"], None, None, None

    company = None
    if company_payload:
        company_match = _match_company(business_profile=business, payload=company_payload)
        if company_match:
            company = update_company(business_profile=business, actor=actor, company=company_match, payload=company_payload)
            company_status = CrmImportRowStatus.UPDATED
        else:
            company = create_company(business_profile=business, actor=actor, payload=company_payload)
            company_status = CrmImportRowStatus.CREATED
    else:
        company_status = CrmImportRowStatus.SKIPPED

    if not contact_payload:
        return company_status, [], None, getattr(company, "id", None), None

    external_identity = _external_identity_payload(mapping.get("contact") or {}, row)
    contact_match, duplicate_match = _match_contact(
        business_profile=business,
        payload=contact_payload,
        external_identity=external_identity,
    )
    if duplicate_match and not contact_match:
        reasons = []
        if contact_payload.get("primary_email"):
            reasons.append("exact_email")
        if contact_payload.get("primary_phone"):
            reasons.append("exact_phone")
        CrmDuplicateSuggestion.objects.update_or_create(
            business_profile=business,
            import_job=job,
            record_type=CrmRecordType.CONTACT,
            source_row_number=row_number,
            candidate_record_id=duplicate_match.id,
            defaults={
                "record_id": None,
                "incoming_snapshot": row,
                "match_reasons": reasons or ["exact_identifier"],
                "status": CrmDuplicateSuggestionStatus.OPEN,
            },
        )
        return CrmImportRowStatus.DUPLICATE, ["Duplicate suggestion created"], None, getattr(company, "id", None), duplicate_match.id
    if contact_match:
        contact = update_contact(business_profile=business, actor=actor, contact=contact_match, payload=contact_payload)
        status = CrmImportRowStatus.UPDATED
    else:
        contact = create_contact(business_profile=business, actor=actor, payload=contact_payload)
        status = CrmImportRowStatus.CREATED
    if company:
        _attach_company(contact=contact, company=company)
    if external_identity:
        CrmExternalIdentity.objects.update_or_create(
            business_profile=business,
            source_system=external_identity["source_system"],
            source_account_ref=external_identity["source_account_ref"],
            external_object_type=external_identity["external_object_type"],
            external_id=external_identity["external_id"],
            defaults={
                "record_type": CrmRecordType.CONTACT,
                "contact": contact,
                "external_label": external_identity["external_label"],
                "last_seen_at": timezone.now(),
            },
        )
    company_external_identity = _external_identity_payload(mapping.get("company") or {}, row)
    if company and company_external_identity:
        CrmExternalIdentity.objects.update_or_create(
            business_profile=business,
            source_system=company_external_identity["source_system"],
            source_account_ref=company_external_identity["source_account_ref"],
            external_object_type=company_external_identity["external_object_type"],
            external_id=company_external_identity["external_id"],
            defaults={
                "record_type": CrmRecordType.COMPANY,
                "company": company,
                "external_label": company_external_identity["external_label"],
                "last_seen_at": timezone.now(),
            },
        )
    return status, [], contact.id, getattr(company, "id", None), None


def _extract_payload(row: dict[str, Any], mapping: dict[str, Any]) -> dict[str, Any]:
    payload: dict[str, Any] = {}
    custom_fields: dict[str, Any] = {}
    for source_column, target in (mapping.get("columns") or {}).items():
        value = row.get(source_column)
        if value in (None, ""):
            continue
        if isinstance(target, str):
            payload[target] = value
        elif isinstance(target, dict) and target.get("kind") == "custom":
            custom_fields[str(target.get("key") or "")] = value
    if custom_fields:
        payload["custom_fields"] = custom_fields
    return payload


def _match_contact(
    *,
    business_profile: BusinessProfile,
    payload: dict[str, Any],
    external_identity: dict[str, str] | None = None,
) -> tuple[CrmContact | None, CrmContact | None]:
    if external_identity:
        identity_match = (
            CrmExternalIdentity.objects.filter(
                business_profile=business_profile,
                record_type=CrmRecordType.CONTACT,
                source_system=external_identity["source_system"],
                source_account_ref=external_identity["source_account_ref"],
                external_object_type=external_identity["external_object_type"],
                external_id=external_identity["external_id"],
                contact__isnull=False,
            )
            .select_related("contact")
            .first()
        )
        if identity_match and identity_match.contact:
            return identity_match.contact, None
    email = (payload.get("primary_email") or "").strip().lower()
    phone = "".join(ch for ch in str(payload.get("primary_phone") or "") if ch.isdigit())
    if email:
        match = CrmContact.objects.filter(business_profile=business_profile, primary_email_normalized=email).first()
        if match:
            return None, match
    if phone:
        match = CrmContact.objects.filter(business_profile=business_profile, primary_phone_normalized=phone).first()
        if match:
            return None, match
    return None, None


def _match_company(*, business_profile: BusinessProfile, payload: dict[str, Any]) -> CrmCompany | None:
    name = (payload.get("name") or "").strip()
    website = (payload.get("website") or "").strip().lower()
    if website:
        match = CrmCompany.objects.filter(business_profile=business_profile, website_domain_normalized__icontains=website.split("//")[-1].split("/")[0]).first()
        if match:
            return match
    if name:
        return CrmCompany.objects.filter(business_profile=business_profile, name__iexact=name).first()
    return None


def _attach_company(*, contact: CrmContact, company: CrmCompany) -> None:
    from .models import CrmContactCompanyLink

    CrmContactCompanyLink.objects.get_or_create(
        business_profile=contact.business_profile,
        contact=contact,
        company=company,
        defaults={"is_primary": True},
    )


def _external_identity_payload(mapping: dict[str, Any], row: dict[str, Any]) -> dict[str, str] | None:
    identity = mapping.get("external_identity")
    if not isinstance(identity, dict):
        return None
    source_column = identity.get("source_column")
    if not source_column:
        return None
    external_id = str(row.get(source_column) or "").strip()
    if not external_id:
        return None
    return {
        "source_system": str(identity.get("source_system") or "import").strip() or "import",
        "source_account_ref": str(identity.get("source_account_ref") or "").strip(),
        "external_object_type": str(identity.get("external_object_type") or "contact").strip() or "contact",
        "external_id": external_id,
        "external_label": str(identity.get("external_label") or source_column).strip(),
    }


def _infer_default_mapping(columns: Sequence[str]) -> dict[str, Any]:
    normalized = {str(column): _normalize_column_name(column) for column in columns or []}
    contact_columns: dict[str, Any] = {}
    company_columns: dict[str, Any] = {}
    external_identity: dict[str, Any] | None = None

    contact_targets = {
        "displayname": "display_name",
        "name": "display_name",
        "fullname": "display_name",
        "firstname": "first_name",
        "lastname": "last_name",
        "email": "primary_email",
        "primaryemail": "primary_email",
        "phone": "primary_phone",
        "mobile": "primary_phone",
        "phonenumber": "primary_phone",
        "title": "title",
    }
    company_targets = {
        "company": "name",
        "companyname": "name",
        "account": "name",
        "website": "website",
        "domain": "website",
        "companyphone": "primary_phone",
    }
    external_keys = {"id", "customerid", "contactid", "recordid", "externalid", "companyid", "accountid"}

    for source_column, key in normalized.items():
        if key in contact_targets:
            contact_columns[source_column] = contact_targets[key]
        elif key in company_targets:
            company_columns[source_column] = company_targets[key]
        elif key in external_keys and external_identity is None:
            external_identity = {
                "source_column": source_column,
                "source_system": "csv_import",
                "external_object_type": "contact",
                "external_label": source_column,
            }

    mapping: dict[str, Any] = {}
    if contact_columns or external_identity:
        mapping["contact"] = {"columns": contact_columns}
        if external_identity:
            mapping["contact"]["external_identity"] = external_identity
    if company_columns:
        mapping["company"] = {"columns": company_columns}
    return mapping


def _normalize_column_name(value: str) -> str:
    return "".join(ch for ch in str(value or "").strip().lower() if ch.isalnum())


def _load_rows(handle, *, file_format: str, limit: int | None = None) -> Iterable[dict[str, str]]:
    if file_format == CrmImportSourceFormat.XLSX:
        workbook = load_workbook(handle, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator, [])]
        if len(headers) > MAX_IMPORT_COLUMNS:
            raise ValidationError({"source_file": f"Import exceeds the maximum of {MAX_IMPORT_COLUMNS} columns."})
        for index, row in enumerate(iterator, start=1):
            if limit is None and index > MAX_IMPORT_ROWS:
                raise ValidationError({"source_file": f"Import exceeds the maximum of {MAX_IMPORT_ROWS} rows."})
            if limit is not None and index > limit:
                break
            yield {headers[col_index]: str(value or "").strip() for col_index, value in enumerate(row) if col_index < len(headers) and headers[col_index]}
        return
    if hasattr(handle, "read"):
        data = handle.read()
    else:
        data = handle
    if isinstance(data, bytes):
        text = data.decode("utf-8-sig")
    else:
        text = str(data)
    reader = csv.DictReader(io.StringIO(text))
    headers = [str(field or "").strip() for field in (reader.fieldnames or []) if str(field or "").strip()]
    if len(headers) > MAX_IMPORT_COLUMNS:
        raise ValidationError({"source_file": f"Import exceeds the maximum of {MAX_IMPORT_COLUMNS} columns."})
    for index, row in enumerate(reader, start=1):
        if limit is None and index > MAX_IMPORT_ROWS:
            raise ValidationError({"source_file": f"Import exceeds the maximum of {MAX_IMPORT_ROWS} rows."})
        if limit is not None and index > limit:
            break
        yield {str(key or "").strip(): str(value or "").strip() for key, value in row.items() if str(key or "").strip()}


def _validate_uploaded_file(uploaded_file) -> None:
    if uploaded_file is None:
        raise ValidationError({"source_file": "A source file is required."})
    size = int(getattr(uploaded_file, "size", 0) or 0)
    if size <= 0:
        raise ValidationError({"source_file": "Uploaded file is empty."})
    if size > MAX_IMPORT_FILE_BYTES:
        raise ValidationError({"source_file": f"File exceeds the maximum size of {MAX_IMPORT_FILE_BYTES} bytes."})


def _finalize_job(job: CrmImportJob) -> None:
    summary = {
        "created": job.row_results.filter(status=CrmImportRowStatus.CREATED).count(),
        "updated": job.row_results.filter(status=CrmImportRowStatus.UPDATED).count(),
        "duplicates": job.row_results.filter(status=CrmImportRowStatus.DUPLICATE).count(),
        "failed": job.row_results.filter(status=CrmImportRowStatus.FAILED).count(),
        "skipped": job.row_results.filter(status=CrmImportRowStatus.SKIPPED).count(),
    }
    job.status = CrmImportJobStatus.SUCCEEDED if summary["failed"] == 0 else CrmImportJobStatus.FAILED
    job.finished_at = timezone.now()
    job.lease_expires_at = None
    job.summary = summary
    job.save(update_fields=["status", "finished_at", "lease_expires_at", "summary", "updated_at"])
    log_crm_activity(
        business_profile=job.business_profile,
        actor=job.initiated_by,
        actor_type=CrmActivityActor.IMPORT,
        activity_type=CrmActivityType.IMPORTED if job.status == CrmImportJobStatus.SUCCEEDED else CrmActivityType.OTHER,
        summary="Import job finished" if job.status == CrmImportJobStatus.SUCCEEDED else "Import job failed",
        detail=f"job={job.id} summary={summary}",
    )


def _handle_job_failure(job: CrmImportJob, exc: Exception) -> None:
    now = timezone.now()
    job.refresh_from_db(fields=["attempt_count", "max_attempts", "summary"])
    job.error_detail = str(exc)
    job.lease_expires_at = None
    if job.attempt_count >= job.max_attempts:
        job.status = CrmImportJobStatus.FAILED
        job.finished_at = now
        job.run_after = None
        job.save(update_fields=["status", "finished_at", "run_after", "lease_expires_at", "error_detail", "updated_at"])
        log_crm_activity(
            business_profile=job.business_profile,
            actor=job.initiated_by,
            actor_type=CrmActivityActor.IMPORT,
            activity_type=CrmActivityType.OTHER,
            summary="Import job failed",
            detail=f"job={job.id} error={job.error_detail}",
        )
        return
    job.status = CrmImportJobStatus.QUEUED
    job.run_after = now + timedelta(minutes=RETRY_BACKOFF_MINUTES)
    job.save(update_fields=["status", "run_after", "lease_expires_at", "error_detail", "updated_at"])
    log_crm_activity(
        business_profile=job.business_profile,
        actor=job.initiated_by,
        actor_type=CrmActivityActor.IMPORT,
        activity_type=CrmActivityType.OTHER,
        summary="Import job rescheduled",
        detail=f"job={job.id} error={job.error_detail}",
    )
